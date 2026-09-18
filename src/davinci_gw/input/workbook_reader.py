"""按表头读取标准配置表，规范化字段并保留工作表及行号。"""

from __future__ import annotations

import re
import warnings
from collections.abc import Iterable, Sequence
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from davinci_gw.domain.models import (
    DirectRouteChange,
    OperationType,
    ReferenceDataEntry,
    SignalRouteChange,
    SourceLocation,
    ValidationIssue,
    ValidationSeverity,
    WorkbookData,
    WorkbookReadResult,
)
from davinci_gw.domain.route_keys import DirectRouteKey, SignalRouteKey

from .workbook_schema import (
    DIRECT_ADD_REQUIRED_FIELDS,
    DIRECT_HEADERS,
    DIRECT_IDENTITY_FIELDS,
    DIRECT_SHEET,
    REFERENCE_HEADERS,
    REFERENCE_SHEET,
    SIGNAL_HEADERS,
    SIGNAL_IDENTITY_FIELDS,
    SIGNAL_SHEET,
)

VERSION_PATTERN = re.compile(r"_v(\d+(?:\.\d+)*)\.xlsx$", re.IGNORECASE)
MAX_CAN_ID = 0x1FFFFFFF


def _clean_string(value: object) -> object:
    """仅裁剪字符串首尾空白，不改变业务标识的大小写。"""
    return value.strip() if isinstance(value, str) else value


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _display(value: object) -> str:
    return "空" if value is None else repr(value)


def _issue(
    path: Path, sheet: str | None, row: int | None, field: str | None,
    value: object, detail: str, code: str,
) -> ValidationIssue:
    location = SourceLocation(sheet, row)
    prefix = f"配置表“{path}”"
    if sheet:
        prefix += f"的“{sheet}”工作表"
    if row is not None:
        prefix += f"第{row}行"
    if field:
        prefix += f"的“{field}”"
    return ValidationIssue(
        code=code, message=f"{prefix}{detail}", file_path=path,
        location=location, field_name=field, actual_value=value,
    )


def parse_target_version(path: Path) -> tuple[str | None, ValidationIssue | None]:
    """从大小写不敏感的 *_vX.x.xlsx 文件名解析并规范化版本。"""
    match = VERSION_PATTERN.search(path.name)
    if not match:
        return None, ValidationIssue(
            code="WORKBOOK_VERSION_MISSING",
            message=f"配置表“{path}”无法从配置表文件名识别目标版本，请使用 *_vX.x.xlsx 格式。",
            file_path=path,
        )
    version = ".".join(str(int(part)) for part in match.group(1).split("."))
    return version, None


def _header_map(
    sheet: Worksheet, expected: Sequence[str], path: Path, issues: list[ValidationIssue],
) -> dict[str, int] | None:
    cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
    values = tuple(_cell_value(cell) for cell in cells)
    found: dict[str, int] = {}
    for index, value in enumerate(values):
        cleaned = _clean_string(value)
        if isinstance(cleaned, str) and cleaned:
            found[cleaned] = index
    missing = [header for header in expected if header not in found]
    for header in missing:
        issues.append(_issue(
            path, sheet.title, 1, header, None,
            "缺少必需表头，请补充该列后重试。", "WORKBOOK_HEADER_MISSING",
        ))
    return None if missing else found


def _cell_value(cell: Any) -> object:
    """保留显式空字符串（inlineStr）与从未填写单元格（None）的区别。"""
    if cell.value is None and getattr(cell, "data_type", None) == "inlineStr":
        return ""
    return cell.value


def _rows(sheet: Worksheet, headers: Sequence[str], mapping: dict[str, int]) -> Iterable[tuple[int, dict[str, object]]]:
    """遍历全部候选行；格式化空行会被忽略，中间空行不会终止扫描。"""
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = tuple(_cell_value(cell) for cell in cells)
        row = {header: _clean_string(values[index] if index < len(values) else None)
               for header, index in mapping.items() if header in headers}
        if any(not _is_blank(value) for value in row.values()):
            yield row_number, row


def _required(
    row: dict[str, object], fields: Sequence[str], path: Path, sheet: str,
    row_number: int, issues: list[ValidationIssue],
) -> bool:
    valid = True
    for field in fields:
        if _is_blank(row.get(field)):
            issues.append(_issue(
                path, sheet, row_number, field, row.get(field),
                "为空，该字段为必填项，请填写后重试。", "WORKBOOK_REQUIRED_VALUE",
            ))
            valid = False
    return valid


def _operation(
    value: object, path: Path, sheet: str, row_number: int, issues: list[ValidationIssue],
) -> OperationType | None:
    normalized = value.strip().upper() if isinstance(value, str) else value
    try:
        return OperationType(normalized)
    except (ValueError, TypeError):
        issues.append(_issue(
            path, sheet, row_number, "操作类型", value,
            f"为{_display(value)}，仅支持ADD或DELETE，请修改后重试。", "WORKBOOK_OPERATION_INVALID",
        ))
        return None


def _integer(
    value: object, *, minimum: int, maximum: int, path: Path, sheet: str,
    row_number: int, field: str, issues: list[ValidationIssue], optional: bool = False,
) -> int | None:
    if _is_blank(value) and optional:
        return None
    parsed: int | None = None
    if isinstance(value, int) and not isinstance(value, bool):
        parsed = value
    elif isinstance(value, float) and value.is_integer():
        parsed = int(value)
    elif isinstance(value, str):
        try:
            parsed = int(value, 16 if value.lower().startswith("0x") else 10)
        except ValueError:
            parsed = None
    if parsed is None or not minimum <= parsed <= maximum:
        issues.append(_issue(
            path, sheet, row_number, field, value,
            f"为{_display(value)}，必须是{minimum}至{maximum}范围内的整数，请修改后重试。",
            "WORKBOOK_INTEGER_INVALID",
        ))
        return None
    return parsed


def _decimal(
    value: object, path: Path, sheet: str, row_number: int, field: str,
    issues: list[ValidationIssue],
) -> Decimal | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        parsed = None
    else:
        try:
            parsed = (Decimal(int(value, 16))
                      if field == "超时值" and isinstance(value, str)
                      and value.lower().startswith("0x") else Decimal(str(value)))
        except (InvalidOperation, ValueError):
            parsed = None
    if parsed is None or not parsed.is_finite() or parsed < 0:
        issues.append(_issue(
            path, sheet, row_number, field, value,
            f"为{_display(value)}，非空时必须是非负数，请修改后重试。", "WORKBOOK_TIMEOUT_INVALID",
        ))
        return None
    return parsed


def _text(row: dict[str, object], field: str) -> str | None:
    value = row.get(field)
    return value if isinstance(value, str) and value else None


def _optional_text(row: dict[str, object], field: str) -> str | None:
    """返回可选字符串，同时保留显式空字符串。"""
    value = row.get(field)
    return value if isinstance(value, str) else None


def _read_references(
    sheet: Worksheet, mapping: dict[str, int], path: Path, issues: list[ValidationIssue],
) -> list[ReferenceDataEntry]:
    entries: list[ReferenceDataEntry] = []
    seen: dict[str, int] = {}
    for number, row in _rows(sheet, REFERENCE_HEADERS, mapping):
        if not _required(row, ("CAN通道名称",), path, sheet.title, number, issues):
            continue
        channel = _text(row, "CAN通道名称")
        if channel is None:
            continue
        if channel in seen:
            issues.append(_issue(
                path, sheet.title, number, "CAN通道名称", channel,
                f"为{_display(channel)}，与第{seen[channel]}行重复；CAN通道名称必须唯一，请删除或修改重复定义。",
                "REFERENCE_CHANNEL_DUPLICATE",
            ))
        else:
            seen[channel] = number
        entries.append(ReferenceDataEntry(
            channel, _optional_text(row, "CanIfTxBuffer名称"), _optional_text(row, "CanIfHrh名称"),
            SourceLocation(sheet.title, number),
        ))
    return entries


def _read_direct(
    sheet: Worksheet, mapping: dict[str, int], path: Path, issues: list[ValidationIssue],
) -> list[DirectRouteChange]:
    changes: list[DirectRouteChange] = []
    seen: dict[DirectRouteKey, dict[OperationType, int]] = {}
    for number, row in _rows(sheet, DIRECT_HEADERS, mapping):
        operation = _operation(row.get("操作类型"), path, sheet.title, number, issues)
        required = DIRECT_ADD_REQUIRED_FIELDS if operation is OperationType.ADD else DIRECT_IDENTITY_FIELDS
        required_ok = _required(row, required, path, sheet.title, number, issues)
        source_id = _integer(row.get("源网段报文CANID"), minimum=0, maximum=MAX_CAN_ID,
                             path=path, sheet=sheet.title, row_number=number,
                             field="源网段报文CANID", issues=issues,
                             optional=_is_blank(row.get("源网段报文CANID")))
        target_id = _integer(row.get("目标网段报文CANID"), minimum=0, maximum=MAX_CAN_ID,
                             path=path, sheet=sheet.title, row_number=number,
                             field="目标网段报文CANID", issues=issues,
                             optional=_is_blank(row.get("目标网段报文CANID")))
        source_length = _integer(row.get("源网段报文Length"), minimum=0, maximum=64,
                                 path=path, sheet=sheet.title, row_number=number,
                                 field="源网段报文Length", issues=issues,
                                 optional=_is_blank(row.get("源网段报文Length")))
        target_length = _integer(row.get("目标网段报文Length"), minimum=0, maximum=64,
                                 path=path, sheet=sheet.title, row_number=number,
                                 field="目标网段报文Length", issues=issues,
                                 optional=_is_blank(row.get("目标网段报文Length")))
        identity = tuple(_text(row, field) for field in (
            "源网段报文名称", "源网段CAN通道", "目标网段报文名称", "目标网段CAN通道"))
        if (operation is None or not required_ok or source_id is None or target_id is None
                or any(v is None for v in identity)):
            continue
        source_name, source_channel, target_name, target_channel = identity
        key = DirectRouteKey(source_name, source_id, source_channel, target_name, target_id, target_channel)
        previous = seen.setdefault(key, {})
        if operation in previous:
            issues.append(_issue(
                path, sheet.title, number, "操作类型", operation.value,
                f"与第{previous[operation]}行的同一路由操作重复；同一操作只能保留一条。",
                "DIRECT_ROUTE_DUPLICATE",
            ))
        else:
            # 同一身份允许一条 DELETE 与一条 ADD 组成替换对；执行阶段固定先删后增。
            previous[operation] = number
        changes.append(DirectRouteChange(
            operation, key, source_length, _text(row, "源网段报文类型"),
            _text(row, "源网段RxIndicationUL"), _text(row, "源网段报文Checksum使能"),
            _text(row, "源网段报文Dlc Check使能"), target_length,
            _text(row, "目标网段报文类型"), _text(row, "目标网段报文Checksum使能"),
            _text(row, "目标网段报文PnFilter使能"), _text(row, "目标网段报文Truncation使能"),
            _text(row, "路由Length Strategy功能选择"), SourceLocation(sheet.title, number),
        ))
    return changes


def _read_signals(
    sheet: Worksheet, mapping: dict[str, int], path: Path, issues: list[ValidationIssue],
) -> list[SignalRouteChange]:
    changes: list[SignalRouteChange] = []
    seen: dict[SignalRouteKey, dict[OperationType, int]] = {}
    for number, row in _rows(sheet, SIGNAL_HEADERS, mapping):
        operation = _operation(row.get("操作类型"), path, sheet.title, number, issues)
        required_ok = _required(row, SIGNAL_IDENTITY_FIELDS, path, sheet.title, number, issues)
        timeout_value = _decimal(row.get("超时值"), path, sheet.title, number, "超时值", issues)
        timeout_time = _decimal(row.get("超时时间"), path, sheet.title, number, "超时时间", issues)
        identity = tuple(_text(row, field) for field in SIGNAL_IDENTITY_FIELDS)
        if operation is None or not required_ok or any(value is None for value in identity):
            continue
        key = SignalRouteKey(*identity)
        previous = seen.setdefault(key, {})
        if operation in previous:
            issues.append(_issue(
                path, sheet.title, number, "操作类型", operation.value,
                f"与第{previous[operation]}行的同一路由操作重复；同一操作只能保留一条。",
                "SIGNAL_ROUTE_DUPLICATE",
            ))
        else:
            previous[operation] = number
        changes.append(SignalRouteChange(
            operation, key, _text(row, "字节序"), timeout_value, timeout_time,
            _text(row, "源信号组合名"), SourceLocation(sheet.title, number),
        ))
    return changes


def read_workbook(config_path: str | Path) -> WorkbookReadResult:
    """读取三张执行输入页并返回规范化数据和全部可定位契约问题。"""
    from .diagnostic_reader import read_diagnostics
    from .workbook_schema import DIAGNOSTIC_HEADERS, DIAGNOSTIC_SHEET

    path = Path(config_path).expanduser().resolve()
    issues: list[ValidationIssue] = []
    version, version_issue = parse_target_version(path)
    if version_issue:
        issues.append(version_issue)
    with warnings.catch_warnings():
        # 真实模板含 openpyxl 不支持的扩展验证元数据；只读且不保存，不会影响输入数据。
        warnings.filterwarnings(
            "ignore", message="Data Validation extension is not supported.*",
            category=UserWarning, module=r"openpyxl\.worksheet\._reader",
        )
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            required_sheets = {
                REFERENCE_SHEET: REFERENCE_HEADERS,
                DIRECT_SHEET: DIRECT_HEADERS,
                SIGNAL_SHEET: SIGNAL_HEADERS,
            }
            mappings: dict[str, dict[str, int]] = {}
            for name, headers in required_sheets.items():
                if name not in workbook.sheetnames:
                    issues.append(_issue(
                        path, name, None, None, None,
                        "不存在，该工作表是执行输入，请补充后重试。", "WORKBOOK_SHEET_MISSING",
                    ))
                    continue
                mapping = _header_map(workbook[name], headers, path, issues)
                if mapping is not None:
                    mappings[name] = mapping
            references = _read_references(workbook[REFERENCE_SHEET], mappings[REFERENCE_SHEET], path, issues) if REFERENCE_SHEET in mappings else []
            direct = _read_direct(workbook[DIRECT_SHEET], mappings[DIRECT_SHEET], path, issues) if DIRECT_SHEET in mappings else []
            signals = _read_signals(workbook[SIGNAL_SHEET], mappings[SIGNAL_SHEET], path, issues) if SIGNAL_SHEET in mappings else []
            diagnostics = []
            functional_ids = []
            if DIAGNOSTIC_SHEET in workbook.sheetnames:
                sheet = workbook[DIAGNOSTIC_SHEET]
                headers = {_clean_string(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))}
                if not {"诊断入口类型", "操作类型"} & headers:
                    # 旧版模板的示例页从未作为执行输入，不能替它推断 ADD。
                    issues.append(ValidationIssue(
                        code="DIAGNOSTIC_LEGACY_SHEET_IGNORED",
                        message="旧版诊断页缺少诊断入口类型和操作类型，本次不执行该页；请使用新版标准表提交诊断需求。",
                        severity=ValidationSeverity.WARNING, file_path=path,
                        location=SourceLocation(DIAGNOSTIC_SHEET, 1)))
                else:
                    mapping = _header_map(sheet, DIAGNOSTIC_HEADERS, path, issues)
                    if mapping is not None:
                        diagnostics = read_diagnostics(sheet, mapping, path, issues, references)
            if diagnostics and "诊断报文路由参数" in workbook.sheetnames:
                sheet = workbook["诊断报文路由参数"]
                headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
                if "通用功能寻址CANID" in headers:
                    column = headers.index("通用功能寻址CANID")
                    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                        value = values[column] if column < len(values) else None
                        if _is_blank(value):
                            continue
                        if isinstance(value, str) and not value.lower().startswith("0x"):
                            value = "0x" + value
                        parsed = _integer(value, minimum=0, maximum=MAX_CAN_ID, path=path, sheet=sheet.title,
                                          row_number=number, field="通用功能寻址CANID", issues=issues)
                        if parsed is not None:
                            functional_ids.append(parsed)
        finally:
            workbook.close()
    data = WorkbookData(path, version, tuple(references), tuple(direct), tuple(signals), tuple(diagnostics), tuple(functional_ids))
    return WorkbookReadResult(data, tuple(issues))
