import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from apply_signal_invalid_values import apply_invalid_values
from extract_candidates import generated_signal_records, signal_records, timeout_from_cycle
from review_requirements import review_manifest, review_route_inputs


class RequirementReviewTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        source = Workbook()
        history = source.active
        history.title = "History"
        history.cell(3, 2, "V1")
        history.cell(3, 3, "增加报文Example从FL_CANFD_EP网段路由到FL_CANFD_IC/FL_CANFD_DK网段")
        routing = source.create_sheet("Routing(FLZCU)")
        for column, value in {2: "Message", 3: "FL_CANFD_EP", 5: "Example", 7: "StandardCAN_FD", 8: "0x123", 9: 8, 16: "Example", 17: "StandardCAN_FD", 18: "0x123", 19: 8, 32: "√", 33: "√", 50: "√"}.items():
            routing.cell(5, column, value)
        self.source = self.root / "source.xlsx"
        source.save(self.source)
        source.close()
        self.matrix = self.root / "matrix.xlsx"
        matrix = Workbook()
        sheet = matrix.active
        sheet.title = "FLZCU_VCU_EP"
        for column, value in {4: "Msg Name", 12: "Signal Name", 27: "Invalid Value(Hex)"}.items():
            sheet.cell(1, column, value)
        sheet.cell(2, 4, "Source")
        sheet.cell(2, 12, "Value")
        sheet.cell(2, 27, "0x0")
        matrix.save(self.matrix)
        matrix.close()

    def tearDown(self):
        self.directory.cleanup()

    def manifest(self):
        routes = [{"source_message_name": "Example", "source_can_id": "0x123", "source_length": 8, "source_channel": "EPCAN", "target_message_name": "Example", "target_can_id": "0x123", "target_length": 8, "target_channel": channel, "operation": "ADD", "history_rows": [3]} for channel in ["ICCAN", "DKCAN"]]
        return {"source_workbook": str(self.source), "selected_versions": ["V1"], "project_mode": "AX", "routes": routes, "signal_routes": [], "requirement_review": [{"id": "3:0", "status": "output", "reason": "核对完整目标", "route_refs": ["routes:0", "routes:1"], "expected_route_refs": ["routes:0", "routes:1"]}]}

    def test_complete_requirement_passes(self):
        self.assertEqual(review_manifest(self.manifest()), [])

    def test_missing_parameter_blocks(self):
        manifest = self.manifest()
        manifest["routes"][0]["source_length"] = ""
        self.assertTrue(any("缺少必要参数" in error for error in review_manifest(manifest)))

    def test_missing_history_disposition_blocks(self):
        manifest = self.manifest()
        manifest["requirement_review"] = []
        self.assertTrue(any("缺少审核结论" in error for error in review_manifest(manifest)))

    def test_missing_target_blocks_even_with_matching_review_references(self):
        manifest = self.manifest()
        manifest["routes"].pop()
        manifest["requirement_review"][0].update(route_refs=["routes:0"], expected_route_refs=["routes:0"])
        self.assertTrue(any("全部有效目标" in error for error in review_manifest(manifest)))

    def test_false_exclusion_blocks(self):
        manifest = self.manifest()
        manifest["routes"] = []
        manifest["requirement_review"][0].update(status="excluded", route_refs=[], expected_route_refs=[])
        self.assertTrue(any("全部有效目标" in error for error in review_manifest(manifest)))

    def test_deferred_requirement_cannot_be_output(self):
        manifest = self.manifest()
        manifest["user_decisions"] = {"deferred_history_rows": [3]}
        self.assertTrue(any("暂缓" in error for error in review_manifest(manifest)))

    def test_reversed_operation_blocks(self):
        manifest = self.manifest()
        for route in manifest["routes"]:
            route["operation"] = "DELETE"
        self.assertTrue(any("全部有效目标" in error for error in review_manifest(manifest)))

    def test_missing_version_cannot_skip_review(self):
        manifest = self.manifest()
        del manifest["selected_versions"]
        self.assertTrue(any("缺少所选History版本" in error for error in review_manifest(manifest)))

    def update_source(self, update):
        from openpyxl import load_workbook
        source = load_workbook(self.source)
        update(source)
        source.save(self.source)
        source.close()

    def test_deletion_requires_all_branches(self):
        self.update_source(lambda book: book["History"].cell(3, 3, "删除报文Example从FL_CANFD_EP网段路由到FL_CANFD_IC/FL_CANFD_DK网段"))
        manifest = self.manifest()
        for route in manifest["routes"]:
            route["operation"] = "DELETE"
        self.assertEqual(review_manifest(manifest), [])
        manifest["routes"].pop()
        manifest["requirement_review"][0].update(route_refs=["routes:0"], expected_route_refs=["routes:0"])
        self.assertTrue(any("全部有效目标" in error for error in review_manifest(manifest)))

    def test_deleting_notes_is_not_deleting_a_route(self):
        self.update_source(lambda book: book["History"].cell(3, 3, "删除报文Example从FL_CANFD_EP网段路由到FL_CANFD_IC/FL_CANFD_DK的备注"))
        self.assertEqual(review_manifest(self.manifest()), [])

    def test_later_deletion_cancels_only_matching_branch(self):
        def update(book):
            book["History"].cell(4, 2, "V2")
            book["History"].cell(4, 3, "删除报文Example从FL_CANFD_EP网段路由到FL_CANFD_DK网段")
        self.update_source(update)
        manifest = self.manifest()
        manifest["selected_versions"].append("V2")
        manifest["routes"].pop()
        manifest["requirement_review"][0].update(route_refs=["routes:0"], expected_route_refs=["routes:0"])
        manifest["requirement_review"].append({"id": "4:0", "status": "cancelled", "reason": "后续删除抵消较早的DK新增", "route_refs": []})
        self.assertEqual(review_manifest(manifest), [])

    def test_explicit_can_id_does_not_select_same_name_other_id(self):
        def update(book):
            book["History"].cell(3, 3, "增加报文Example(0x123)从FL_CANFD_EP网段路由到FL_CANFD_IC/FL_CANFD_DK网段")
            routing = book["Routing(FLZCU)"]
            for column in range(1, 52):
                routing.cell(6, column).value = routing.cell(5, column).value
            routing.cell(6, 8, "0x456")
            routing.cell(6, 18, "0x456")
        self.update_source(update)
        self.assertEqual(review_manifest(self.manifest()), [])

    def test_fd_dm_demand_does_not_become_standard_dm(self):
        def update(book):
            book["History"].cell(3, 3, "增加报文Example从FL_CANFD_EP网段路由到FL_CANFD_DM网段")
            book["Routing(FLZCU)"].cell(5, 35, "√")
            book["Routing(FLZCU)"].cell(5, 36, "√")
        self.update_source(update)
        manifest = self.manifest()
        manifest["routes"] = []
        manifest["requirement_review"][0].update(status="excluded", reason="FD DM目标排除", route_refs=[], expected_route_refs=[])
        self.assertEqual(review_manifest(manifest), [])

    def test_project_transition_uses_actual_project_marks(self):
        manifest = self.manifest()
        manifest["project_mode"] = "AX_TO_AY"
        for route in manifest["routes"]:
            route.update(operation="DELETE", history_rows=[])
        manifest["transition_review"] = [{"route_ref": f"routes:{index}", "flzcu_row": 5} for index in range(2)]
        manifest["requirement_review"][0].update(status="excluded", reason="目标项目AY未启用新增需求", route_refs=[], expected_route_refs=[])
        self.assertEqual(review_manifest(manifest), [])

    def add_second_source(self):
        def update(book):
            routing = book["Routing(FLZCU)"]
            for column in range(1, 52):
                routing.cell(6, column).value = routing.cell(5, column).value
            routing.cell(6, 3, "FL_CANFD_CH")
            routing.cell(6, 33).value = None
        self.update_source(update)

    def test_multiple_sources_are_checked_against_full_requirements(self):
        self.add_second_source()
        manifest = self.manifest()
        self.assertTrue(any("多个源端汇入同一目标端" in error for error in review_manifest(manifest)))
        manifest["routes"].pop(0)
        manifest["requirement_review"][0].update(route_refs=["routes:0"], expected_route_refs=["routes:0"], reason="多源IC分支不做，仅保留无冲突DK分支")
        self.assertEqual(review_manifest(manifest), [])

    def test_confirmed_source_selection_is_preserved(self):
        self.add_second_source()
        manifest = self.manifest()
        manifest["user_decisions"] = {"excluded_source_rows": [{"flzcu_row": 6, "reason": "用户明确不采用CH源", "evidence": "本次用户回复"}]}
        self.assertEqual(review_manifest(manifest), [])

    def test_frzcu_cannot_be_output(self):
        from openpyxl import load_workbook
        source = load_workbook(self.source)
        source["History"].cell(3, 3, "FRZCU:增加报文Example从FR_CANFD_RE网段路由到FL_CANFD_GL网段")
        source.save(self.source)
        source.close()
        self.assertTrue(any("FRZCU" in error for error in review_manifest(self.manifest())))

    def test_every_send_type_uses_source_cycle(self):
        for send_type in ["Cyclic", "Cycle", "CE", ""]:
            for cycle, expected in [(10, 0.5), (20, 0.5), (21, 0.42), (100, 2), (101, 4)]:
                self.assertEqual(timeout_from_cycle(send_type, cycle)[1], expected)
        for cycle in [None, 0, -1, "bad", "nan", "inf"]:
            self.assertEqual(timeout_from_cycle("CE", cycle), (None, None))

    def test_blank_and_missing_values_do_not_get_defaults(self):
        signals = [{"operation": "ADD", "source_network": "EP", "source_message_name": "Source", "source_signal_name": name, "timeout_value": "0xFF"} for name in ["Value", "Absent"]]
        manifest = {"signal_routes": signals}
        apply_invalid_values(manifest, self.matrix)
        self.assertEqual(signals[0]["timeout_value"], "0x0")
        self.assertEqual(signals[1]["timeout_value"], "")
        self.assertEqual(signals[1]["timeout_value_status"], "missing_signal")

    def test_timeout_exception_requires_user_evidence(self):
        route = {"source_network": "EP", "source_message_name": "Source", "source_signal_name": "Value", "byte_order": "Intel", "target_network": "IC", "target_message_name": "Target", "target_signal_name": "Value", "source_cycle_time_ms": 100, "timeout_time": 1200, "timeout_value": "0x0", "operation": "ADD"}
        manifest = {"signal_routes": [route], "signal_matrix": str(self.matrix)}
        self.assertTrue(review_route_inputs(manifest))
        manifest["user_decisions"] = {"timeout_overrides": [{"source_key": ["EP", "Source", "Value"], "seconds": 1200, "reason": "用户指定", "evidence": "本次明确回复"}]}
        self.assertEqual(review_route_inputs(manifest), [])
        del manifest["user_decisions"]["timeout_overrides"][0]["evidence"]
        self.assertTrue(review_route_inputs(manifest))

    def test_pure_event_only_maps_and_inherits_message_type(self):
        from openpyxl import load_workbook
        matrix = load_workbook(self.matrix)
        sheet = matrix.active
        sheet.cell(2, 7, "Event")
        sheet.cell(2, 32, "OnWrite")
        sheet.cell(3, 12, "Next")
        sheet.cell(3, 27, "0xFF")
        sheet.cell(4, 4, "Combined")
        sheet.cell(4, 7, "CE")
        sheet.cell(4, 12, "Value")
        sheet.cell(4, 27, "0xFF")
        sheet.cell(4, 32, "OnChangeWithRepetition")
        matrix.save(self.matrix)
        matrix.close()
        signals = [{"source_network": "EP", "source_message_name": message, "source_signal_name": signal, "byte_order": "Intel", "target_network": "IC", "target_message_name": "Target", "target_signal_name": signal, "source_cycle_time_ms": 0 if message == "Source" else 100, "timeout_time": 2, "timeout_value": "0x0", "operation": "ADD"} for message, signal in [("Source", "Value"), ("Source", "Next"), ("Combined", "Value")]]
        manifest = {"signal_routes": signals}
        apply_invalid_values(manifest, self.matrix)
        self.assertEqual([route["routing_mode"] for route in signals], ["MAPPING_ONLY", "MAPPING_ONLY", "STANDARD"])
        self.assertTrue(all(route["timeout_time"] == "" and route["timeout_value"] == "" for route in signals[:2]))
        self.assertEqual(signals[2]["timeout_time"], 2)
        self.assertEqual(signals[2]["timeout_value"], "0xFF")
        self.assertEqual(review_route_inputs(manifest), [])
        signals[0]["timeout_time"] = 0.5
        self.assertTrue(any("纯Event只能做Mapping" in error for error in review_route_inputs(manifest)))

    def test_mapping_only_cannot_be_faked_for_a_cyclic_source(self):
        route = {"source_network": "EP", "source_message_name": "Source", "source_signal_name": "Value", "target_network": "IC", "target_message_name": "Target", "target_signal_name": "Value", "timeout_time": "", "timeout_value": "", "operation": "ADD", "routing_mode": "MAPPING_ONLY"}
        errors = review_route_inputs({"signal_routes": [route], "signal_matrix": str(self.matrix)})
        self.assertTrue(any("模式与源矩阵不一致" in error for error in errors))

    def test_lin_signal_is_rejected(self):
        errors = review_route_inputs({"signal_routes": [{"source_network": "EP", "target_network": "PSMM", "source_cycle_time_ms": 100, "timeout_time": 2}]})
        self.assertTrue(any("LIN" in error for error in errors))

    def test_generated_flag_is_retained_for_review(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(5, 15, "Example_Timeout_Flag")
        sheet.cell(5, 16, "Target")
        sheet.cell(5, 22, 12)
        sheet.cell(5, 23, 2)
        sheet.cell(5, 32, "√")
        self.assertEqual(signal_records(sheet), [])
        self.assertEqual(generated_signal_records(sheet)[0]["target_start_bit"], 12)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
