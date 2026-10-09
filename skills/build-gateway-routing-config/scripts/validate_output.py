#!/usr/bin/env python3
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from doip_mapping import HEADERS as DOIP_HEADERS, mapping_key, mapping_values, normalize_mapping


HEADERS = [
    "源网段报文名称",
    "源网段报文CANID",
    "源网段报文Length",
    "源网段报文类型",
    "源网段CAN通道",
    "源网段RxIndicationUL",
    "源网段报文Checksum使能",
    "源网段报文Dlc Check使能",
    "目标网段报文名称",
    "目标网段报文CANID",
    "目标网段报文Length",
    "目标网段报文类型",
    "目标网段CAN通道",
    "目标网段报文Checksum使能",
    "目标网段报文PnFilter使能",
    "目标网段报文Truncation使能",
    "路由Length Strategy功能选择",
    "操作类型",
]

SIGNAL_HEADERS = [
    "源网段",
    "源报文名",
    "源信号名",
    "字节序",
    "超时值",
    "超时时间",
    "源信号组合名",
    "目标网段",
    "目标报文名",
    "目标信号名",
    "操作类型",
]

DIAGNOSTIC_HEADERS = [
    "诊断请求端报文名称",
    "诊断请求端CANID_REQ",
    "诊断请求端接收报文类型",
    "诊断请求端CANID_RES",
    "诊断请求端发送报文类型",
    "诊断请求端Length",
    "诊断请求端CAN通道",
    "诊断应答端报文名称",
    "诊断应答端CANID_REQ",
    "诊断应答端发送报文类型",
    "诊断应答端CANID_RES",
    "诊断应答端接收报文类型",
    "诊断应答端Length",
    "诊断应答端CAN通道",
    "诊断请求端N_As",
    "诊断请求端N_Bs",
    "诊断请求端N_Cs",
    "诊断请求端N_Ar",
    "诊断请求端N_Br",
    "诊断请求端N_Cr",
    "诊断请求端BlockSize",
    "诊断请求端STmin",
    "诊断应答端N_As",
    "诊断应答端N_Bs",
    "诊断应答端N_Cs",
    "诊断应答端N_Ar",
    "诊断应答端N_Br",
    "诊断应答端N_Cr",
    "诊断应答端BlockSize",
    "诊断应答端STmin",
    "诊断入口类型",
    "操作类型",
]

CLASSIC_DIAGNOSTIC_CHANNELS = {"BDCAN", "DMCAN", "DGCAN"}


def operation(route):
    value = comparable(route.get("operation")).upper()
    return value


def find_sheet(workbook, expected):
    normalized = expected.strip().casefold()
    for name in workbook.sheetnames:
        if name.strip().casefold() == normalized:
            return workbook[name]
    raise KeyError(f"Missing worksheet: {expected}")


def expected_row(route):
    target_channel = str(route["target_channel"])
    target_type = "STANDARD_CAN" if target_channel in {"BDCAN", "DMCAN"} else "STANDARD_FD_CAN"
    return [
        str(route["source_message_name"]),
        str(route["source_can_id"]),
        route["source_length"],
        "STANDARD_CAN",
        str(route["source_channel"]),
        "PDUR",
        None,
        "Enable",
        str(route["target_message_name"]),
        str(route["target_can_id"]),
        route["target_length"],
        target_type,
        target_channel,
        None,
        None,
        "Enable",
        "IGNORE",
        operation(route),
    ]


def comparable(value):
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def expected_signal_row(route):
    return [
        str(route["source_network"]),
        str(route["source_message_name"]),
        str(route["source_signal_name"]),
        str(route.get("byte_order", "")),
        route.get("timeout_value", ""),
        route.get("timeout_time", ""),
        str(route.get("source_signal_combined_name", "")),
        str(route["target_network"]),
        str(route["target_message_name"]),
        str(route["target_signal_name"]),
        operation(route),
    ]


def expected_diagnostic_row(route):
    return [
        route.get("request_message_name", ""),
        route.get("request_can_id_req", ""),
        route.get("request_rx_type", ""),
        route.get("request_can_id_res", ""),
        route.get("request_tx_type", ""),
        route.get("request_length", ""),
        route.get("request_channel", ""),
        route.get("response_message_name", ""),
        route.get("response_can_id_req", ""),
        route.get("response_tx_type", ""),
        route.get("response_can_id_res", ""),
        route.get("response_rx_type", ""),
        route.get("response_length", ""),
        route.get("response_channel", ""),
        route.get("request_n_as", ""),
        route.get("request_n_bs", ""),
        route.get("request_n_cs", ""),
        route.get("request_n_ar", ""),
        route.get("request_n_br", ""),
        route.get("request_n_cr", ""),
        route.get("request_block_size", ""),
        route.get("request_st_min", ""),
        route.get("response_n_as", ""),
        route.get("response_n_bs", ""),
        route.get("response_n_cs", ""),
        route.get("response_n_ar", ""),
        route.get("response_n_br", ""),
        route.get("response_n_cr", ""),
        route.get("response_block_size", ""),
        route.get("response_st_min", ""),
        str(route.get("entry_type", "")).strip().upper(),
        operation(route),
    ]


def diagnostic_channel_settings(channel):
    if comparable(channel).upper() in CLASSIC_DIAGNOSTIC_CHANNELS:
        return "STANDARD_NO_FD_CAN", "STANDARD_CAN", "8"
    return "STANDARD_FD_CAN", "STANDARD_FD_CAN", "64"


def require_fields(route, fields, row, errors):
    missing = [field for field in fields if comparable(route.get(field)) == ""]
    if missing:
        errors.append(f"Diagnostic route row {row} missing fields: {', '.join(missing)}")


def computed_timeout_seconds(send_type, cycle_time_ms):
    if "cyclic" not in comparable(send_type).casefold():
        return ""
    try:
        cycle = float(cycle_time_ms)
    except (TypeError, ValueError):
        return ""
    if cycle <= 0:
        return ""
    if cycle <= 20:
        timeout_ms = 500
    elif cycle <= 100:
        timeout_ms = cycle * 20
    else:
        timeout_ms = 4000
    return comparable(timeout_ms / 1000)


def main():
    parser = argparse.ArgumentParser(description="Validate generated direct-routing workbook.")
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    args = parser.parse_args()

    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    sheet = find_sheet(workbook, "直接报文路由")
    signal_sheet = find_sheet(workbook, "信号路由")
    diagnostic_sheet = find_sheet(workbook, "诊断报文路由")
    doip_sheet = find_sheet(workbook, "DoIP_to_CAN")
    reference_sheet = find_sheet(workbook, "引用数据")
    reference_channels = {
        comparable(reference_sheet.cell(row=row, column=1).value).upper()
        for row in range(2, reference_sheet.max_row + 1)
        if comparable(reference_sheet.cell(row=row, column=1).value)
    }
    errors = []
    if [doip_sheet.cell(1, column).value for column in range(1, 8)] != list(DOIP_HEADERS):
        errors.append("DoIP_to_CAN 表头与七列契约不一致")
    mappings = manifest.get("doip_can_mappings", [])
    if not isinstance(mappings, list):
        errors.append("Manifest doip_can_mappings must be a list")
        mappings = []
    doip_keys = set()
    doip_rows = list(doip_sheet.iter_rows(min_row=2, max_col=7, values_only=True))
    for index, record in enumerate(mappings, 2):
        try:
            record = normalize_mapping(record, reference_channels)
            key = mapping_key(record)
            if key in doip_keys:
                errors.append(f"DoIP_to_CAN 第 {index} 行身份重复：{key}")
            doip_keys.add(key)
            values = doip_rows[index - 2] if index - 2 < len(doip_rows) else (None,) * 7
            actual = [comparable(value) for value in values]
            if actual != [comparable(v) for v in mapping_values(record)]:
                errors.append(f"DoIP_to_CAN 第 {index} 行与 manifest 不一致")
        except ValueError as error:
            errors.append(f"DoIP_to_CAN 第 {index} 行：{error}")
    for row in doip_rows[len(mappings):]:
        if any(comparable(value) for value in row):
            errors.append("DoIP_to_CAN 存在 manifest 之外的额外数据")
            break

    actual_headers = [sheet.cell(row=1, column=column).value for column in range(1, 19)]
    if actual_headers != HEADERS:
        errors.append("Direct-routing headers do not match the bundled template schema.")
    actual_signal_headers = [
        signal_sheet.cell(row=1, column=column).value for column in range(1, 12)
    ]
    if actual_signal_headers != SIGNAL_HEADERS:
        errors.append("Signal-routing headers do not match the bundled template schema.")
    actual_diagnostic_headers = [
        diagnostic_sheet.cell(row=1, column=column).value for column in range(1, 33)
    ]
    if actual_diagnostic_headers != DIAGNOSTIC_HEADERS:
        errors.append("Diagnostic-routing headers do not match the bundled template schema.")

    routes = manifest.get("routes", [])
    keys = set()
    for index, route in enumerate(routes, start=2):
        if operation(route) not in {"ADD", "DELETE"}:
            errors.append(f"Invalid direct-route operation at output row {index}: {route.get('operation')!r}")
        key = (
            comparable(route.get("source_message_name")),
            comparable(route.get("source_can_id")).casefold(),
            comparable(route.get("source_channel")),
            comparable(route.get("target_message_name")),
            comparable(route.get("target_can_id")).casefold(),
            comparable(route.get("target_channel")),
        )
        if key in keys:
            errors.append(f"Duplicate manifest route at output row {index}: {key}")
        keys.add(key)

        expected = expected_row(route)
        actual = [sheet.cell(row=index, column=column).value for column in range(1, 19)]
        for column, (expected_value, actual_value) in enumerate(zip(expected, actual), start=1):
            if comparable(expected_value) != comparable(actual_value):
                errors.append(
                    f"Row {index} column {column} mismatch: expected {expected_value!r}, got {actual_value!r}"
                )

    first_extra_row = len(routes) + 2
    if any(sheet.cell(row=first_extra_row, column=column).value not in (None, "") for column in range(1, 19)):
        errors.append(f"Unexpected data after manifest routes at row {first_extra_row}.")

    signal_routes = manifest.get("signal_routes", [])
    signal_keys = set()
    source_settings = {}
    for index, route in enumerate(signal_routes, start=2):
        if operation(route) not in {"ADD", "DELETE"}:
            errors.append(f"Invalid signal-route operation at output row {index}: {route.get('operation')!r}")
        source_key = (
            comparable(route.get("source_network")).casefold(),
            comparable(route.get("source_message_name")).casefold(),
            comparable(route.get("source_signal_name")).casefold(),
        )
        key = source_key + (
            comparable(route.get("target_network")).casefold(),
            comparable(route.get("target_message_name")).casefold(),
            comparable(route.get("target_signal_name")).casefold(),
        )
        if key in signal_keys:
            errors.append(f"Duplicate manifest signal route at output row {index}: {key}")
        signal_keys.add(key)

        settings = (
            comparable(route.get("byte_order")),
            comparable(route.get("timeout_value")),
            comparable(route.get("timeout_time")),
            comparable(route.get("source_signal_combined_name")),
        )
        previous = source_settings.setdefault(source_key, settings)
        if previous != settings:
            errors.append(f"Inconsistent source settings for signal route at output row {index}: {source_key}")

        if "source_send_type" in route or "source_cycle_time_ms" in route:
            expected_timeout = computed_timeout_seconds(
                route.get("source_send_type"), route.get("source_cycle_time_ms")
            )
            if expected_timeout != comparable(route.get("timeout_time")):
                errors.append(
                    f"Signal route row {index} timeout mismatch: expected {expected_timeout!r}, "
                    f"got {route.get('timeout_time')!r}"
                )
        source_sheet = comparable(route.get("timeout_value_source_sheet"))
        source_row = route.get("timeout_value_source_row")
        if not source_sheet or not isinstance(source_row, int) or source_row < 2:
            errors.append(
                f"Signal route row {index} lacks valid source-matrix evidence for timeout_value."
            )

        expected = expected_signal_row(route)
        actual = [signal_sheet.cell(row=index, column=column).value for column in range(1, 12)]
        for column, (expected_value, actual_value) in enumerate(zip(expected, actual), start=1):
            if comparable(expected_value) != comparable(actual_value):
                errors.append(
                    f"Signal row {index} column {column} mismatch: "
                    f"expected {expected_value!r}, got {actual_value!r}"
                )

    first_extra_signal_row = len(signal_routes) + 2
    if any(
        signal_sheet.cell(row=first_extra_signal_row, column=column).value not in (None, "")
        for column in range(1, 12)
    ):
        errors.append(
            f"Unexpected data after manifest signal routes at row {first_extra_signal_row}."
        )

    diagnostic_routes = manifest.get("diagnostic_routes", [])
    if not isinstance(diagnostic_routes, list):
        errors.append("Manifest diagnostic_routes must be a list.")
        diagnostic_routes = []
    diagnostic_keys = set()
    response_parameter_fields = [
        "response_n_as",
        "response_n_bs",
        "response_n_cs",
        "response_n_ar",
        "response_n_br",
        "response_n_cr",
        "response_block_size",
        "response_st_min",
    ]
    request_parameter_fields = [
        "request_n_as",
        "request_n_bs",
        "request_n_cs",
        "request_n_ar",
        "request_n_br",
        "request_n_cr",
        "request_block_size",
        "request_st_min",
    ]
    for index, route in enumerate(diagnostic_routes, start=2):
        entry_type = comparable(route.get("entry_type")).upper()
        if entry_type not in {"OBD_CAN", "OBD_ETH"}:
            errors.append(
                f"Invalid diagnostic entry_type at output row {index}: {route.get('entry_type')!r}"
            )
        if operation(route) not in {"ADD", "DELETE"}:
            errors.append(
                f"Invalid diagnostic operation at output row {index}: {route.get('operation')!r}"
            )
        key = (
            entry_type,
            operation(route),
            comparable(route.get("request_message_name")).casefold(),
            comparable(route.get("response_can_id_req")).casefold(),
            comparable(route.get("response_message_name")).casefold(),
            comparable(route.get("response_can_id_res")).casefold(),
            comparable(route.get("response_channel")).casefold(),
        )
        if key in diagnostic_keys:
            errors.append(f"Duplicate manifest diagnostic route at output row {index}: {key}")
        diagnostic_keys.add(key)

        require_fields(
            route,
            [
                "request_message_name",
                "response_message_name",
                "response_can_id_req",
                "response_tx_type",
                "response_can_id_res",
                "response_rx_type",
                "response_length",
                "response_channel",
                *response_parameter_fields,
            ],
            index,
            errors,
        )
        response_channel = comparable(route.get("response_channel")).upper()
        if response_channel and response_channel not in reference_channels:
            errors.append(
                f"Diagnostic route row {index} response channel is not in 引用数据: "
                f"{route.get('response_channel')!r}"
            )
        expected_rx, expected_tx, expected_length = diagnostic_channel_settings(
            route.get("response_channel")
        )
        actual_response = (
            comparable(route.get("response_rx_type")),
            comparable(route.get("response_tx_type")),
            comparable(route.get("response_length")),
        )
        if actual_response != (expected_rx, expected_tx, expected_length):
            errors.append(
                f"Diagnostic route row {index} response channel settings mismatch: "
                f"expected {(expected_rx, expected_tx, expected_length)!r}, got {actual_response!r}"
            )

        request_identity_fields = (
            "request_can_id_req", "request_can_id_res", "request_channel"
        )
        has_can_request = entry_type == "OBD_CAN" or any(
            comparable(route.get(field)) for field in request_identity_fields
        )
        if has_can_request:
            require_fields(
                route,
                [
                    "request_can_id_req",
                    "request_rx_type",
                    "request_can_id_res",
                    "request_tx_type",
                    "request_length",
                    "request_channel",
                    *request_parameter_fields,
                ],
                index,
                errors,
            )
            request_channel = comparable(route.get("request_channel")).upper()
            if request_channel and request_channel not in reference_channels:
                errors.append(
                    f"Diagnostic route row {index} request channel is not in 引用数据: "
                    f"{route.get('request_channel')!r}"
                )
            request_rx, request_tx, request_length = diagnostic_channel_settings(
                route.get("request_channel")
            )
            actual_request = (
                comparable(route.get("request_rx_type")),
                comparable(route.get("request_tx_type")),
                comparable(route.get("request_length")),
            )
            if actual_request != (request_rx, request_tx, request_length):
                errors.append(
                    f"Diagnostic route row {index} request channel settings mismatch: "
                    f"expected {(request_rx, request_tx, request_length)!r}, got {actual_request!r}"
                )
            for request_field, response_field in (
                ("request_can_id_req", "response_can_id_req"),
                ("request_can_id_res", "response_can_id_res"),
            ):
                if comparable(route.get(request_field)).casefold() != comparable(
                    route.get(response_field)
                ).casefold():
                    errors.append(
                        f"Diagnostic route row {index} CAN IDs differ between endpoints: "
                        f"{request_field} and {response_field}"
                    )
        elif entry_type == "OBD_ETH":
            forbidden_fields = [
                "request_can_id_req",
                "request_rx_type",
                "request_can_id_res",
                "request_tx_type",
                "request_length",
                "request_channel",
                *request_parameter_fields,
            ]
            populated = [
                field for field in forbidden_fields if comparable(route.get(field)) != ""
            ]
            if populated:
                errors.append(
                    f"Diagnostic route row {index} OBD_ETH must leave request-side CAN fields blank: "
                    f"{', '.join(populated)}"
                )

        expected = expected_diagnostic_row(route)
        actual = [
            diagnostic_sheet.cell(row=index, column=column).value
            for column in range(1, 33)
        ]
        for column, (expected_value, actual_value) in enumerate(zip(expected, actual), start=1):
            if comparable(expected_value) != comparable(actual_value):
                errors.append(
                    f"Diagnostic row {index} column {column} mismatch: "
                    f"expected {expected_value!r}, got {actual_value!r}"
                )

    first_extra_diagnostic_row = len(diagnostic_routes) + 2
    if any(
        diagnostic_sheet.cell(row=first_extra_diagnostic_row, column=column).value
        not in (None, "")
        for column in range(1, 33)
    ):
        errors.append(
            f"Unexpected data after manifest diagnostic routes at row {first_extra_diagnostic_row}."
        )

    result = {
        "workbook": str(args.workbook.resolve()),
        "route_count": len(routes),
        "signal_route_count": len(signal_routes),
        "diagnostic_route_count": len(diagnostic_routes),
        "doip_can_mapping_count": len(mappings),
        "valid": not errors,
        "errors": errors,
    }
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if not errors else 1)


if __name__ == "__main__":
    main()
