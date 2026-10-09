#!/usr/bin/env python3
import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


TARGET_COLUMNS = {
    28: ("AB", "FL_CANFD_DG", "DGCAN", True),
    29: ("AC", "FL_CANFD_EP", "EPCAN", True),
    30: ("AD", "FL_CANFD_PT", "PTCAN", True),
    31: ("AE", "FL_CANFD_CH", "CHCAN", True),
    32: ("AF", "FL_CANFD_IC", "ICCAN", True),
    33: ("AG", "FL_CANFD_DK", "DKCAN", True),
    34: ("AH", "FL_CAN_BD", "BDCAN", True),
    35: ("AI", "FL_CAN_DM", "DMCAN", True),
    36: ("AJ", "FL_CANFD_DM", "DMCAN", False),
    37: ("AK", "FL_CANFD_LC", "LCCAN", True),
    38: ("AL", "FL_CANFD_DA", "DACAN", True),
    39: ("AM", "FL_CANFD_SU", "SUCAN", True),
    40: ("AN", "FL_CANFD_GL", "GLCAN", True),
}
SIGNAL_TARGET_COLUMNS = {
    28: ("AB", "FL_CANFD_DG", "DG", True),
    29: ("AC", "FL_CANFD_EP", "EP", True),
    30: ("AD", "FL_CANFD_PT", "PT", True),
    31: ("AE", "FL_CANFD_CH", "CH", True),
    32: ("AF", "FL_CANFD_IC", "IC", True),
    33: ("AG", "FL_CANFD_DK", "DK", True),
    34: ("AH", "FL_CAN_BD", "BD", True),
    35: ("AI", "FL_CAN_DM", "DM", True),
    36: ("AJ", "FL_CANFD_DM", "DM", False),
    37: ("AK", "FL_CANFD_LC", "LC", True),
    38: ("AL", "FL_CANFD_DA", "DA", True),
    39: ("AM", "FL_CANFD_SU", "SU", True),
    40: ("AN", "FL_CANFD_GL", "GL", True),
    41: ("AO", "FL_LIN_TDL1", "TDL1", True),
    42: ("AP", "FL_LIN_TDL2", "TDL2", True),
    43: ("AQ", "FL_LIN_RLHS", "RLHS", True),
    44: ("AR", "FL_LIN_EBS", "EBS", True),
    45: ("AS", "FL_LIN_DDSP", "DDSP", True),
    46: ("AT", "FL_LIN_LSMM", "LSMM", True),
    47: ("AU", "FL_LIN_PSMM", "PSMM", True),
    48: ("AV", "FL_LIN_SRF", "SRF", True),
    49: ("AW", "FL_LIN_DLM", "DLM", True),
}
PROJECT_COLUMNS = {
    "AX": 50,
    "AY": 51,
}
PROJECT_MODES = ("AX", "AY", "AX_TO_AY", "AY_TO_AX")
KNOWN_CODES = {
    "DG", "EP", "PT", "CH", "IC", "DK", "BD", "DM", "LC", "DA", "SU", "GL",
    "TDL1", "TDL2", "RLHS", "EBS", "DDSP", "LSMM", "PSMM", "SRF", "DLM",
}


def find_sheet(workbook, expected):
    normalized = expected.strip().casefold()
    for name in workbook.sheetnames:
        if name.strip().casefold() == normalized:
            return workbook[name]
    raise KeyError(f"Missing worksheet: {expected}")


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def project_mark(value):
    raw = clean(value)
    symbols = re.findall(r"[√×xX]", raw)
    last_symbol = symbols[-1] if symbols else ""
    return {
        "raw": raw,
        "last_symbol": last_symbol,
        "active": last_symbol == "√",
    }


def project_marks(row):
    return {
        column: project_mark(row[index - 1])
        for column, index in PROJECT_COLUMNS.items()
    }


def transition_action(marks, mode):
    if mode == "AX_TO_AY":
        old_column, new_column = "AX", "AY"
    elif mode == "AY_TO_AX":
        old_column, new_column = "AY", "AX"
    else:
        return None
    old_active = marks[old_column]["active"]
    new_active = marks[new_column]["active"]
    if old_active == new_active:
        return None
    return "ADD" if new_active else "DELETE"


def normalize_version(value):
    text = clean(value).casefold().replace(" ", "")
    return text[1:] if text.startswith("v") else text


def normalize_can_id(value):
    text = clean(value)
    match = re.search(r"0[xX]([0-9A-Fa-f]+)", text)
    if match:
        return str(int(match.group(1), 16))
    if text.isdigit():
        return str(int(text))
    return ""


def source_channel(subnet):
    text = clean(subnet).upper()
    special = {
        "FL_CAN_BD": "BDCAN",
        "FL_CAN_DM": "DMCAN",
        "FL_CANFD_DM": "DMCAN",
    }
    if text in special:
        return special[text]
    match = re.fullmatch(r"FL_CANFD_([A-Z0-9]+)", text)
    return f"{match.group(1)}CAN" if match else ""


def network_code(subnet):
    text = clean(subnet).upper()
    match = re.fullmatch(r"FL_(?:CANFD|CAN|LIN)_([A-Z0-9]+)", text)
    return match.group(1) if match else ""


def numeric_value(value):
    if value is None or clean(value) == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else number


def timeout_from_cycle(send_type, cycle_time_ms):
    if "cyclic" not in clean(send_type).casefold():
        return None, None
    cycle = numeric_value(cycle_time_ms)
    if cycle is None or cycle <= 0:
        return None, None
    if cycle <= 20:
        timeout_ms = 500
    elif cycle <= 100:
        timeout_ms = cycle * 20
    else:
        timeout_ms = 4000
    timeout_seconds = timeout_ms / 1000
    if isinstance(timeout_seconds, float) and timeout_seconds.is_integer():
        timeout_seconds = int(timeout_seconds)
    return timeout_ms, timeout_seconds


def split_scopes(text):
    pattern = re.compile(r"\b(FLZCU|FRZCU)\s*[:：]", re.IGNORECASE)
    matches = list(pattern.finditer(text))
    if not matches:
        return [{"scope": "unknown", "text": text.strip()}]
    parts = []
    prefix = text[: matches[0].start()].strip()
    if prefix:
        parts.append({"scope": "unknown", "text": prefix})
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        parts.append({"scope": match.group(1).upper(), "text": text[match.end() : end].strip()})
    return [part for part in parts if part["text"]]


def action_flags(text):
    flags = []
    if re.search(r"新增|增加|添加", text):
        flags.append("add")
    if re.search(r"删除|取消|移除", text):
        flags.append("delete")
    if re.search(r"修改|改变|更改|调整|改为|改成|由.+改", text):
        flags.append("modify")
    return flags or ["unknown"]


def mentioned_codes(text):
    upper = text.upper()
    codes = set()
    for code in KNOWN_CODES:
        if re.search(rf"(?<![A-Z0-9_]){code}(?![A-Z0-9_])", upper):
            codes.add(code)
    for full in re.findall(r"FL_(?:CANFD|CAN|LIN)_([A-Z0-9]+)", upper):
        if full in KNOWN_CODES:
            codes.add(full)
    return sorted(codes)


def direct_records(sheet):
    records = []
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=5, max_col=51, values_only=True),
        start=5,
    ):
        routing_type = clean(row[1])
        if routing_type != "Message":
            continue
        source_frame = clean(row[6])
        target_frame = clean(row[16])
        target_legs = []
        for column_number, (letter, subnet, channel, allowed) in TARGET_COLUMNS.items():
            mark = clean(row[column_number - 1])
            if "√" in mark:
                target_legs.append(
                    {
                        "column": letter,
                        "subnet": subnet,
                        "channel": channel,
                        "allowed": allowed,
                    }
                )
        records.append(
            {
                "row": row_number,
                "source_subnet": clean(row[2]),
                "source_channel": source_channel(row[2]),
                "source_message": clean(row[4]),
                "source_frame": source_frame,
                "source_id": clean(row[7]),
                "source_id_normalized": normalize_can_id(row[7]),
                "source_length": row[8],
                "target_message": clean(row[15]),
                "target_frame": target_frame,
                "target_id": clean(row[17]),
                "target_id_normalized": normalize_can_id(row[17]),
                "target_length": row[18],
                "project_marks": project_marks(row),
                "lin_message_route": source_frame.upper() == "LIN" or target_frame.upper() == "LIN",
                "target_legs": target_legs,
            }
        )
    return records


def signal_records(sheet):
    records = []
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=5, max_col=51, values_only=True),
        start=5,
    ):
        if clean(row[1]).casefold() != "signal":
            continue
        target_legs = []
        for column_number, (letter, subnet, network, allowed) in SIGNAL_TARGET_COLUMNS.items():
            mark = clean(row[column_number - 1])
            if "√" in mark:
                target_legs.append(
                    {
                        "column": letter,
                        "subnet": subnet,
                        "network": network,
                        "allowed": allowed,
                    }
                )
        timeout_ms, timeout_seconds = timeout_from_cycle(row[9], row[10])
        records.append(
            {
                "row": row_number,
                "source_subnet": clean(row[2]),
                "source_network": network_code(row[2]),
                "source_message": clean(row[4]),
                "source_frame": clean(row[6]),
                "source_id": clean(row[7]),
                "source_id_normalized": normalize_can_id(row[7]),
                "source_length": row[8],
                "source_send_type": clean(row[9]),
                "source_cycle_time_ms": numeric_value(row[10]),
                "source_signal": clean(row[0]),
                "source_start_bit": numeric_value(row[11]),
                "source_signal_length": numeric_value(row[12]),
                "source_byte_order": clean(row[13]),
                "target_signal": clean(row[14]),
                "target_message": clean(row[15]),
                "target_frame": clean(row[16]),
                "target_id": clean(row[17]),
                "target_id_normalized": normalize_can_id(row[17]),
                "target_length": row[18],
                "timeout_ms": timeout_ms,
                "timeout_time": timeout_seconds,
                "timeout_value": "",
                "source_signal_combined_name": "",
                "project_marks": project_marks(row),
                "target_legs": target_legs,
            }
        )
    return records


def all_route_names_and_ids(sheet):
    names = set()
    ids = set()
    max_col = min(sheet.max_column, 40)
    for row in sheet.iter_rows(min_row=5, max_col=max_col, values_only=True):
        for index in (0, 4, 14, 15):
            if index < len(row) and clean(row[index]):
                names.add(clean(row[index]))
        for index in (7, 17):
            if index < len(row):
                normalized = normalize_can_id(row[index])
                if normalized:
                    ids.add(normalized)
    return names, ids


def version_groups(history):
    markers = []
    for row in range(3, history.max_row + 1):
        version = clean(history.cell(row=row, column=2).value)
        if version:
            markers.append({"version": version, "start_row": row})
    for index, marker in enumerate(markers):
        marker["end_row"] = (
            markers[index + 1]["start_row"] - 1 if index + 1 < len(markers) else history.max_row
        )
    return markers


def main():
    parser = argparse.ArgumentParser(description="Extract selected History text and FLZCU candidates.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--versions", nargs="+", required=True)
    parser.add_argument("--project-mode", choices=PROJECT_MODES, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    history = find_sheet(workbook, "History")
    flzcu = find_sheet(workbook, "Routing(FLZCU)")
    frzcu = find_sheet(workbook, "Routing(FRZCU)")

    requested = {normalize_version(value) for value in args.versions}
    groups = [group for group in version_groups(history) if normalize_version(group["version"]) in requested]
    missing = sorted(requested - {normalize_version(group["version"]) for group in groups})
    if missing:
        raise SystemExit(f"Versions not found: {', '.join(missing)}")

    records = direct_records(flzcu)
    signals = signal_records(flzcu)
    flzcu_names = {
        name
        for record in records
        for name in (record["source_message"], record["target_message"])
        if name
    }
    frzcu_names, frzcu_ids = all_route_names_and_ids(frzcu)
    flzcu_signal_names = {
        name
        for record in signals
        for name in (record["source_signal"], record["target_signal"])
        if name
    }
    flzcu_signal_message_names = {
        name
        for record in signals
        for name in (record["source_message"], record["target_message"])
        if name
    }

    units = []
    candidate_use = defaultdict(int)
    for group in groups:
        for history_row in range(group["start_row"], group["end_row"] + 1):
            text = clean(history.cell(row=history_row, column=3).value)
            if not text:
                continue
            for part in split_scopes(text):
                part_text = part["text"]
                folded = part_text.casefold()
                signal_only = "信号" in part_text
                project_only = "打点" in part_text
                excluded_before_matching = part["scope"] == "FRZCU"
                mentioned_names = sorted(
                    (name for name in flzcu_names if name.casefold() in folded),
                    key=lambda value: (-len(value), value),
                )
                ids = sorted(set(re.findall(r"0[xX][0-9A-Fa-f]+", part_text)))
                normalized_ids = {normalize_can_id(value) for value in ids}
                mentioned_signal_names = sorted(
                    (name for name in flzcu_signal_names if name.casefold() in folded),
                    key=lambda value: (-len(value), value),
                )
                mentioned_signal_message_names = sorted(
                    (name for name in flzcu_signal_message_names if name.casefold() in folded),
                    key=lambda value: (-len(value), value),
                )
                message_candidates = []
                signal_candidates = []
                if not excluded_before_matching and not signal_only:
                    for record in records:
                        name_match = any(
                            name.casefold() in {
                                record["source_message"].casefold(),
                                record["target_message"].casefold(),
                            }
                            for name in mentioned_names
                        )
                        id_match = bool(
                            normalized_ids
                            & {record["source_id_normalized"], record["target_id_normalized"]}
                        )
                        if name_match or id_match:
                            message_candidates.append(record)
                            candidate_use[("message", record["row"])] += 1
                if not excluded_before_matching and signal_only:
                    for record in signals:
                        signal_name_match = any(
                            name.casefold() in {
                                record["source_signal"].casefold(),
                                record["target_signal"].casefold(),
                            }
                            for name in mentioned_signal_names
                        )
                        message_name_match = any(
                            name.casefold() in {
                                record["source_message"].casefold(),
                                record["target_message"].casefold(),
                            }
                            for name in mentioned_signal_message_names
                        )
                        id_match = bool(
                            normalized_ids
                            & {record["source_id_normalized"], record["target_id_normalized"]}
                        )
                        if signal_name_match or (message_name_match and id_match) or (
                            id_match and not mentioned_signal_names
                        ):
                            signal_candidates.append(record)
                            candidate_use[("signal", record["row"])] += 1

                frzcu_hits = sorted(
                    name for name in frzcu_names if name and name.casefold() in folded
                )
                frzcu_id_hits = sorted(
                    value for value in ids if normalize_can_id(value) in frzcu_ids
                )
                units.append(
                    {
                        "version": group["version"],
                        "history_row": history_row,
                        "scope": part["scope"],
                        "text": part_text,
                        "actions": action_flags(part_text),
                        "signal_only": signal_only,
                        "routing_kind": "signal" if signal_only else "message",
                        "project_only": project_only,
                        "project_mode": args.project_mode,
                        "excluded_before_matching": excluded_before_matching,
                        "mentioned_names": mentioned_names,
                        "mentioned_signal_names": mentioned_signal_names,
                        "mentioned_signal_message_names": mentioned_signal_message_names,
                        "mentioned_ids": ids,
                        "mentioned_network_codes": mentioned_codes(part_text),
                        "frzcu_hits": frzcu_hits,
                        "frzcu_id_hits": frzcu_id_hits,
                        "candidate_flzcu_rows": message_candidates,
                        "candidate_message_rows": message_candidates,
                        "candidate_signal_rows": signal_candidates,
                    }
                )

    transition_candidates = {"message": [], "signal": []}
    if args.project_mode in {"AX_TO_AY", "AY_TO_AX"}:
        for kind, route_records in (("message", records), ("signal", signals)):
            for record in route_records:
                action = transition_action(record["project_marks"], args.project_mode)
                if action:
                    transition_candidates[kind].append(
                        {
                            "operation": action,
                            "candidate": record,
                        }
                    )

    payload = {
        "source_workbook": str(workbook_path),
        "selected_versions": [group["version"] for group in groups],
        "output_version": groups[-1]["version"].lstrip("Vv") if groups else "",
        "project_mode": args.project_mode,
        "project_transition_candidates": transition_candidates,
        "history_units": units,
        "summary": {
            "history_units": len(units),
            "units_with_candidates": sum(
                bool(unit["candidate_message_rows"] or unit["candidate_signal_rows"])
                for unit in units
            ),
            "units_with_message_candidates": sum(
                bool(unit["candidate_message_rows"]) for unit in units
            ),
            "units_with_signal_candidates": sum(
                bool(unit["candidate_signal_rows"]) for unit in units
            ),
            "explicit_frzcu_units": sum(unit["scope"] == "FRZCU" for unit in units),
            "units_with_frzcu_hits": sum(
                bool(unit["frzcu_hits"] or unit["frzcu_id_hits"]) for unit in units
            ),
        },
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(args.output.resolve())


if __name__ == "__main__":
    main()
