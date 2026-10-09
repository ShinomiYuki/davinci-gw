#!/usr/bin/env python3
import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

SHEET_PREFIX = "FLZCU_VCU_"
EXCLUDED_NETWORKS = {"DM_FD"}


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def header_name(value):
    return clean(value).splitlines()[0].casefold()


def signal_sheets(workbook):
    sheets = {}
    for worksheet in workbook.worksheets:
        if not worksheet.title.upper().startswith(SHEET_PREFIX):
            continue
        network = worksheet.title[len(SHEET_PREFIX) :].strip().upper()
        if network in EXCLUDED_NETWORKS:
            continue
        if network in sheets:
            raise ValueError(f"信号矩阵中网段 {network!r} 对应多个工作表。")
        sheets[network] = worksheet
    return sheets


def build_signal_index(worksheet):
    expected_headers = {
        4: "msg name",
        12: "signal name",
        27: "invalid value(hex)",
    }
    for column, expected in expected_headers.items():
        actual = header_name(worksheet.cell(row=1, column=column).value)
        if actual != expected:
            raise ValueError(
                f"工作表 {worksheet.title!r} 第 {column} 列表头应为 {expected!r}，"
                f"实际为 {actual!r}。"
            )

    index = defaultdict(list)
    current_message = ""
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, max_col=27, values_only=True),
        start=2,
    ):
        message = clean(row[3])
        if message:
            current_message = message
        signal = clean(row[11])
        if not current_message or not signal:
            continue
        invalid_value = clean(row[26])
        key = (current_message.casefold(), signal.casefold())
        index[key].append(
            {
                "row": row_number,
                "message": current_message,
                "signal": signal,
                "invalid_value": invalid_value,
            }
        )
    return index


def apply_invalid_values(manifest, matrix_path):
    signal_routes = manifest.get("signal_routes", [])
    if not isinstance(signal_routes, list):
        raise TypeError("Manifest signal_routes 必须是列表。")

    workbook = load_workbook(matrix_path, read_only=True, data_only=True)
    try:
        sheets = signal_sheets(workbook)
        indexes = {}
        errors = []
        resolved = 0
        nonblank = 0

        for position, route in enumerate(signal_routes, start=1):
            operation = clean(route.get("operation")).upper()
            if operation not in {"ADD", "DELETE"}:
                errors.append(f"signal_routes[{position}] 操作类型不是 ADD 或 DELETE。")
                continue

            network = clean(route.get("source_network")).upper()
            message = clean(route.get("source_message_name"))
            signal = clean(route.get("source_signal_name"))
            identity = f"{network}/{message}/{signal}"

            if network in EXCLUDED_NETWORKS:
                errors.append(f"signal_routes[{position}] {identity} 属于已排除的 DM_FD 网段。")
                continue
            worksheet = sheets.get(network)
            if worksheet is None:
                errors.append(
                    f"signal_routes[{position}] {identity} 找不到源网段工作表 "
                    f"{SHEET_PREFIX}{network}。"
                )
                continue
            if network not in indexes:
                indexes[network] = build_signal_index(worksheet)
            index = indexes[network]
            matches = index.get((message.casefold(), signal.casefold()), [])
            if not matches:
                errors.append(
                    f"signal_routes[{position}] {identity} 在源网段工作表 "
                    f"{worksheet.title!r} 中没有精确匹配。"
                )
                continue
            if len(matches) > 1:
                rows = ", ".join(str(match["row"]) for match in matches)
                errors.append(
                    f"signal_routes[{position}] {identity} 在源网段工作表 "
                    f"{worksheet.title!r} 中匹配到多行：{rows}。"
                )
                continue

            match = matches[0]
            route["timeout_value"] = match["invalid_value"]
            route["timeout_value_source_sheet"] = worksheet.title
            route["timeout_value_source_row"] = match["row"]
            resolved += 1
            if match["invalid_value"]:
                nonblank += 1

        if errors:
            raise ValueError("\n".join(errors))

        manifest["signal_matrix"] = str(matrix_path.resolve())
        return {
            "signal_route_count": len(signal_routes),
            "resolved_count": resolved,
            "nonblank_timeout_value_count": nonblank,
            "blank_timeout_value_count": resolved - nonblank,
        }
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(
        description="按源网段、源报文和源信号，从客户信号矩阵 AA 列补全超时值。"
    )
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--signal-matrix", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    manifest_path = args.manifest.resolve()
    matrix_path = args.signal_matrix.resolve()
    if not manifest_path.is_file():
        parser.error(f"Manifest 不存在：{manifest_path}")
    if not matrix_path.is_file():
        parser.error(f"信号矩阵不存在：{matrix_path}")

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    try:
        summary = apply_invalid_values(manifest, matrix_path)
    except (TypeError, ValueError) as error:
        parser.error(str(error))

    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    result = {
        "manifest": str(output),
        "signal_matrix": str(matrix_path),
        **summary,
    }
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
