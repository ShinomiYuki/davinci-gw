#!/usr/bin/env python3
"""从原始需求汇总 logicalAddress→Tester→CAN 映射，保留真实来源。"""

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from doip_mapping import TESTER_TYPES, hex_value, mapping_key, normalize_mapping
from extract_diagnostic_candidates import clean, is_missing, merged_values, reference_channels


def header_columns(sheet):
    for row in sheet.iter_rows(min_row=1, max_row=min(10, sheet.max_row)):
        columns = {}
        testers = []
        for cell in row:
            text = clean(cell.value)
            first = text.split("\n")[0].strip()
            if first in {"ECU DoIP Logical Address", "ECU DoIP逻辑地址"}:
                columns["logical"] = cell.column
            elif first in {"Diagnostic Tool's Logical Address", "OTA Master Logical Address", "诊断仪逻辑地址", "OTA Master逻辑地址"}:
                testers.append(cell.column)
            elif first in {"Destination Network Segment", "目标网段"}:
                columns["network"] = cell.column
            elif first in {"Diagnostic Request CANID", "诊断请求CANID"}:
                columns["request"] = cell.column
            elif first in {"Diagnostic Response CANID", "诊断响应CANID"}:
                columns["response"] = cell.column
            elif first in {"ECU Name", "控制器名称"}:
                columns["ecu"] = cell.column
        if {"logical", "network", "request", "response"} <= columns.keys() and testers:
            columns["testers"] = testers
            columns["header_row"] = row[0].row
            return columns
    return None


def extract_mappings(workbook, channels, *, selections=None, all_rows=False):
    if all_rows == bool(selections):
        raise ValueError("必须明确选择全量提取或逐 Sheet 项目列/行号，不能同时选择")
    if selections is not None and (not isinstance(selections, list) or not selections):
        raise ValueError("selections 必须是非空列表")
    if all_rows:
        selections = [{"sheet": s.title} for s in workbook if header_columns(s)]
        if not selections:
            raise ValueError("未找到具有 DoIP→CAN 映射表头的工作表")
    records = {}
    issues = []
    statistics = {"duplicates_removed": 0, "skipped_non_can": 0, "skipped_unmarked": 0}
    for selection in selections:
        if not isinstance(selection, dict):
            issues.append("每项 selection 必须是包含 sheet 的对象")
            continue
        name = selection.get("sheet")
        if name not in workbook.sheetnames:
            issues.append(f"缺少工作表：{name}")
            continue
        sheet = workbook[name]
        columns = header_columns(sheet)
        if columns is None:
            issues.append(f"{name} 缺少 DoIP→CAN 地址、Tester 或 CAN 映射表头")
            continue
        mark_column = selection.get("project_column")
        rows = selection.get("rows")
        if not all_rows:
            if bool(mark_column) == (rows is not None):
                issues.append(f"{name} 必须且只能指定 project_column 或 rows")
                continue
            if rows is not None and (not isinstance(rows, list) or not rows or any(
                    isinstance(r, bool) or not isinstance(r, int) or not columns["header_row"] < r <= sheet.max_row for r in rows)):
                issues.append(f"{name} rows 必须是表头之后的有效行号列表")
                continue
        if mark_column:
            try:
                mark_column = column_index_from_string(str(mark_column).upper())
                if mark_column <= max(columns["testers"] + [columns["response"], columns["network"], columns["logical"]]) or mark_column > sheet.max_column:
                    raise ValueError("不是项目打点列")
                if is_missing(sheet.cell(columns["header_row"], mark_column).value):
                    raise ValueError("项目打点列表头为空")
            except ValueError as error:
                issues.append(f"{name} 项目列无效：{error}")
                continue
        merged = merged_values(sheet)
        for row in sorted(set(rows)) if rows is not None else range(columns["header_row"] + 1, sheet.max_row + 1):
            def get(col):
                value = sheet.cell(row, col).value
                return merged.get((row, col)) if value is None else value
            if mark_column:
                mark = clean(get(mark_column))
                if mark not in {"", "●"}:
                    issues.append(f"{name} 第 {row} 行项目打点无效：{mark!r}")
                    continue
                if not mark:
                    statistics["skipped_unmarked"] += 1
                    continue
            network = clean(get(columns["network"])).upper()
            match = re.fullmatch(r"FL_CAN(?:FD)?_([A-Z]+)", network)
            channel = match.group(1) + "CAN" if match else network
            if not match and channel not in channels:
                statistics["skipped_non_can"] += 1
                continue
            logical = get(columns["logical"])
            tester_values = [(col, get(col)) for col in columns["testers"] if not is_missing(get(col))]
            if is_missing(logical) and not tester_values:
                continue
            try:
                logical = hex_value(logical, 0xFFFF, 4)
                if not tester_values:
                    raise ValueError("CAN 映射缺少 Tester")
                response = get(columns["response"])
                functional = is_missing(response)
                request = get(columns["request"])
                ecu = clean(get(columns["ecu"])) if "ecu" in columns else ""
                if functional and "功能" not in ecu:
                    raise ValueError("物理寻址 CAN 映射缺少响应 ID；不能推断功能寻址")
                for col, tester in tester_values:
                    tester = hex_value(tester, 0xFFFF, 4)
                    record = normalize_mapping({"Tester": tester, "logicalAddress": logical,
                        "RouterType": TESTER_TYPES.get(tester), "CANBus": channel,
                        "RequestCanId": request, "RespCanId": "" if functional else response,
                        "Functional": "YES" if functional else ""}, channels)
                    source = {"sheet": name, "row": row, "tester_column": col, "network": network}
                    key = mapping_key(record)
                    existing = records.get(key)
                    if existing is not None:
                        if any(existing[field] != record[field] for field in record):
                            raise ValueError(f"相同逻辑地址/Tester/CANBus 的 CAN 映射冲突：{key}，既有来源 {existing['sources']}")
                        existing["sources"].append(source)
                        statistics["duplicates_removed"] += 1
                    else:
                        record["sources"] = [source]
                        records[key] = record
            except ValueError as error:
                issues.append(f"{name} 第 {row} 行：{error}")
    mappings = sorted(records.values(), key=lambda r: (int(r["logicalAddress"], 16), r["CANBus"], int(r["Tester"], 16)))
    return mappings, statistics, issues


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    scope = parser.add_mutually_exclusive_group(required=True)
    scope.add_argument("--all", action="store_true", help="明确忽略项目打点，汇总所有可映射的 CAN 需求")
    scope.add_argument("--selection", type=Path, help="JSON 列表：每项指定 sheet 与 project_column 或 rows")
    parser.add_argument("--template", type=Path, default=Path(__file__).resolve().parent.parent / "assets" / "网关路由配置表空白模板.xlsx")
    parser.add_argument("--output-version", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        parser.error("输出 manifest 已存在，请使用新的路径")
    selections = json.loads(args.selection.read_text(encoding="utf-8")) if args.selection else None
    workbook = load_workbook(args.workbook, data_only=True)
    try:
        mappings, statistics, issues = extract_mappings(workbook, reference_channels(args.template), selections=selections, all_rows=args.all)
    finally:
        workbook.close()
    result = {"output_version": args.output_version, "source_workbook": str(args.workbook.resolve()),
        "routes": [], "signal_routes": [], "diagnostic_routes": [], "doip_can_mappings": mappings,
        "doip_selection": "ALL" if args.all else selections, "statistics": statistics, "issues": issues}
    if issues:
        print(json.dumps({"valid": False, "issues": issues}, ensure_ascii=False, indent=2))
        raise SystemExit(1)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"valid": True, "mapping_count": len(mappings), **statistics}, ensure_ascii=False))


if __name__ == "__main__":
    main()
