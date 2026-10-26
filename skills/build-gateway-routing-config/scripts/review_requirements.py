#!/usr/bin/env python3
"""Check requirement coverage and route inputs before a workbook is written."""

import argparse
import copy
import json
import math
import re
from pathlib import Path

from openpyxl import load_workbook

from apply_signal_invalid_values import apply_invalid_values
from extract_candidates import clean, direct_records, find_sheet, normalize_can_id, normalize_version, signal_records, split_scopes, timeout_from_cycle, transition_action

CAN_NETWORKS = {"DG", "EP", "PT", "CH", "IC", "DK", "BD", "DM", "LC", "DA", "SU", "GL"}
STATUSES = {"output", "excluded", "user_deferred", "external_implementation", "metadata_only", "cancelled"}


def requirement_units(manifest):
    selected = {normalize_version(value) for value in manifest.get("selected_versions", [])}
    if not selected:
        return []
    workbook = load_workbook(manifest["source_workbook"], read_only=True, data_only=True)
    try:
        history = find_sheet(workbook, "History")
        current = ""
        found = set()
        units = []
        for number, row in enumerate(history.iter_rows(min_row=3, max_col=3, values_only=True), 3):
            if clean(row[1]):
                current = normalize_version(row[1])
                found.add(current)
            if current not in selected or not clean(row[2]):
                continue
            for index, part in enumerate(split_scopes(clean(row[2]))):
                units.append({"id": f"{number}:{index}", "history_row": number, **part})
        missing = selected - found
        if missing:
            raise ValueError(f"History 版本不存在：{sorted(missing)}")
        return units
    finally:
        workbook.close()


def route_reference(manifest, reference):
    kind, index = reference.split(":")
    if kind not in {"routes", "signal_routes"}:
        raise ValueError(f"需求引用类型无效：{reference}")
    return manifest[kind][int(index)]


def route_key(route, kind):
    if kind == "routes":
        identity = tuple(clean(route.get(field)).casefold() if "can_id" not in field else normalize_can_id(route.get(field)) for field in (
            "source_message_name", "source_can_id", "source_channel", "target_message_name", "target_can_id", "target_channel"
        ))
    else:
        identity = tuple(clean(route.get(field)).casefold() for field in (
            "source_network", "source_message_name", "source_signal_name", "target_network", "target_message_name", "target_signal_name"
        ))
    return identity + (clean(route.get("operation")).upper(),)


def source_route_records(manifest):
    workbook = load_workbook(manifest["source_workbook"], read_only=True, data_only=True)
    try:
        sheet = find_sheet(workbook, "Routing(FLZCU)")
        return {"routes": direct_records(sheet), "signal_routes": signal_records(sheet)}
    finally:
        workbook.close()


def record_route(record, leg, kind, operation):
    if kind == "routes":
        fields = ("source_message_name", "source_can_id", "source_channel", "target_message_name", "target_can_id", "target_channel")
        values = (record["source_message"], record["source_id"], record["source_channel"], record["target_message"], record["target_id"], leg["channel"])
    else:
        fields = ("source_network", "source_message_name", "source_signal_name", "target_network", "target_message_name", "target_signal_name")
        values = (record["source_network"], record["source_message"], record["source_signal"], leg["network"], record["target_message"], record["target_signal"])
    return {**dict(zip(fields, values)), "operation": operation}


def excluded_source_rows(manifest):
    excluded = manifest.get("user_decisions", {}).get("excluded_source_rows", [])
    return {entry["flzcu_row"] for entry in excluded if clean(entry.get("reason")) and clean(entry.get("evidence"))}


def multi_source_targets(manifest, source_records):
    column = "AY" if manifest.get("project_mode") in {"AY", "AX_TO_AY"} else "AX"
    excluded_rows = excluded_source_rows(manifest)
    targets = {}
    for kind, records in source_records.items():
        for record in records:
            if record["row"] in excluded_rows or not record["project_marks"][column]["active"] or record.get("lin_message_route"):
                continue
            for leg in record["target_legs"]:
                if not leg["allowed"]:
                    continue
                route = record_route(record, leg, kind, "ADD")
                key = route_key(route, kind)
                destination = key[3:6]
                if kind == "routes":
                    destination = (key[5], key[4])
                targets.setdefault((kind, destination), set()).add(key[:3])
    return {destination for destination, sources in targets.items() if len(sources) > 1}


def is_multi_source(route, kind, targets):
    key = route_key(route, kind)
    destination = (key[5], key[4]) if kind == "routes" else key[3:6]
    return (kind, destination) in targets


def simple_requirement_keys(manifest, units, source_records, shared_targets):
    result = {}
    errors = []
    events = []
    decisions = manifest.get("user_decisions", {})
    aliases = decisions.get("name_aliases", {})
    subnet_aliases = decisions.get("subnet_aliases", {})
    deferred = set(decisions.get("deferred_history_rows", []))
    excluded_rows = excluded_source_rows(manifest)
    mode = manifest.get("project_mode")
    column = "AY" if mode in {"AY", "AX_TO_AY"} else "AX"
    for unit in units:
        match = re.fullmatch(r"(增加|新增|添加|删除|取消|移除)(报文|信号)(\w+)(?:[（(]([^）)]*)[）)])?从(FL_(?:CANFD|CAN|LIN)_\w+?)(?:网段)?(?:报文|信号)?路由到(.+?)(?:网段)?", unit["text"])
        if not match or unit["scope"] == "FRZCU" or unit["history_row"] in deferred:
            continue
        action, label, name, qualifier, source, target_text = match.groups()
        if not re.fullmatch(r"FL_(?:CANFD|CAN|LIN)_[A-Z0-9]+(?:\s*[/、，,和及]\s*FL_(?:CANFD|CAN|LIN)_[A-Z0-9]+)*", target_text):
            continue
        operation = "ADD" if action in {"增加", "新增", "添加"} else "DELETE"
        kind = "routes" if label == "报文" else "signal_routes"
        name = aliases.get(name, name)
        targets = {subnet_aliases.get(subnet, subnet) for subnet in re.findall(r"FL_(?:CANFD|CAN|LIN)_[A-Z0-9]+", target_text)}
        source = subnet_aliases.get(source, source)
        required_id = normalize_can_id(qualifier) if qualifier else ""
        matches = [record for record in source_records[kind] if clean(record["source_message"] if kind == "routes" else record["source_signal"]).casefold() == name.casefold() and record["source_subnet"] == source and (not required_id or record["source_id_normalized"] == required_id)]
        if not matches and operation == "DELETE":
            continue
        keys = set()
        per_target = {}
        for record in matches:
            if operation == "ADD" and record["row"] in excluded_rows:
                continue
            if record.get("lin_message_route") or (operation == "ADD" and not record["project_marks"][column]["active"]):
                continue
            for leg in record["target_legs"]:
                if not leg["allowed"] or leg["subnet"] not in targets:
                    continue
                route = record_route(record, leg, kind, operation)
                if operation == "ADD" and is_multi_source(route, kind, shared_targets):
                    continue
                key = route_key(route, kind)
                keys.add(key)
                per_target.setdefault(leg["subnet"], set()).add(key)
        if any(len(values) > 1 for values in per_target.values()):
            errors.append(f"需求存在多个完整身份候选，不能取候选并集：{unit['id']}")
        result[unit["id"]] = (kind, set())
        events.extend((unit["id"], kind, key) for key in keys)
    state = {}
    for identity, kind, key in events:
        base = (kind, key[:-1])
        previous = state.get(base)
        if previous and previous[0] == "ADD" and key[-1] == "DELETE":
            state[base] = (None, set())
        elif previous and previous[0] == key[-1]:
            previous[1].add(identity)
        else:
            state[base] = (key[-1], {identity})
    for (kind, identity_key), (operation, providers) in state.items():
        for identity in providers:
            result[identity][1].add(identity_key + (operation,))
    return result, errors


def review_transitions(manifest, source_records, shared_targets):
    mode = manifest.get("project_mode")
    if mode not in {"AX_TO_AY", "AY_TO_AX"}:
        return set(), []
    expected = set()
    excluded_rows = excluded_source_rows(manifest)
    for kind, records in source_records.items():
        for record in records:
            operation = transition_action(record["project_marks"], mode)
            if not operation or record.get("lin_message_route"):
                continue
            for leg in record["target_legs"]:
                if leg["allowed"]:
                    route = record_route(record, leg, kind, operation)
                    if operation == "ADD" and (record["row"] in excluded_rows or is_multi_source(route, kind, shared_targets)):
                        continue
                    expected.add((kind, record["row"], route_key(route, kind)))
    actual = set()
    referenced = set()
    errors = []
    for entry in manifest.get("transition_review", []):
        try:
            reference = entry["route_ref"]
            kind = reference.split(":")[0]
            route = route_reference(manifest, reference)
            evidence = (kind, entry["flzcu_row"], route_key(route, kind))
            if evidence in actual:
                errors.append(f"项目转换审核重复：{reference}")
            actual.add(evidence)
            referenced.add(reference)
        except (ValueError, KeyError, IndexError, TypeError):
            errors.append("项目转换审核引用无效")
    if actual != expected:
        errors.append("项目转换输出与原始AX/AY打点不一致")
    return referenced, errors


def review_route_inputs(manifest):
    errors = []
    for route in manifest.get("routes", []):
        missing = [field for field in ("source_message_name", "source_can_id", "source_length", "source_channel", "target_message_name", "target_can_id", "target_length", "target_channel", "operation") if not clean(route.get(field))]
        if missing:
            errors.append(f"整帧路由缺少必要参数：{missing}")
        if clean(route.get("operation")).upper() not in {"ADD", "DELETE"}:
            errors.append("整帧路由的操作类型必须是ADD或DELETE")
        for field in ("source_channel", "target_channel"):
            if clean(route.get(field)).upper() not in {f"{network}CAN" for network in CAN_NETWORKS}:
                errors.append(f"整帧路由包含不支持的通道：{route.get(field)}")
    signals = manifest.get("signal_routes", [])
    overrides = manifest.get("user_decisions", {}).get("timeout_overrides", [])
    resolved_signals = []
    if signals:
        matrix = manifest.get("signal_matrix")
        if not matrix:
            errors.append("信号路由缺少源信号矩阵路径")
        else:
            try:
                resolved = copy.deepcopy(manifest)
                apply_invalid_values(resolved, Path(matrix))
                resolved_signals = resolved["signal_routes"]
            except (ValueError, KeyError, OSError) as error:
                errors.append(str(error))
    source_records = None
    if signals and manifest.get("source_workbook"):
        source_book = load_workbook(manifest["source_workbook"], read_only=True, data_only=True)
        try:
            source_records = {record["row"]: record for record in signal_records(find_sheet(source_book, "Routing(FLZCU)"))}
        finally:
            source_book.close()
    for index, route in enumerate(signals):
        mapping_only = route.get("routing_mode") == "MAPPING_ONLY"
        if resolved_signals:
            matrix_route = resolved_signals[index]
            if mapping_only != (matrix_route.get("routing_mode") == "MAPPING_ONLY"):
                errors.append(f"纯Event只做Mapping的模式与源矩阵不一致：{route.get('source_signal_name')}")
            if clean(route.get("timeout_value")) != clean(matrix_route.get("timeout_value")):
                errors.append(f"超时值与源信号矩阵不一致：{route.get('source_signal_name')}")
        required = ("source_network", "source_message_name", "source_signal_name", "target_network", "target_message_name", "target_signal_name", "operation")
        if not mapping_only:
            required += ("byte_order", "source_cycle_time_ms", "timeout_time")
        missing = [field for field in required if not clean(route.get(field))]
        if missing:
            errors.append(f"信号路由缺少必要参数：{missing}")
        if clean(route.get("operation")).upper() not in {"ADD", "DELETE"}:
            errors.append("信号路由的操作类型必须是ADD或DELETE")
        key = [clean(route.get(field)) for field in ("source_network", "source_message_name", "source_signal_name")]
        if any(clean(route.get(field)).upper() not in CAN_NETWORKS for field in ("source_network", "target_network")):
            errors.append(f"禁止输出 LIN 或 FRZCU 信号路由：{key}")
        if source_records is not None:
            source = source_records.get(route.get("flzcu_row"))
            if source is None or key != [source["source_network"], source["source_message"], source["source_signal"]]:
                errors.append(f"源信号与路由需求行不一致：{key}")
            else:
                if not mapping_only and route.get("source_cycle_time_ms") != source["source_cycle_time_ms"]:
                    errors.append(f"源周期与原始路由表不一致：{key}")
                if not any(leg["allowed"] and leg["network"] == route.get("target_network") for leg in source["target_legs"]) or route.get("target_message_name") != source["target_message"] or route.get("target_signal_name") != source["target_signal"]:
                    errors.append(f"目标信号与路由需求行不一致：{key}")
        approved = [entry for entry in overrides if entry.get("source_key") == key]
        if mapping_only:
            if clean(route.get("timeout_time")) or clean(route.get("timeout_value")) or approved:
                errors.append(f"纯Event只能做Mapping，不能设置超时或超时特例：{key}")
            continue
        _, expected = timeout_from_cycle(route.get("source_send_type"), route.get("source_cycle_time_ms"))
        if approved:
            if len(approved) != 1 or not clean(approved[0].get("reason")) or not clean(approved[0].get("evidence")):
                errors.append(f"超时特例缺少唯一的用户决定和依据：{key}")
            else:
                expected = approved[0].get("seconds")
        if not isinstance(expected, (int, float)) or not math.isfinite(expected) or expected <= 0:
            errors.append(f"源周期或超时特例无效：{key}")
        elif route.get("timeout_time") != expected:
            errors.append(f"超时时间不符合源周期或用户特例：{key}，应为 {expected} 秒")
    return errors


def review_manifest(manifest):
    errors = review_route_inputs(manifest)
    forwarding = bool(manifest.get("routes") or manifest.get("signal_routes"))
    if forwarding and not manifest.get("selected_versions"):
        errors.append("报文或信号转发缺少所选History版本，不能跳过需求审核")
        return errors
    if forwarding and manifest.get("project_mode") not in {"AX", "AY", "AX_TO_AY", "AY_TO_AX"}:
        errors.append("报文或信号转发缺少有效的项目列模式")
        return errors
    units = requirement_units(manifest)
    if not units and not forwarding:
        return errors
    expected = {unit["id"]: unit for unit in units}
    source_records = source_route_records(manifest)
    shared_targets = multi_source_targets(manifest, source_records)
    excluded_keys = {(kind, route_key(record_route(record, leg, kind, "ADD"), kind)) for kind, records in source_records.items() for record in records if record["row"] in excluded_source_rows(manifest) for leg in record["target_legs"] if leg["allowed"]}
    for kind in ("routes", "signal_routes"):
        for route in manifest.get(kind, []):
            if (kind, route_key(route, kind)) in excluded_keys:
                errors.append(f"用户已排除的源路径仍被输出：{kind} {route_key(route, kind)}")
            if clean(route.get("operation")).upper() == "ADD" and is_multi_source(route, kind, shared_targets):
                errors.append(f"多个源端汇入同一目标端，本条必须告知用户并排除：{kind} {route_key(route, kind)}")
    simple_requirements, requirement_errors = simple_requirement_keys(manifest, units, source_records, shared_targets)
    errors.extend(requirement_errors)
    records = manifest.get("requirement_review", [])
    seen = set()
    referenced, transition_errors = review_transitions(manifest, source_records, shared_targets)
    errors.extend(transition_errors)
    deferred = {int(row) for row in manifest.get("user_decisions", {}).get("deferred_history_rows", [])}
    for record in records:
        identity = record.get("id")
        if identity in seen:
            errors.append(f"需求重复审核：{identity}")
        seen.add(identity)
        unit = expected.get(identity)
        if unit is None:
            errors.append(f"审核记录不属于所选需求：{identity}")
            continue
        status = record.get("status")
        if status not in STATUSES or not clean(record.get("reason")):
            errors.append(f"需求未完成审核或缺少原因：{identity}")
        references = record.get("route_refs", [])
        if status == "output" and not references:
            errors.append(f"需求没有关联输出路由：{identity}")
        if status != "output" and references:
            errors.append(f"非输出需求仍关联输出路由：{identity}")
        if unit["scope"] == "FRZCU" and (status != "excluded" or references):
            errors.append(f"FRZCU 需求必须明确排除：{identity}")
        if status == "user_deferred" and unit["history_row"] not in deferred:
            errors.append(f"暂缓需求缺少用户决定：{identity}")
        if unit["history_row"] in deferred and (status != "user_deferred" or references):
            errors.append(f"已暂缓的需求仍被输出或未保留暂缓结论：{identity}")
        if status == "external_implementation" and not record.get("requirement_details"):
            errors.append(f"其他实现需求缺少具体信号或字段：{identity}")
        for reference in references:
            try:
                route = route_reference(manifest, reference)
                if unit["history_row"] not in route.get("history_rows", []):
                    errors.append(f"输出路由未保留对应 History 行：{identity}/{reference}")
                referenced.add(reference)
            except (ValueError, KeyError, IndexError, TypeError):
                errors.append(f"需求关联的输出不存在：{identity}/{reference}")
        expected_refs = set(record.get("expected_route_refs", []))
        if status == "output" and (not expected_refs or expected_refs != set(references)):
            errors.append(f"审核期望的路由分支与输出不一致：{identity}")
        if identity in simple_requirements and status != "user_deferred":
            kind, required_keys = simple_requirements[identity]
            actual_keys = set()
            for reference in references:
                try:
                    if reference.startswith(kind + ":"):
                        actual_keys.add(route_key(route_reference(manifest, reference), kind))
                except (ValueError, KeyError, IndexError, TypeError):
                    pass
            if actual_keys != required_keys:
                errors.append(f"原始需求的全部有效目标分支未被准确覆盖：{identity}")
    for identity in sorted(expected.keys() - seen):
        errors.append(f"选定需求缺少审核结论：{identity} {expected[identity]['text']}")
    for kind in ("routes", "signal_routes"):
        for index, route in enumerate(manifest.get(kind, [])):
            if deferred.intersection(route.get("history_rows", [])):
                errors.append(f"输出路由关联了用户暂缓需求：{kind}:{index}")
            if f"{kind}:{index}" not in referenced:
                errors.append(f"输出路由没有需求依据：{kind}:{index}")
    return errors


def main():
    parser = argparse.ArgumentParser(description="审核选定 History 的覆盖、排除依据、源周期和矩阵值")
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    errors = review_manifest(manifest)
    if errors:
        raise SystemExit("\n".join(errors))
    print(json.dumps({"review": "passed", "requirement_units": len(manifest.get("requirement_review", []))}, ensure_ascii=False))


if __name__ == "__main__":
    main()
