#!/usr/bin/env python3
import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

OBD_CAN_SHEET = "DIAG Message routing(OBD CAN)"
OBD_ETH_SHEET = "DIAG Message routing(OBD ETH)"
OBD_CAN_PROJECT_COLUMNS = {"P", "Q", "R"}
OBD_ETH_PROJECT_COLUMNS = {"T", "U", "V", "W"}
CLASSIC_CHANNELS = {"BDCAN", "DMCAN", "DGCAN"}


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def is_missing(value):
    return clean(value).upper() in {"", "NA", "N/A", "/", "-"}


def parse_project_mode(value, allowed_columns):
    mode = clean(value).upper()
    if mode == "NONE":
        return None
    match = re.fullmatch(r"([A-Z]+)(?:_TO_([A-Z]+))?", mode)
    if not match:
        raise ValueError(f"无效的项目打点模式：{value!r}")
    old_column, new_column = match.groups()
    if new_column is None:
        if old_column not in allowed_columns:
            raise ValueError(f"项目列 {old_column} 不在允许范围 {sorted(allowed_columns)} 内。")
        return {"old": None, "new": old_column, "mode": mode}
    if old_column not in allowed_columns or new_column not in allowed_columns:
        raise ValueError(
            f"项目转换 {mode} 的列必须位于 {sorted(allowed_columns)} 内。"
        )
    if old_column == new_column:
        raise ValueError("项目转换的旧列和新列不能相同。")
    return {"old": old_column, "new": new_column, "mode": mode}


def parse_explicit_rows(value):
    text = clean(value)
    if not text:
        return None
    rows = set()
    for part in text.split(","):
        match = re.fullmatch(r"(\d+)(?:-(\d+))?", part.strip())
        if not match:
            raise ValueError(f"无效的行号或行号范围：{part!r}")
        first = int(match.group(1))
        last = int(match.group(2) or first)
        if first < 4 or last < first:
            raise ValueError(f"行号范围无效：{part!r}")
        rows.update(range(first, last + 1))
    return sorted(rows)


def mark_active(value, sheet_name, coordinate):
    mark = clean(value)
    if not mark:
        return False
    if mark == "●":
        return True
    raise ValueError(f"{sheet_name}!{coordinate} 出现不支持的打点值：{mark!r}")


def operation_for_row(sheet, row, mode):
    if mode is None:
        return None
    new_index = column_index_from_string(mode["new"])
    new_active = mark_active(
        sheet.cell(row=row, column=new_index).value,
        sheet.title,
        f"{mode['new']}{row}",
    )
    if mode["old"] is None:
        return "ADD" if new_active else None
    old_index = column_index_from_string(mode["old"])
    old_active = mark_active(
        sheet.cell(row=row, column=old_index).value,
        sheet.title,
        f"{mode['old']}{row}",
    )
    if old_active == new_active:
        return None
    return "ADD" if new_active else "DELETE"


def merged_values(sheet):
    values = {}
    for merged_range in sheet.merged_cells.ranges:
        top_left = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                values[(row, column)] = top_left
    return values


def cell_value(sheet, merged, row, column):
    value = sheet.cell(row=row, column=column).value
    if value is None:
        value = merged.get((row, column))
    return clean(value)


def can_id(value):
    raw = clean(value)
    if not raw:
        raise ValueError("CAN ID为空。")
    try:
        number = int(raw, 16)
    except ValueError as error:
        raise ValueError(f"CAN ID格式无效：{raw!r}") from error
    if not 0 <= number <= 0x1FFFFFFF:
        raise ValueError(f"CAN ID超出范围：{raw!r}")
    return f"0x{number:X}"


def reference_channels(template_path):
    workbook = load_workbook(template_path, read_only=True, data_only=True)
    try:
        sheet = workbook["引用数据"]
        return {
            clean(row[0]).upper()
            for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True)
            if clean(row[0])
        }
    finally:
        workbook.close()


def channel_from_network(network, channels):
    text = clean(network).upper()
    if is_missing(text):
        raise ValueError("目标网段为空。")
    code = text.rsplit("_", 1)[-1]
    candidate = f"{code}CAN"
    if candidate not in channels:
        raise ValueError(f"目标网段 {network!r} 无法映射到模板“引用数据”中的CAN通道。")
    return candidate


def channel_settings(channel):
    if channel in CLASSIC_CHANNELS:
        return {
            "rx_type": "STANDARD_NO_FD_CAN",
            "tx_type": "STANDARD_CAN",
            "length": 8,
            "protocol": "DoCAN",
        }
    return {
        "rx_type": "STANDARD_FD_CAN",
        "tx_type": "STANDARD_FD_CAN",
        "length": 64,
        "protocol": "DoCANFD",
    }


def extract_obd_can(sheet, mode, channels):
    if mode is None:
        return [], {"selected": 0, "skipped_non_route": 0}, []
    merged = merged_values(sheet)
    candidates = []
    issues = []
    skipped_non_route = 0
    for row in range(4, sheet.max_row + 1):
        try:
            operation = operation_for_row(sheet, row, mode)
        except ValueError as error:
            issues.append(str(error))
            continue
        if operation is None:
            continue
        request_name = cell_value(sheet, merged, row, 6)
        request_id = cell_value(sheet, merged, row, 7)
        response_name = cell_value(sheet, merged, row, 8)
        response_id = cell_value(sheet, merged, row, 9)
        target_network = cell_value(sheet, merged, row, 12)
        if any(is_missing(value) for value in (request_name, request_id, response_name, response_id, target_network)):
            skipped_non_route += 1
            continue
        identity = f"{sheet.title}第{row}行 {request_name}/{request_id} → {response_name}/{response_id}"
        try:
            target_channel = channel_from_network(target_network, channels)
            settings = channel_settings(target_channel)
            request_can_id = can_id(request_id)
            response_can_id = can_id(response_id)
        except ValueError as error:
            issues.append(f"{identity}：{error}")
            continue
        candidates.append(
            {
                "entry_type": "OBD_CAN",
                "operation": operation,
                "request_message_name": request_name,
                "request_can_id": request_can_id,
                "response_message_name": response_name,
                "response_can_id": response_can_id,
                "target_network": target_network,
                "target_channel": target_channel,
                "target_rx_type": settings["rx_type"],
                "target_tx_type": settings["tx_type"],
                "target_length": settings["length"],
                "diagnostic_protocol": settings["protocol"],
                "routing_type": cell_value(sheet, merged, row, 11),
                "source_sheet": sheet.title,
                "source_row": row,
                "project_mode": mode["mode"],
            }
        )
    return (
        candidates,
        {
            "selected": len(candidates),
            "skipped_non_route": skipped_non_route,
        },
        issues,
    )


def extract_obd_eth(sheet, mode, channels, request_channel=None, explicit_rows=None, explicit_operation=None):
    if mode is None and explicit_rows is None:
        return [], {"selected": 0, "skipped_non_can": 0, "skipped_non_route": 0}, []
    merged = merged_values(sheet)
    candidates = []
    issues = []
    skipped_non_can = 0
    skipped_non_route = 0
    rows = explicit_rows if explicit_rows is not None else range(4, sheet.max_row + 1)
    request_settings = channel_settings(request_channel)
    for row in rows:
        if row > sheet.max_row:
            issues.append(f"{sheet.title} 不存在第{row}行。")
            continue
        try:
            operation = (explicit_operation if explicit_rows is not None
                         else operation_for_row(sheet, row, mode))
        except ValueError as error:
            issues.append(str(error))
            continue
        if operation is None:
            continue
        protocol = cell_value(sheet, merged, row, 4)
        if protocol.casefold() not in {"docan", "docanfd"}:
            if explicit_rows is not None:
                issues.append(f"{sheet.title}第{row}行不是DoCAN/DoCANFD，无法按明确指定的CAN诊断路由新增或删除。")
            else:
                skipped_non_can += 1
            continue
        target_network = cell_value(sheet, merged, row, 11)
        request_name = cell_value(sheet, merged, row, 12)
        request_id = cell_value(sheet, merged, row, 13)
        response_name = cell_value(sheet, merged, row, 15)
        response_id = cell_value(sheet, merged, row, 16)
        if any(is_missing(value) for value in (target_network, request_name, request_id, response_name, response_id)):
            if explicit_rows is not None:
                issues.append(f"{sheet.title}第{row}行缺少K/L/M/O/P中的CAN诊断身份字段。")
            else:
                skipped_non_route += 1
            continue
        identity = f"{sheet.title}第{row}行 {request_name}/{request_id} → {response_name}/{response_id}"
        try:
            target_channel = channel_from_network(target_network, channels)
            settings = channel_settings(target_channel)
            request_can_id = can_id(request_id)
            response_can_id = can_id(response_id)
        except ValueError as error:
            issues.append(f"{identity}：{error}")
            continue
        candidates.append(
            {
                "entry_type": "OBD_ETH",
                "operation": operation,
                "request_message_name": request_name,
                "request_can_id": request_can_id,
                "request_channel": request_channel,
                "request_rx_type": request_settings["rx_type"],
                "request_tx_type": request_settings["tx_type"],
                "request_length": request_settings["length"],
                "response_message_name": response_name,
                "response_can_id": response_can_id,
                "target_network": target_network,
                "target_channel": target_channel,
                "target_rx_type": settings["rx_type"],
                "target_tx_type": settings["tx_type"],
                "target_length": settings["length"],
                "diagnostic_protocol": settings["protocol"],
                "source_sheet": sheet.title,
                "source_row": row,
                "project_mode": mode["mode"] if mode else "EXPLICIT_ROWS",
            }
        )
    return (
        candidates,
        {
            "selected": len(candidates),
            "skipped_non_can": skipped_non_can,
            "skipped_non_route": skipped_non_route,
        },
        issues,
    )


def deduplicate(candidates):
    unique = []
    seen = set()
    for candidate in candidates:
        key = (
            candidate["entry_type"],
            candidate["operation"],
            candidate["request_message_name"].casefold(),
            candidate["request_can_id"].casefold(),
            candidate["response_message_name"].casefold(),
            candidate["response_can_id"].casefold(),
            candidate.get("request_channel", "").casefold(),
            candidate["target_channel"].casefold(),
        )
        if key in seen:
            continue
        seen.add(key)
        unique.append(candidate)
    return unique


def main():
    parser = argparse.ArgumentParser(
        description="读取OBD CAN和OBD ETH项目打点，提取CAN诊断路由候选。"
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--obd-can-project-mode", default="NONE")
    parser.add_argument("--obd-eth-project-mode", default="NONE")
    parser.add_argument("--obd-eth-rows", help="逗号分隔的明确行号或范围；指定时忽略OBD ETH打点")
    parser.add_argument("--obd-eth-operation", choices=("ADD", "DELETE"))
    parser.add_argument("--obd-eth-request-channel", help="OBD ETH CAN诊断路由的请求端CAN通道")
    parser.add_argument(
        "--template",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "assets" / "网关路由配置表空白模板.xlsx",
    )
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    try:
        can_mode = parse_project_mode(args.obd_can_project_mode, OBD_CAN_PROJECT_COLUMNS)
        eth_mode = parse_project_mode(args.obd_eth_project_mode, OBD_ETH_PROJECT_COLUMNS)
        eth_rows = parse_explicit_rows(args.obd_eth_rows)
    except ValueError as error:
        parser.error(str(error))
    if eth_rows is not None and eth_mode is not None:
        parser.error("明确指定OBD ETH行号时，不得同时指定该页项目打点模式。")
    if eth_rows is not None and args.obd_eth_operation is None:
        parser.error("明确指定OBD ETH行号时，必须指定--obd-eth-operation。")
    if eth_rows is None and args.obd_eth_operation is not None:
        parser.error("--obd-eth-operation仅用于明确指定OBD ETH行号。")
    if can_mode is None and eth_mode is None and eth_rows is None:
        parser.error("OBD CAN和OBD ETH至少选择一个项目打点模式。")
    if not args.workbook.is_file():
        parser.error(f"需求工作簿不存在：{args.workbook}")
    if not args.template.is_file():
        parser.error(f"模板不存在：{args.template}")

    channels = reference_channels(args.template.resolve())
    request_channel = clean(args.obd_eth_request_channel).upper()
    if eth_mode is not None or eth_rows is not None:
        if not request_channel or request_channel not in channels:
            parser.error("OBD ETH CAN诊断路由必须提供模板引用数据中存在的--obd-eth-request-channel。")
    workbook = load_workbook(args.workbook.resolve(), read_only=False, data_only=True)
    try:
        try:
            if can_mode is not None:
                can_candidates, can_summary, can_issues = extract_obd_can(
                    workbook[OBD_CAN_SHEET], can_mode, channels
                )
            else:
                can_candidates, can_summary, can_issues = [], {"selected": 0, "skipped_non_route": 0}, []
            if eth_mode is not None or eth_rows is not None:
                eth_candidates, eth_summary, eth_issues = extract_obd_eth(
                    workbook[OBD_ETH_SHEET], eth_mode, channels, request_channel,
                    eth_rows, args.obd_eth_operation,
                )
            else:
                eth_candidates, eth_summary, eth_issues = [], {"selected": 0, "skipped_non_can": 0, "skipped_non_route": 0}, []
        except (KeyError, ValueError) as error:
            parser.error(str(error))
    finally:
        workbook.close()

    issues = can_issues + eth_issues
    if issues:
        parser.error("诊断路由候选存在以下问题：\n- " + "\n- ".join(issues))

    combined = can_candidates + eth_candidates
    candidates = deduplicate(combined)
    result = {
        "source_workbook": str(args.workbook.resolve()),
        "obd_can_project_mode": can_mode["mode"] if can_mode else "NONE",
        "obd_eth_project_mode": eth_mode["mode"] if eth_mode else ("EXPLICIT_ROWS" if eth_rows is not None else "NONE"),
        "obd_eth_explicit_rows": eth_rows,
        "diagnostic_candidates": candidates,
        "summary": {
            "obd_can": can_summary,
            "obd_eth": eth_summary,
            "deduplicated_count": len(combined) - len(candidates),
            "candidate_count": len(candidates),
            "add_count": sum(item["operation"] == "ADD" for item in candidates),
            "delete_count": sum(item["operation"] == "DELETE" for item in candidates),
        },
    }
    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(result["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
