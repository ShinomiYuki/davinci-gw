#!/usr/bin/env python3
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def find_sheet(workbook, expected):
    normalized = expected.strip().casefold()
    for name in workbook.sheetnames:
        if name.strip().casefold() == normalized:
            return workbook[name]
    raise KeyError(f"Missing worksheet: {expected}")


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def main():
    parser = argparse.ArgumentParser(description="List History version groups.")
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    if not workbook_path.is_file():
        parser.error(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    history = find_sheet(workbook, "History")

    markers = []
    for row in range(3, history.max_row + 1):
        version = clean(history.cell(row=row, column=2).value)
        if version:
            markers.append(
                {
                    "version": version,
                    "start_row": row,
                    "date": clean(history.cell(row=row, column=4).value),
                    "note": clean(history.cell(row=row, column=5).value),
                }
            )

    for index, marker in enumerate(markers):
        next_row = markers[index + 1]["start_row"] if index + 1 < len(markers) else history.max_row + 1
        marker["end_row"] = next_row - 1
        marker["change_rows"] = sum(
            1
            for row in range(marker["start_row"], next_row)
            if clean(history.cell(row=row, column=3).value)
        )

    payload = {
        "workbook": str(workbook_path),
        "history_sheet": history.title,
        "versions": markers,
    }
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
