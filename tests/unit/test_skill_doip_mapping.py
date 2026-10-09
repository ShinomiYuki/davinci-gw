"""Skill 映射提取、模板和写表流程的契约测试，不包含客户数据。"""
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

TEST_FILE = Path(__file__).resolve()
ROOT = TEST_FILE.parents[2] if TEST_FILE.parent.name == "unit" else TEST_FILE.parent
SKILL = ROOT / "skills" / "build-gateway-routing-config" if TEST_FILE.parent.name == "unit" else ROOT / "build-gateway-routing-config"
sys.path.insert(0, str(SKILL / "scripts"))
from doip_mapping import HEADERS, normalize_mapping
from extract_doip_can_mappings import extract_mappings
from write_gateway_config import write_workbook, worksheet_path

TEMPLATE = SKILL / "assets" / "网关路由配置表空白模板.xlsx"
CHANNELS = {"PTCAN", "DGCAN", "DMCAN"}


def sheet(workbook, title, testers=1):
    s = workbook.create_sheet(title)
    s.append(["ECU Name", "ECU DoIP Logical Address"] + ["Diagnostic Tool's Logical Address"] * testers +
             ["Destination Network Segment", "Diagnostic Request CANID", "Diagnostic Response CANID", "SD"])
    return s


def book():
    w = Workbook()
    w.remove(w.active)
    return w


def test_actual_tester_sets_and_duplicate_sources():
    w = book()
    a = sheet(w, "OBD")
    a.append(["ECU", "011F", "0E80", "FL_CANFD_PT", "7E0", "7E8", "●"])
    b = sheet(w, "OTA", 3)
    b.append(["ECU", "011F", "0F00", "0F01", "0E81", "FL_CANFD_PT", "7E0", "7E8", "●"])
    b.append(["ICC", "0300", "0F00", "0F01", "0E81", "DoIP", "NA", "NA", "●"])
    a.append(["ICC", "0300", "0E80", "FL_CAN_DG", "725", "7A5", ""])
    c = sheet(w, "重复")
    c.append(["ECU", "011F", "0E80", "FL_CANFD_PT", "7E0", "7E8", "●"])
    records, stats, issues = extract_mappings(w, CHANNELS, all_rows=True)
    assert not issues
    assert len(records) == 5
    assert len([r for r in records if r["logicalAddress"] == "0x0300"]) == 1
    assert stats["duplicates_removed"] == 1
    assert stats["skipped_non_can"] == 1
    assert len(records[0]["sources"]) == 2


@pytest.mark.parametrize("count", [1, 2, 3, 4])
def test_never_fills_missing_testers(count):
    w = book()
    s = sheet(w, "需求", count)
    s.append(["ECU", "011F"] + ["0E80", "0E81", "0F00", "0F01"][:count] + ["FL_CANFD_DM", "7E0", "7E8", "●"])
    records, _, issues = extract_mappings(w, CHANNELS, all_rows=True)
    assert not issues
    assert len(records) == count
    assert all(r["CANBus"] == "DMCAN" for r in records)


def test_merged_functional_and_project_selection():
    w = book()
    s = sheet(w, "功能")
    s.append(["功能寻址", "E400", "0E80", "FL_CANFD_PT", "7DF", "/", "●"])
    s.append([None, None, None, "FL_CAN_DG", None, None, "●"])
    s.append(["ECU", "011F", "0F00", "FL_CANFD_PT", "7E0", "7E8", ""])
    for col in ("A", "B", "C", "E", "F"):
        s.merge_cells(f"{col}2:{col}3")
    records, stats, issues = extract_mappings(w, CHANNELS, selections=[{"sheet": "功能", "project_column": "G"}])
    assert not issues
    assert len(records) == 2
    assert stats["skipped_unmarked"] == 1
    assert {r["CANBus"] for r in records} == {"PTCAN", "DGCAN"}
    assert all(r["Functional"] == "YES" and r["RespCanId"] == "" for r in records)


def test_conflict_and_invalid_mark_block():
    w = book()
    s = sheet(w, "需求")
    s.append(["ECU", "011F", "0E80", "FL_CANFD_PT", "7E0", "7E8", "●"])
    s.append(["ECU", "011F", "0E80", "FL_CANFD_PT", "7E1", "7E9", "●"])
    _, _, issues = extract_mappings(w, CHANNELS, all_rows=True)
    assert len(issues) == 1 and "冲突" in issues[0]
    s["G3"] = "x"
    _, _, issues = extract_mappings(w, CHANNELS, selections=[{"sheet": "需求", "project_column": "G"}])
    assert len(issues) == 1 and "打点无效" in issues[0]


def test_missing_response_does_not_become_functional():
    w = book()
    s = sheet(w, "需求")
    s.append(["ECU", "011F", "0E80", "FL_CANFD_PT", "7E0", "", "●"])
    records, _, issues = extract_mappings(w, CHANNELS, all_rows=True)
    assert not records and "缺少响应" in issues[0]


def test_explicit_scope_is_required():
    w = book()
    sheet(w, "需求")
    with pytest.raises(ValueError):
        extract_mappings(w, CHANNELS)
    _, _, issues = extract_mappings(w, CHANNELS, selections=[{"sheet": "需求"}, "bad"])
    assert len(issues) == 2


def test_numeric_and_merged_numeric_addresses_keep_value():
    w = book()
    s = sheet(w, "数字地址")
    s.append(["ECU", 287, 3712, "FL_CANFD_PT", 2016, 2024, "●"])
    s.append([None, None, None, "FL_CAN_DG", None, None, "●"])
    for col in ("A", "B", "C", "E", "F"):
        s.merge_cells(f"{col}2:{col}3")
    records, _, issues = extract_mappings(w, CHANNELS, all_rows=True)
    assert not issues
    assert len(records) == 2
    assert all((r["logicalAddress"], r["Tester"], r["RequestCanId"], r["RespCanId"]) ==
               ("0x011F", "0x0E80", "0x7E0", "0x7E8") for r in records)


def record():
    return dict(zip(HEADERS, ["0E80", "011F", "OBD", "PTCAN", "7E0", "7E8", ""]))


@pytest.mark.parametrize("updates", [{"Tester": "1234"}, {"RouterType": "OTA"}, {"logicalAddress": "10000"},
                                      {"RequestCanId": "20000000"}, {"CANBus": "UNKNOWN"}, {"Functional": "YES"}])
def test_invalid_mapping_rejected(updates):
    r = record()
    r.update(updates)
    with pytest.raises(ValueError):
        normalize_mapping(r, CHANNELS)


def test_template_and_roundtrip(tmp_path):
    w = load_workbook(TEMPLATE, read_only=True)
    assert "需求描述" in w.sheetnames and "DoIP_to_CAN" in w.sheetnames
    assert not {"诊断报文路由参数", "Sheet1", "命名规则", "直接报文路由需求描述"} & set(w.sheetnames)
    assert tuple(c.value for c in w["DoIP_to_CAN"][1]) == HEADERS
    w.close()
    r = record()
    output = tmp_path / "映射_v1.0.xlsx"
    write_workbook(TEMPLATE, output, [], [], [], [r])
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps({"routes": [], "doip_can_mappings": [r]}), encoding="utf-8")
    command = [sys.executable, str(SKILL / "scripts" / "validate_output.py"), "--manifest", str(manifest), "--workbook", str(output)]
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stdout + result.stderr
    assert json.loads(result.stdout)["doip_can_mapping_count"] == 1
    with zipfile.ZipFile(TEMPLATE) as source, zipfile.ZipFile(output) as target:
        assert source.read("xl/styles.xml") == target.read("xl/styles.xml")
        path = worksheet_path(source, "需求描述")
        assert source.read(path) == target.read(path)
    # 校验必须发现映射页中被人工改坏的响应 ID。
    w = load_workbook(output)
    w["DoIP_to_CAN"]["F2"] = "0x7E9"
    w.save(output)
    w.close()
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 1 and "manifest" in result.stdout


def test_writer_rejects_duplicates_before_output(tmp_path):
    output = tmp_path / "重复_v1.0.xlsx"
    with pytest.raises(ValueError, match="重复"):
        write_workbook(TEMPLATE, output, [], [], [], [record(), record()])
    assert not output.exists()
