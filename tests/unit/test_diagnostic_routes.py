"""诊断输入、引用定位、两层 PDU 与共享删除的针对性验证。"""

from dataclasses import replace
from pathlib import Path

import pytest
from lxml import etree
from openpyxl import load_workbook

from davinci_gw.arxml.document import ArxmlDocument
from davinci_gw.arxml.index import autosar_path
from davinci_gw.domain.errors import OutputValidationError
from davinci_gw.domain.models import (
    DiagnosticEndpoint, DiagnosticRouteChange, MutationAction, OperationType,
    ReferenceDataEntry, SourceLocation, WorkbookData,
)
from davinci_gw.input.workbook_reader import read_workbook
from davinci_gw.input.workbook_schema import DIAGNOSTIC_HEADERS, DIAGNOSTIC_TRANSPORT_FIELDS
from davinci_gw.modules import definitions as d
from davinci_gw.modules.common import semantic_values
from davinci_gw.routing.add import AddCoordinator
from davinci_gw.routing.delete import DeleteCoordinator
from davinci_gw.routing.diagnostic_route import DiagnosticRoutePlanner
from davinci_gw.routing.transaction import TransactionCoordinator
from davinci_gw.validation.output import validate_generated_output


def diagnostic_row(entry="OBD_CAN"):
    row = {"诊断入口类型": entry, "操作类型": "ADD",
           "诊断请求端报文名称": "REQ", "诊断应答端报文名称": "RES"}
    for role, channel in (("请求", "SRC_CAN"), ("应答", "DST_CAN")):
        if entry == "OBD_ETH" and role == "请求":
            continue
        prefix = f"诊断{role}端"
        row.update({prefix + "CANID_REQ": "7E0", prefix + "CANID_RES": "7E8",
                    prefix + "CAN通道": channel, prefix + "Length": 64,
                    prefix + "接收报文类型": "STANDARD_FD_CAN", prefix + "发送报文类型": "STANDARD_FD_CAN"})
        row.update({prefix + field: (0 if field in {"BlockSize", "STmin"} else 100)
                    for field in DIAGNOSTIC_TRANSPORT_FIELDS})
    return row


@pytest.mark.parametrize("entry", ["OBD_CAN", "OBD_ETH"])
def test_read_diagnostic_milliseconds_and_can_only(workbook_factory, entry):
    path = workbook_factory(direct_rows=(), signal_rows=())
    book = load_workbook(path)
    sheet = book.create_sheet("诊断报文路由")
    sheet.append(DIAGNOSTIC_HEADERS)
    row = diagnostic_row(entry)
    sheet.append([row.get(field) for field in DIAGNOSTIC_HEADERS])
    book.save(path)
    result = read_workbook(path)
    assert result.is_valid, result.issues
    route = result.data.diagnostic_routes[0]
    assert dict(route.response_endpoint.transport_parameters)["N_As"] == "0.1"
    assert route.response_endpoint.request_id == 0x7E0
    assert (route.request_endpoint is None) == (entry == "OBD_ETH")


def test_missing_transport_parameter_has_no_default(workbook_factory):
    path = workbook_factory(direct_rows=(), signal_rows=())
    book = load_workbook(path)
    sheet = book.create_sheet("诊断报文路由")
    sheet.append(DIAGNOSTIC_HEADERS)
    row = diagnostic_row()
    row.pop("诊断应答端N_As")
    sheet.append([row.get(field) for field in DIAGNOSTIC_HEADERS])
    book.save(path)
    result = read_workbook(path)
    assert not result.is_valid
    assert any(i.field_name == "诊断应答端N_As" and i.location.row_number == 2 for i in result.issues)


def test_legacy_diagnostic_example_sheet_is_not_executed(workbook_factory):
    path = workbook_factory(direct_rows=(), signal_rows=())
    book = load_workbook(path)
    sheet = book.create_sheet("诊断报文路由")
    sheet.append(DIAGNOSTIC_HEADERS[:-2])
    sheet.append([diagnostic_row().get(field) for field in DIAGNOSTIC_HEADERS[:-2]])
    book.save(path)
    result = read_workbook(path)
    assert result.is_valid, result.issues
    assert not result.data.diagnostic_routes
    assert any(issue.code == "DIAGNOSTIC_LEGACY_SHEET_IGNORED" for issue in result.issues)


@pytest.fixture
def diagnostic_case(arxml_factory):
    """构造名称故意误导但引用完整的两个 CAN 端点。"""
    document = ArxmlDocument.load(arxml_factory())
    ns = document.namespace
    tag = lambda value: f"{{{ns}}}{value}"
    def group(parent, name):
        found = parent.find(tag(name))
        return found if found is not None else etree.SubElement(parent, tag(name))
    def container(parent, name, definition, parameters=None, references=None, module=False):
        node = etree.SubElement(parent, tag("ECUC-MODULE-CONFIGURATION-VALUES" if module else "ECUC-CONTAINER-VALUE"))
        etree.SubElement(node, tag("SHORT-NAME")).text = name
        etree.SubElement(node, tag("DEFINITION-REF"), DEST="ECUC-PARAM-CONF-CONTAINER-DEF").text = definition
        for key, value in (parameters or {}).items():
            numeric = str(value).replace(".", "").isdigit() or value in {"true", "false"}
            entry = etree.SubElement(group(node, "PARAMETER-VALUES"), tag("ECUC-NUMERICAL-PARAM-VALUE" if numeric else "ECUC-TEXTUAL-PARAM-VALUE"))
            etree.SubElement(entry, tag("DEFINITION-REF"), DEST="ECUC-FLOAT-PARAM-DEF" if numeric else "ECUC-ENUMERATION-PARAM-DEF").text = key
            etree.SubElement(entry, tag("VALUE")).text = str(value)
        for key, value in (references or {}).items():
            entry = etree.SubElement(group(node, "REFERENCE-VALUES"), tag("ECUC-REFERENCE-VALUE"))
            etree.SubElement(entry, tag("DEFINITION-REF"), DEST="ECUC-REFERENCE-DEF").text = key
            etree.SubElement(entry, tag("VALUE-REF"), DEST="ECUC-CONTAINER-VALUE").text = value
        group(node, "SUB-CONTAINERS")
        return node
    loc = SourceLocation("诊断报文路由", 2)
    parameters = tuple((key, "0" if key in {"STmin", "BlockSize"} else "0.1") for key in DIAGNOSTIC_TRANSPORT_FIELDS)
    route = DiagnosticRouteChange(OperationType.ADD, "OBD_CAN", "REQ", "RES",
        DiagnosticEndpoint("SRC_CAN", 0x7e0, 0x7e8, parameters), DiagnosticEndpoint("DST_CAN", 0x7e0, 0x7e8, parameters), loc)
    refs = (ReferenceDataEntry("SRC_CAN", "TX", "HRH", loc), ReferenceDataEntry("DST_CAN", "TX2", "HRH2", loc))
    workbook = WorkbookData(Path("diag_v1.1.xlsx"), "1.1", refs, diagnostic_routes=(route,))
    coordinator = AddCoordinator(document, workbook)
    index = coordinator.index
    node = lambda path: index.find_by_path(path)[0]
    path = lambda value: autosar_path(value, ns)
    canif_parent = node(coordinator.canif.rx_parent_path)
    module = index.find_by_short_name("CanIf")[0]
    tp = container(module.getparent(), "Transport", "/MICROSAR/CanTp", module=True)
    config = container(group(tp, "CONTAINERS"), "Config", "/MICROSAR/CanTp/CanTpConfig")
    bsw_parent = index.find_by_definition_ref(d.PDUR_BSW_MODULE)[0].getparent()
    bsw = container(bsw_parent, "TransportModule", d.PDUR_BSW_MODULE,
                    references={d.PDUR_BSW_MODULE_REF: path(tp)})
    upper = {}
    for channel, hardware, request_side in (("SRC_CAN", ("HRH", "TX"), True), ("DST_CAN", ("HRH2", "TX2"), False)):
        lower = {}
        upper[channel] = {}
        channel_node = container(group(config, "SUB-CONTAINERS"), "MisleadingName" + str(request_side), d.CANTP_CHANNEL,
                                 {d.CANTP_CHANNEL + "/CanTpChannelMode": "CANTP_MODE_FULL_DUPLEX"})
        for rx in (True, False):
            suffix = channel + ("Rx" if rx else "Tx")
            can_id = 0x7e0 if rx == request_side else 0x7e8
            for target, name in ((lower, "Low"), (upper[channel], "High")):
                target[rx] = path(container(group(node(coordinator.ecuc.parent_path), "SUB-CONTAINERS"), name + suffix,
                    d.ECUC_PDU, {d.ECUC_PDU_LENGTH: "64", d.ECUC_PDU_J1939: "false"}))
            definition = d.CANIF_RX if rx else d.CANIF_TX
            pp = {definition + "/CanIf" + ("RxPduCanId" if rx else "TxPduCanId"): str(can_id),
                  (d.CANIF_RX_CAN_ID_TYPE if rx else d.CANIF_TX_CAN_ID_TYPE): "STANDARD_FD_CAN",
                  (d.CANIF_RX_DLC if rx else d.CANIF_TX_DLC): "64",
                  (d.CANIF_RX_INDICATION_UL if rx else d.CANIF_TX_CONFIRM_UL): "CAN_TP",
                  (d.CANIF_RX_INDICATION_NAME if rx else d.CANIF_TX_CONFIRM_NAME): "CanTp_RxIndication" if rx else "CanTp_TxConfirmation",
                  (d.CANIF_RX_DLC_CHECK if rx else d.CANIF_TX_TRUNCATION): "false" if rx else "true"}
            hw = index.find_by_short_name(hardware[0 if rx else 1])[0]
            container(group(canif_parent, "SUB-CONTAINERS"), "Frame" + suffix, definition, pp,
                {(d.CANIF_RX_PDU_REF if rx else d.CANIF_TX_PDU_REF): lower[rx],
                 (d.CANIF_RX_HRH_REF if rx else d.CANIF_TX_BUFFER_REF): path(hw)})
        for rx in (True, False):
            definition = d.CANTP_RX if rx else d.CANTP_TX
            params = DiagnosticRoutePlanner.transport_values(route.request_endpoint, rx)
            prefix = "CanTpRx" if rx else "CanTpTx"
            params.update({definition + "/" + prefix + "Dl": "64", definition + "/" + prefix + "TaType": "CANTP_CANFD_PHYSICAL",
                           definition + "/" + prefix + "NSduId": "1"})
            sdu = container(group(channel_node, "SUB-CONTAINERS"), prefix, definition, params,
                            {(d.CANTP_RX_SDU_REF if rx else d.CANTP_TX_SDU_REF): upper[channel][rx]})
            for child, ref, receive in ((d.CANTP_RX_NPDU, "CanTpRxNPduRef", True), (d.CANTP_TX_FC, "CanTpTxFcNPduRef", False)) if rx else (
                    (d.CANTP_RX_FC, "CanTpRxFcNPduRef", True), (d.CANTP_TX_NPDU, "CanTpTxNPduRef", False)):
                handle = {d.CANTP_RX_NPDU: "CanTpRxNPduId", d.CANTP_RX_FC: "CanTpRxFcNPduId",
                          d.CANTP_TX_FC: "CanTpTxFcNPduConfirmationPduId", d.CANTP_TX_NPDU: "CanTpTxNPduConfirmationPduId"}[child]
                container(group(sdu, "SUB-CONTAINERS"), child.rsplit("/", 1)[-1], child,
                          {child + "/" + handle: "1"}, references={child + "/" + ref: lower[receive]})
    queue = container(group(node(coordinator.pdur.path_parent).getparent().getparent(), "SUB-CONTAINERS"), "Queue",
                      "/MICROSAR/PduR/PduRRoutingTables/PduRQueue")
    container(group(queue, "SUB-CONTAINERS"), "PduRSharedBufferQueue", d.PDUR_SHARED_QUEUE,
              {d.PDUR_SHARED_QUEUE + "/PduRQueueDepth": "3", d.PDUR_SHARED_QUEUE + "/PduRTpThreshold": "0"},
              {d.PDUR_SHARED_QUEUE + "/PduRImplicitTxBufferRef": upper["SRC_CAN"][True]})
    for label, source, target in (("Req", upper["SRC_CAN"][True], upper["DST_CAN"][False]),
                                  ("Res", upper["DST_CAN"][True], upper["SRC_CAN"][False])):
        route_path = container(group(node(coordinator.pdur.path_parent), "SUB-CONTAINERS"), label, d.PDUR_PATH,
                               {d.PDUR_PATH_COMM_TYPE: "TRANSPORT_PROTOCOL"})
        container(group(route_path, "SUB-CONTAINERS"), "Src", d.PDUR_SRC,
                  {d.PDUR_SRC_HANDLE: "20", d.PDUR_SRC_DIRECTION: "RECEIVE"},
                  references={d.PDUR_SRC_PDU_REF: source, d.PDUR_SRC_MODULE_REF: path(bsw)})
        dest = container(group(route_path, "SUB-CONTAINERS"), "Dest", d.PDUR_DEST,
                        {d.PDUR_DEST_HANDLE: "20", d.PDUR_DEST_DIRECTION: "TRANSMIT",
                         d.PDUR_DEST_ROUTING_TYPE: "GATEWAY_ROUTING", d.PDUR_DEST_PROCESSING: "IMMEDIATE",
                         d.PDUR_DEST_LENGTH_STRATEGY: "UNUSED", d.PDUR_DEST_CROSS_PARTITION: "false"},
                        references={d.PDUR_DEST_PDU_REF: target, d.PDUR_DEST_MODULE_REF: path(bsw), d.PDUR_QUEUE_REF: path(queue)})
        if label == "Res":
            container(bsw_parent, "ArbitraryDiagnosticGroup", d.PDUR_ROUTING_GROUP,
                      references={d.PDUR_ROUTING_GROUP_DEST_REF: path(dest)})
    return document, workbook, container, group


def test_complete_reference_chain_is_idempotent_despite_misleading_names(diagnostic_case):
    document, workbook, _, _ = diagnostic_case
    plan = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not plan.errors, plan.issues
    assert plan.diagnostic_existing_count == 1
    assert not plan.operations


def test_delete_can_pair_preserves_other_destination_and_shared_endpoint(diagnostic_case):
    document, workbook, container, group = diagnostic_case
    index = document.build_index()
    response = index.find_by_short_name("Res")[0]
    destination = index.find_by_path(autosar_path(response, document.namespace) + "/Dest")[0]
    # 外部模块目标作为不透明共享对象保留，不需要读取其模块或协议配置。
    container(group(response, "SUB-CONTAINERS"), "OtherDestination", d.PDUR_DEST,
              references={d.PDUR_DEST_PDU_REF: "/Unrelated/OpaquePdu"})
    route = replace(workbook.diagnostic_routes[0], operation=OperationType.DELETE,
                    request_endpoint=replace(workbook.diagnostic_routes[0].request_endpoint, transport_parameters=()),
                    response_endpoint=replace(workbook.diagnostic_routes[0].response_endpoint, transport_parameters=()))
    plan = TransactionCoordinator(document, replace(workbook, diagnostic_routes=(route,))).plan_and_apply()
    assert not plan.errors, plan.issues
    assert plan.diagnostic_deleted_count == 1
    index = document.build_index()
    assert index.find_by_short_name("OtherDestination")
    assert index.find_by_short_name("Res")
    assert not index.find_by_path(autosar_path(response, document.namespace) + "/Dest")
    assert plan.decisions


@pytest.mark.parametrize("channel,length,rx_type,tx_type,ta", [
    ("BDCAN", 8, "STANDARD_NO_FD_CAN", "STANDARD_CAN", "CANTP_PHYSICAL"),
    ("DST_CAN", 64, "STANDARD_FD_CAN", "STANDARD_FD_CAN", "CANTP_CANFD_PHYSICAL"),
])
def test_new_eth_can_side_has_two_pdu_layers_and_no_pdur_changes(diagnostic_case, channel, length, rx_type, tx_type, ta):
    document, workbook, _, _ = diagnostic_case
    original = workbook.diagnostic_routes[0]
    endpoint = replace(original.response_endpoint, channel=channel, request_id=0x700, response_id=0x708)
    route = replace(original, entry_type="OBD_ETH", request_endpoint=None, response_endpoint=endpoint)
    refs = tuple(replace(r, channel_name=channel) if r.channel_name == "DST_CAN" else r for r in workbook.reference_data)
    workbook = replace(workbook, reference_data=refs, diagnostic_routes=(route,))
    plan = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not plan.errors, plan.issues
    assert plan.diagnostic_added_count == 1
    assert not any("PDUR" in op.kind.value for op in plan.operations)
    validate_generated_output(document, plan)
    index = document.build_index()
    rx = index.find_by_short_name(f"GWT_Diag_CanIf_RES_708_{channel}_Rx")[0]
    tx = index.find_by_short_name(f"GWT_Diag_CanIf_REQ_700_{channel}_Tx")[0]
    assert semantic_values(rx, document.namespace)[0][d.CANIF_RX_CAN_ID_TYPE] == (rx_type,)
    assert semantic_values(tx, document.namespace)[0][d.CANIF_TX_CAN_ID_TYPE] == (tx_type,)
    tp = index.find_by_short_name(f"GWT_CanTpChannelGW_{channel}700_708")[0]
    params, refs = semantic_values(tp, document.namespace, recursive=True)
    assert params[d.CANTP_RX + "/CanTpRxDl"] == (str(length),)
    assert params[d.CANTP_TX + "/CanTpTxTaType"] == (ta,)
    assert refs[d.CANTP_RX_SDU_REF] != refs[d.CANTP_RX_NPDU + "/CanTpRxNPduRef"]
    repeated = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not repeated.errors, repeated.issues
    assert not repeated.operations


def test_new_eth_selects_upper_pdu_length_from_target_channel(diagnostic_case):
    document, workbook, _, _ = diagnostic_case
    index = document.build_index()
    ns = document.namespace
    for name in ("HighSRC_CANRx", "HighSRC_CANTx"):
        pdu = index.find_by_short_name(name)[0]
        for entry in pdu.find(f"{{{ns}}}PARAMETER-VALUES"):
            if entry.findtext(f"{{{ns}}}DEFINITION-REF") == d.ECUC_PDU_LENGTH:
                entry.find(f"{{{ns}}}VALUE").text = "8"
    # 两个网段共用 HRH，但 TxBuffer 各自独立；Rx 样板必须结合 Channel 的 Tx 腿归属。
    source_rx = index.find_by_short_name("FrameSRC_CANRx")[0]
    for entry in source_rx.find(f"{{{ns}}}REFERENCE-VALUES"):
        if entry.findtext(f"{{{ns}}}DEFINITION-REF") == d.CANIF_RX_HRH_REF:
            entry.find(f"{{{ns}}}VALUE-REF").text = autosar_path(
                index.find_by_short_name("HRH2")[0], ns)
    original = workbook.diagnostic_routes[0]
    route = replace(original, entry_type="OBD_ETH", request_endpoint=None,
                    response_endpoint=replace(original.response_endpoint, request_id=0x700, response_id=0x708))
    plan = TransactionCoordinator(document, replace(workbook, diagnostic_routes=(route,))).plan_and_apply()
    assert not plan.errors, plan.issues
    assert plan.diagnostic_added_count == 1
    index = document.build_index()
    for name in ("GWT_Diag_Tp_EcuC_RES_708_DST_CAN_Rx", "GWT_Diag_Tp_EcuC_REQ_700_DST_CAN_Tx"):
        pdu = index.find_by_short_name(name)[0]
        assert semantic_values(pdu, ns)[0][d.ECUC_PDU_LENGTH] == ("64",)


def test_multiple_new_eth_channels_have_unique_cantp_symbolic_names(diagnostic_case):
    document, workbook, _, _ = diagnostic_case
    original = workbook.diagnostic_routes[0]
    first = replace(original, entry_type="OBD_ETH", request_endpoint=None,
                    request_name="Diag_ALPHA_REQ", response_name="Diag_ALPHA_RES",
                    response_endpoint=replace(original.response_endpoint, request_id=0x700, response_id=0x708))
    second = replace(first, request_name="Diag_BETA_REQ", response_name="Diag_BETA_RES",
                     response_endpoint=replace(first.response_endpoint,
                     request_id=0x701, response_id=0x709), source=SourceLocation("诊断报文路由", 3))
    plan = TransactionCoordinator(document, replace(workbook, diagnostic_routes=(first, second))).plan_and_apply()
    assert not plan.errors, plan.issues
    assert plan.diagnostic_added_count == 2
    index = document.build_index()
    for ecu, request, response in (("ALPHA", "700", "708"), ("BETA", "701", "709")):
        channel = index.find_by_short_name(f"GWT_CanTpChannelGW_DST_CAN{request}_{response}")[0]
        names = {node.findtext(f"{{{document.namespace}}}SHORT-NAME")
                 for node in channel.iter() if node.findtext(f"{{{document.namespace}}}DEFINITION-REF") in {
                     d.CANTP_RX, d.CANTP_TX, d.CANTP_RX_NPDU,
                     d.CANTP_TX_FC, d.CANTP_RX_FC, d.CANTP_TX_NPDU}}
        assert names == {
            f"CanTpRxNSdu_{ecu}_DST_CAN", f"CanTpTxNSdu_{ecu}_DST_CAN",
            f"CanTpRxNPdu_{response}", f"CanTpTxFcNPdu_{request}",
            f"CanTpRxFcNPdu_{response}", f"CanTpTxNPdu_{request}",
        }
    validate_generated_output(document, plan)
    duplicated = index.find_by_short_name("CanTpRxNSdu_BETA_DST_CAN")[0]
    duplicated.find(f"{{{document.namespace}}}SHORT-NAME").text = "CanTpRxNSdu_ALPHA_DST_CAN"
    with pytest.raises(OutputValidationError, match="CanTp 符号名"):
        validate_generated_output(document, plan)


def test_new_can_destination_gets_dedicated_queue(diagnostic_case):
    document, workbook, _, _ = diagnostic_case
    index = document.build_index()
    request = index.find_by_short_name("Req")[0]
    request.getparent().remove(request)
    plan = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not plan.errors, plan.issues
    validate_generated_output(document, plan)
    index = document.build_index()
    queue = index.find_by_short_name("GWT_Diag_PduRQueue_7E0_SRC_CAN_To_DST_CAN")[0]
    params, _ = semantic_values(queue, document.namespace, recursive=True)
    assert params[d.PDUR_SHARED_QUEUE + "/PduRQueueDepth"] == ("3",)
    assert params[d.PDUR_SHARED_QUEUE + "/PduRTpThreshold"] == ("0",)
    queue_path = autosar_path(queue, document.namespace)
    assert len(index.find_referrers(queue_path)) == 1
    repeated = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not repeated.errors, repeated.issues
    assert not repeated.operations


def test_dedicated_queue_delete_and_readd_cycle(diagnostic_case):
    document, workbook, container, group = diagnostic_case
    index = document.build_index()
    request = index.find_by_short_name("Req")[0]
    request.getparent().remove(request)
    # 独立的存量路由保留响应通道组证据，并与被删除的业务腿共享发送端点。
    response = index.find_by_short_name("Res")[0]
    destination = index.find_by_path(autosar_path(response, document.namespace) + "/Dest")[0]
    parameters, refs = semantic_values(destination, document.namespace)
    another = container(response.getparent(), "IndependentRoute", d.PDUR_PATH)
    member = container(group(another, "SUB-CONTAINERS"), "IndependentDestination", d.PDUR_DEST,
                       parameters={key: values[0] for key, values in parameters.items()},
                       references={key: values[0] for key, values in refs.items()})
    group_node = index.find_by_short_name("ArbitraryDiagnosticGroup")[0]
    entry = etree.fromstring(etree.tostring(group_node.find(f"{{{document.namespace}}}REFERENCE-VALUES")[0]))
    entry.find(f"{{{document.namespace}}}VALUE-REF").text = autosar_path(member, document.namespace)
    group_node.find(f"{{{document.namespace}}}REFERENCE-VALUES").append(entry)
    added = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not added.errors, added.issues
    route = workbook.diagnostic_routes[0]
    deletion = replace(route, operation=OperationType.DELETE,
                       request_endpoint=replace(route.request_endpoint, transport_parameters=()),
                       response_endpoint=replace(route.response_endpoint, transport_parameters=()))
    delete_book = replace(workbook, diagnostic_routes=(deletion,))
    deleted = TransactionCoordinator(document, delete_book).plan_and_apply()
    assert not deleted.errors, deleted.issues
    validate_generated_output(document, deleted)
    index = document.build_index()
    assert not index.find_by_short_name("GWT_Diag_PduRQueue_7E0_SRC_CAN_To_DST_CAN")
    assert index.find_by_short_name("Queue")  # 独立目标腿仍引用的共享队列保留。
    repeated = TransactionCoordinator(document, delete_book).plan_and_apply()
    assert not repeated.operations
    restored = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not restored.errors, [issue.message for issue in restored.issues]
    validate_generated_output(document, restored)


def test_new_functional_route_has_no_fc_or_response_and_depth_five(diagnostic_case):
    document, workbook, _, _ = diagnostic_case
    index = document.build_index()
    ns = document.namespace
    # 从夹具建立明确的功能寻址模板，保留请求 Rx 和目标 Tx 两个单向 N-SDU。
    for rx, channel_name in ((True, "MisleadingNameTrue"), (False, "MisleadingNameFalse")):
        channel = index.find_by_short_name(channel_name)[0]
        for sdu in tuple(channel.find(f"{{{ns}}}SUB-CONTAINERS")):
            keep = d.CANTP_RX if rx else d.CANTP_TX
            if sdu.findtext(f"{{{ns}}}DEFINITION-REF") != keep:
                sdu.getparent().remove(sdu)
                continue
            prefix = "CanTpRx" if rx else "CanTpTx"
            for entry in sdu.find(f"{{{ns}}}PARAMETER-VALUES"):
                if entry.findtext(f"{{{ns}}}DEFINITION-REF") == f"{keep}/{prefix}TaType":
                    entry.find(f"{{{ns}}}VALUE").text = "CANTP_CANFD_FUNCTIONAL"
            for child in tuple(sdu.find(f"{{{ns}}}SUB-CONTAINERS")):
                if child.findtext(f"{{{ns}}}DEFINITION-REF") in {d.CANTP_RX_FC, d.CANTP_TX_FC}:
                    child.getparent().remove(child)
    original = workbook.diagnostic_routes[0]
    route = replace(original, response_name="/",
                    request_endpoint=replace(original.request_endpoint, request_id=0x7DD, response_id=None),
                    response_endpoint=replace(original.response_endpoint, request_id=0x7DD, response_id=None))
    workbook = replace(workbook, diagnostic_routes=(route,))
    plan = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not plan.errors, [issue.message for issue in plan.issues]
    validate_generated_output(document, plan)
    assert sum(op.definition_ref == d.PDUR_DEST for op in plan.operations) == 1
    assert not any(op.definition_ref in {d.CANTP_RX_FC, d.CANTP_TX_FC} for op in plan.operations)
    queue = next(op for op in plan.operations if op.definition_ref == d.PDUR_SHARED_QUEUE)
    assert dict(queue.parameters)[d.PDUR_SHARED_QUEUE + "/PduRQueueDepth"] == "5"
    repeated = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not repeated.errors, [issue.message for issue in repeated.issues]
    assert not repeated.operations


def test_fanout_reuses_source_path_planned_in_same_batch(diagnostic_case):
    document, workbook, container, _ = diagnostic_case
    index = document.build_index()
    request = index.find_by_short_name("Req")[0]
    request.getparent().remove(request)
    for name, original, definition in (("HRH3", "HRH2", d.CANIF_HRH), ("TX3", "TX2", d.CANIF_BUFFER)):
        container(index.find_by_short_name(original)[0].getparent(), name, definition)
    route = workbook.diagnostic_routes[0]
    first = replace(route, response_endpoint=replace(route.response_endpoint, request_id=0x700, response_id=0x708))
    second = replace(first, response_endpoint=replace(first.response_endpoint, channel="THIRD_CAN"),
                     source=SourceLocation("诊断报文路由", 3))
    workbook = replace(workbook, diagnostic_routes=(first, second), reference_data=workbook.reference_data + (
        ReferenceDataEntry("THIRD_CAN", "TX3", "HRH3", second.source),))
    plan = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not plan.errors, [issue.message for issue in plan.issues]
    validate_generated_output(document, plan)
    paths = [op for op in plan.operations if op.definition_ref == d.PDUR_PATH and "7E0_SRC_CAN" in op.short_name]
    assert len(paths) == 1
    assert sum(op.definition_ref == d.PDUR_DEST and op.parent_path == paths[0].object_path for op in plan.operations) == 2
    repeated = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not repeated.errors, [issue.message for issue in repeated.issues]
    assert not repeated.operations
