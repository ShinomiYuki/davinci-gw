"""DoIP 映射、请求响应复用和前置选择的脱敏验证。"""

from dataclasses import replace

import pytest
from openpyxl import load_workbook

from davinci_gw.arxml.index import autosar_path
from davinci_gw.domain.models import DoIPMapping, MutationKind, SourceLocation
from davinci_gw.input.doip_reader import DOIP_HEADERS
from davinci_gw.input.workbook_reader import read_workbook
from davinci_gw.modules import definitions as d
from davinci_gw.modules.common import definition_ref, semantic_values
from davinci_gw.routing.doip_route import ACTIVATION, CHANNEL, INTERFACE, RX, TARGET, TARGET_REF, TESTER, TX
from davinci_gw.routing.transaction import TransactionCoordinator
from davinci_gw.validation.output import validate_generated_output
from tests.unit.test_diagnostic_routes import diagnostic_case


def mapping(address=3, bus="DST_CAN", request=0x7E0, response=0x7E8, functional=False):
    return DoIPMapping(0x0E80, address, "OBD", bus, request, response, functional, SourceLocation("DoIP_to_CAN", 2))


@pytest.fixture
def doip_case(diagnostic_case):
    document, workbook, container, group = diagnostic_case
    ns = document.namespace
    path = lambda n: autosar_path(n, ns)
    index = document.build_index()
    get = lambda name: index.find_by_short_name(name)[0]
    module = container(get("Transport").getparent(), "DoIP", "/MICROSAR/DoIP", module=True)
    bsw = container(get("TransportModule").getparent(), "DoipModule", d.PDUR_BSW_MODULE,
                    references={d.PDUR_BSW_MODULE_REF: path(module)})
    config = container(group(module, "CONTAINERS"), "DoIPConfigSet", "/MICROSAR/DoIP/DoIPConfigSet")
    interface = container(group(config, "SUB-CONTAINERS"), "DoIPInterface", INTERFACE)
    connections = container(group(interface, "SUB-CONTAINERS"), "DoIPConnections", INTERFACE + "/DoIPConnections")
    target = container(group(connections, "SUB-CONTAINERS"), "ExistingTarget", TARGET,
                       {TARGET + "/DoIPTargetAddressValue": "3", TARGET + "/DoIPTargetAddressVerification": "false"})
    activation = container(group(interface, "SUB-CONTAINERS"), "ExistingActivation", ACTIVATION,
                           {ACTIVATION + "/DoIPRoutingActivationNumber": "0"}, {TARGET_REF: path(target)})
    tester = container(group(interface, "SUB-CONTAINERS"), "ExistingTester", TESTER,
                       {TESTER + "/DoIPTesterSA": str(0x0E80)}, {TESTER + "/DoIPRoutingActivationRef": path(activation)})
    channel = container(group(interface, "SUB-CONTAINERS"), "ExistingChannel", CHANNEL,
                        {CHANNEL + "/DoIPChannelId": "5", CHANNEL + "/DoIPTargetMaxBufSize": "4100",
                         CHANNEL + "/DoIPChannelReadyNotification": "false"},
                        {CHANNEL + "/DoIPChannelSARef": path(tester), CHANNEL + "/DoIPChannelTARef": path(target)})
    pdus = {}
    for definition, direction in ((RX, "Rx"), (TX, "Tx")):
        pdu = container(get("HighDST_CANTx").getparent(), "Doip" + direction, d.ECUC_PDU,
                        {d.ECUC_PDU_LENGTH: "1500", d.ECUC_PDU_J1939: "false"})
        parameters = {definition + f"/DoIPPduR{direction}PduId": "5"}
        if definition == TX:
            parameters[definition + "/DoIPPduType"] = "DOIP_TPPDU"
        container(group(channel, "SUB-CONTAINERS"), "Pdu" + direction, definition, parameters,
                  {definition + f"/DoIPPduR{direction}PduRef": path(pdu)})
        pdus[direction] = path(pdu)
    request_path = container(get("Req").getparent(), "ExistingDoipRequest", d.PDUR_PATH,
                             {d.PDUR_PATH_COMM_TYPE: "TRANSPORT_PROTOCOL"})
    container(group(request_path, "SUB-CONTAINERS"), "DoipSource", d.PDUR_SRC,
              {d.PDUR_SRC_HANDLE: "35", d.PDUR_SRC_DIRECTION: "RECEIVE"},
              {d.PDUR_SRC_PDU_REF: pdus["Rx"], d.PDUR_SRC_MODULE_REF: path(bsw)})
    for owner, target_pdu, module_ref, name in ((request_path, path(get("HighDST_CANTx")), path(get("TransportModule")), "CanTarget"),
                                               (get("Res"), pdus["Tx"], path(bsw), "DoipTarget")):
        container(group(owner, "SUB-CONTAINERS"), name, d.PDUR_DEST,
                  {d.PDUR_DEST_HANDLE: "35", d.PDUR_DEST_DIRECTION: "TRANSMIT",
                   d.PDUR_DEST_ROUTING_TYPE: "GATEWAY_ROUTING", d.PDUR_DEST_PROCESSING: "IMMEDIATE",
                   d.PDUR_DEST_LENGTH_STRATEGY: "UNUSED", d.PDUR_DEST_CROSS_PARTITION: "false"},
                  {d.PDUR_DEST_PDU_REF: target_pdu, d.PDUR_DEST_MODULE_REF: module_ref, d.PDUR_QUEUE_REF: path(get("Queue"))})
    return document, replace(workbook, diagnostic_routes=(), doip_mappings=(mapping(),)), container, group


@pytest.mark.parametrize("address", [3, 0x101])
def test_physical_reuse_and_new_channel_are_idempotent(doip_case, address):
    document, workbook, _, _ = doip_case
    workbook = replace(workbook, doip_mappings=(mapping(address),))
    plan = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not plan.errors, plan.issues
    validate_generated_output(document, plan)
    if address == 3:
        assert plan.doip_existing_count == 1
        assert not plan.operations
    else:
        assert plan.doip_added_count == 1
        assert sum(op.definition_ref == CHANNEL for op in plan.operations) == 1
        assert sum(op.definition_ref == TX for op in plan.operations) == 1
        assert sum(op.kind is MutationKind.DOIP_REFERENCE for op in plan.operations) == 1
        assert not any(op.kind is MutationKind.PDUR_QUEUE for op in plan.operations)
        assert all(op.short_name.startswith("GWT_") for op in plan.operations if op.definition_ref in {CHANNEL, TARGET, d.ECUC_PDU, d.PDUR_PATH})
        again = TransactionCoordinator(document, workbook).plan_and_apply()
        assert not again.errors, again.issues
        assert again.doip_existing_count == 1
        assert not again.operations


@pytest.mark.parametrize("policy", [None, "BLOCK", "SKIP"])
def test_missing_prerequisites_require_choice_and_skip_whole_mapping(doip_case, policy):
    document, workbook, _, _ = doip_case
    workbook = replace(workbook, doip_mappings=(mapping(), mapping(0x102, request=0x700, response=0x708)))
    plan = TransactionCoordinator(document, workbook, doip_missing_policy=policy).plan_and_apply()
    assert not plan.operations
    if policy == "SKIP":
        assert not plan.errors
        assert plan.doip_skipped_count == 1
        assert plan.doip_existing_count == 1
    else:
        assert plan.errors
        assert any(i.code == "DOIP_PREREQUISITE_CHOICE_REQUIRED" for i in plan.issues) == (policy is None)
    issue = next(i for i in plan.issues if i.code.startswith("DOIP_PREREQUISITE_") and i.actual_value)
    assert issue.actual_value["logicalAddress"] == "0x0102"
    assert issue.actual_value["CANBus"] == "DST_CAN"


def test_can_diagnostic_is_applied_before_doip_in_same_transaction(doip_case):
    document, workbook, _, _ = doip_case
    # 使用既有诊断参数，明确新增另一对 CAN ID，再建立 DoIP 映射。
    from davinci_gw.domain.models import DiagnosticEndpoint, DiagnosticRouteChange, OperationType
    params = tuple((k, "0" if k in {"BlockSize", "STmin"} else "0.1")
                   for k in ("N_As", "N_Bs", "N_Cs", "N_Ar", "N_Br", "N_Cr", "BlockSize", "STmin"))
    route = DiagnosticRouteChange(OperationType.ADD, "OBD_CAN", "Req", "Res",
        DiagnosticEndpoint("SRC_CAN", 0x700, 0x708, params), DiagnosticEndpoint("DST_CAN", 0x700, 0x708, params),
        SourceLocation("诊断报文路由", 2))
    workbook = replace(workbook, diagnostic_routes=(route,), doip_mappings=(mapping(0x102, request=0x700, response=0x708),))
    plan = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not plan.errors, plan.issues
    assert plan.diagnostic_added_count == 1
    assert plan.doip_added_count == 1
    validate_generated_output(document, plan)


@pytest.mark.parametrize("changes,code", [({"RouterType": "OTA"}, "DOIP_TESTER_INVALID"),
                                         ({"logicalAddress": "0x10000"}, "WORKBOOK_INTEGER_INVALID"),
                                         ({"Functional": "YES"}, "DOIP_FUNCTIONAL_RESPONSE_INVALID"),
                                         ({"CANBus": "UNKNOWN"}, "DOIP_CHANNEL_INVALID")])
def test_reader_rejects_invalid_mapping(workbook_factory, changes, code):
    path = workbook_factory(direct_rows=(), signal_rows=())
    book = load_workbook(path)
    sheet = book.create_sheet("DoIP_to_CAN")
    sheet.append(DOIP_HEADERS)
    row = dict(zip(DOIP_HEADERS, ("0x0E80", "0x0003", "OBD", "DST_CAN", "0x7E0", "0x7E8", None)))
    row.update(changes)
    sheet.append([row.get(k) for k in DOIP_HEADERS])
    book.save(path)
    result = read_workbook(path)
    assert not result.is_valid
    assert any(i.code == code for i in result.issues), result.issues


def test_reader_deduplicates_exact_rows_and_rejects_conflicts(workbook_factory):
    path = workbook_factory(direct_rows=(), signal_rows=())
    book = load_workbook(path)
    sheet = book.create_sheet("DoIP_to_CAN")
    sheet.append(DOIP_HEADERS)
    row = ["0x0E80", "0x0003", "OBD", "DST_CAN", "0x7E0", "0x7E8", None]
    sheet.append(row)
    sheet.append(row)
    book.save(path)
    result = read_workbook(path)
    assert result.is_valid
    assert len(result.data.doip_mappings) == 1
    sheet.append(row[:5] + ["0x709", None])
    book.save(path)
    assert any(i.code == "DOIP_MAPPING_CONFLICT" for i in read_workbook(path).issues)


def functional_sample(doip_case):
    """构造含本机目标的功能广播样例，验证扩展时保留原有目标。"""
    document, workbook, container, group = doip_case
    index = document.build_index()
    get = lambda name: index.find_by_short_name(name)[0]
    path = lambda node: autosar_path(node, document.namespace)
    target = container(get("ExistingTarget").getparent(), "FunctionalTarget", TARGET,
                       {TARGET + "/DoIPTargetAddressValue": str(0xE400), TARGET + "/DoIPTargetAddressVerification": "false"})
    from copy import deepcopy
    activation_refs = get("ExistingActivation").find(f"{{{document.namespace}}}REFERENCE-VALUES")
    target_ref = deepcopy(activation_refs[0])
    target_ref.find(f"{{{document.namespace}}}VALUE-REF").text = path(target)
    activation_refs.append(target_ref)
    channel = container(get("ExistingChannel").getparent(), "FunctionalChannel", CHANNEL,
                        {CHANNEL + "/DoIPChannelId": "8", CHANNEL + "/DoIPTargetMaxBufSize": "4100",
                         CHANNEL + "/DoIPChannelReadyNotification": "false", CHANNEL + "/DoIPTargetMaxPduSize": "4100"},
                        {CHANNEL + "/DoIPChannelSARef": path(get("ExistingTester")), CHANNEL + "/DoIPChannelTARef": path(target)})
    pdu = container(get("DoipRx").getparent(), "FunctionalRx", d.ECUC_PDU,
                    {d.ECUC_PDU_LENGTH: "1500", d.ECUC_PDU_J1939: "false"})
    container(group(channel, "SUB-CONTAINERS"), "Rx", RX, {RX + "/DoIPPduRRxPduId": "8"}, {RX + "/DoIPPduRRxPduRef": path(pdu)})
    for bus, hardware in (("DST_CAN", "TX2"), ("SRC_CAN", "TX")):
        lower = container(pdu.getparent(), bus + "FuncLower", d.ECUC_PDU,
                          {d.ECUC_PDU_LENGTH: "64", d.ECUC_PDU_J1939: "false"})
        upper = container(pdu.getparent(), bus + "FuncUpper", d.ECUC_PDU,
                          {d.ECUC_PDU_LENGTH: "4095", d.ECUC_PDU_J1939: "false"})
        container(get("FrameDST_CANTx").getparent(), bus + "FuncFrame", d.CANIF_TX,
                  {d.CANIF_TX_CAN_ID: str(0x7DF), d.CANIF_TX_CAN_ID_TYPE: "STANDARD_FD_CAN",
                   d.CANIF_TX_CONFIRM_UL: "CAN_TP", d.CANIF_TX_CONFIRM_NAME: "CanTp_TxConfirmation",
                   d.CANIF_TX_TRUNCATION: "true", d.CANIF_TX_DLC: "64"},
                  {d.CANIF_TX_BUFFER_REF: path(get(hardware)), d.CANIF_TX_PDU_REF: path(lower)})
        sdu = container(get("MisleadingNameFalse").find(f"{{{document.namespace}}}SUB-CONTAINERS"),
                        bus + "FuncSdu", d.CANTP_TX,
                        {d.CANTP_TX + "/CanTpTxTaType": "CANTP_CANFD_FUNCTIONAL"}, {d.CANTP_TX_SDU_REF: path(upper)})
        container(group(sdu, "SUB-CONTAINERS"), "NPdu", d.CANTP_TX_NPDU,
                  {d.CANTP_TX_NPDU + "/CanTpTxNPduConfirmationPduId": "9"}, {d.CANTP_TX_NPDU + "/CanTpTxNPduRef": path(lower)})
        leg = container(get("Req").getparent(), bus + "FuncCanPath", d.PDUR_PATH,
                        {d.PDUR_PATH_COMM_TYPE: "TRANSPORT_PROTOCOL"})
        container(group(leg, "SUB-CONTAINERS"), "Target", d.PDUR_DEST,
                  {d.PDUR_DEST_HANDLE: "9", d.PDUR_DEST_DIRECTION: "TRANSMIT", d.PDUR_DEST_ROUTING_TYPE: "GATEWAY_ROUTING",
                   d.PDUR_DEST_PROCESSING: "IMMEDIATE", d.PDUR_DEST_LENGTH_STRATEGY: "UNUSED", d.PDUR_DEST_CROSS_PARTITION: "false"},
                  {d.PDUR_DEST_PDU_REF: path(upper), d.PDUR_DEST_MODULE_REF: path(get("TransportModule")), d.PDUR_QUEUE_REF: path(get("Queue"))})
    func = container(get("Req").getparent(), "FunctionalRequest", d.PDUR_PATH,
                     {d.PDUR_PATH_COMM_TYPE: "TRANSPORT_PROTOCOL"})
    container(group(func, "SUB-CONTAINERS"), "Source", d.PDUR_SRC,
              {d.PDUR_SRC_HANDLE: "9", d.PDUR_SRC_DIRECTION: "RECEIVE"},
              {d.PDUR_SRC_PDU_REF: path(pdu), d.PDUR_SRC_MODULE_REF: path(get("DoipModule"))})
    # 把 DST_CAN 目标作为同业务功能样例，同时留下不使用 Queue 的本机 API 目标。
    updated = document.build_index()
    can_target = next(n for n in updated.find_by_short_name("DST_CANFuncCanPath")[0].iter()
                      if definition_ref(n, document.namespace) == d.PDUR_DEST)
    from copy import deepcopy
    group(func, "SUB-CONTAINERS").append(deepcopy(can_target))
    container(group(func, "SUB-CONTAINERS"), "LocalDcm", d.PDUR_DEST,
              {d.PDUR_DEST_HANDLE: "12", d.PDUR_DEST_DIRECTION: "RECEIVE", d.PDUR_DEST_ROUTING_TYPE: "API_FORWARDING"},
              {d.PDUR_DEST_PDU_REF: path(get("DoipRx")), d.PDUR_DEST_MODULE_REF: path(get("DoipModule"))})
    return document, workbook


@pytest.mark.parametrize("address", [0xE400, 0xE401])
def test_functional_fanout_uses_one_rx_and_preserves_local_destination(doip_case, address):
    document, workbook = functional_sample(doip_case)
    routes = (mapping(address, "DST_CAN", 0x7DF, None, True), mapping(address, "SRC_CAN", 0x7DF, None, True))
    plan = TransactionCoordinator(document, replace(workbook, doip_mappings=routes)).plan_and_apply()
    assert not plan.errors, plan.issues
    validate_generated_output(document, plan)
    assert not any(op.definition_ref == TX for op in plan.operations)
    assert plan.doip_added_count == (1 if address == 0xE400 else 2)
    assert plan.doip_existing_count == (1 if address == 0xE400 else 0)
    assert sum(op.definition_ref == CHANNEL for op in plan.operations) == (0 if address == 0xE400 else 1)
    assert sum(op.definition_ref == RX for op in plan.operations) == (0 if address == 0xE400 else 1)
    assert sum(op.definition_ref == d.PDUR_PATH for op in plan.operations) == (0 if address == 0xE400 else 1)
    assert document.build_index().find_by_short_name("LocalDcm")
    again = TransactionCoordinator(document, replace(workbook, doip_mappings=routes)).plan_and_apply()
    assert not again.errors, again.issues
    assert again.doip_existing_count == 2
    assert not again.operations


def test_new_name_conflict_is_reported_before_apply(doip_case):
    document, workbook, container, group = doip_case
    existing = document.build_index().find_by_short_name("ExistingChannel")[0]
    container(existing.getparent(), "GWT_DoIPChannel_OBD_DoIp101", CHANNEL)
    plan = TransactionCoordinator(document, replace(workbook, doip_mappings=(mapping(0x101),))).plan_and_apply()
    assert any(i.code == "DOIP_CONFIGURATION_CONFLICT" for i in plan.errors)
    assert not document.build_index().find_by_short_name("GWT_Pdu_OBD_DoIp101_Rx")


def test_symbolic_name_in_other_path_blocks_new_destination(doip_case):
    document, workbook, container, group = doip_case
    existing = document.build_index().find_by_short_name("Res")[0]
    path = container(existing.getparent(), "UnrelatedInterface", d.PDUR_PATH,
                     {d.PDUR_PATH_COMM_TYPE: "COMMUNICATION_INTERFACE"})
    container(group(path, "SUB-CONTAINERS"), "PduRDestPdu_OBD_DoIp101ToDST_CAN_Tx", d.PDUR_DEST)
    plan = TransactionCoordinator(document, replace(workbook, doip_mappings=(mapping(0x101),))).plan_and_apply()
    assert any("符号名" in issue.message for issue in plan.errors)
    assert not document.build_index().find_by_short_name("GWT_Pdu_OBD_DoIp101_Rx")


def test_channel_name_uses_existing_safe_name_rule(doip_case):
    import re
    document, workbook, _, _ = doip_case
    references = tuple(replace(r, channel_name="DST-CAN") if r.channel_name == "DST_CAN" else r
                       for r in workbook.reference_data)
    workbook = replace(workbook, reference_data=references, doip_mappings=(mapping(0x101, bus="DST-CAN"),))
    plan = TransactionCoordinator(document, workbook).plan_and_apply()
    assert not plan.errors, plan.issues
    assert all(re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", op.short_name)
               for op in plan.operations if op.short_name)
    validate_generated_output(document, plan)


def test_physical_fanout_names_every_destination_by_bus(doip_case):
    document, workbook, _, _ = doip_case
    routes = (mapping(0x101), mapping(0x101, "SRC_CAN", 0x7E8, 0x7E0))
    plan = TransactionCoordinator(document, replace(workbook, doip_mappings=routes)).plan_and_apply()
    assert not plan.errors, plan.issues
    assert sum(op.definition_ref == CHANNEL for op in plan.operations) == 1
    destinations = [op for op in plan.operations if op.definition_ref == d.PDUR_DEST]
    assert len(destinations) == 4
    assert len({op.short_name for op in destinations}) == 4
    assert sum("ToDST_CAN" in op.short_name for op in destinations) == 1
    assert sum("ToSRC_CAN" in op.short_name for op in destinations) == 1
    validate_generated_output(document, plan)


def test_multiple_addresses_share_response_path_with_unique_destinations(doip_case):
    document, workbook, _, _ = doip_case
    routes = (mapping(0x101), mapping(0x102))
    plan = TransactionCoordinator(document, replace(workbook, doip_mappings=routes)).plan_and_apply()
    assert not plan.errors, plan.issues
    response = document.build_index().find_by_short_name("Res")[0]
    response_path = autosar_path(response, document.namespace)
    added = [op for op in plan.operations if op.definition_ref == d.PDUR_DEST and op.parent_path == response_path]
    assert {op.short_name for op in added} == {
        "PduRDestPdu_OBD_DST_CANToDoIp101_Tx", "PduRDestPdu_OBD_DST_CANToDoIp102_Tx"}
    assert len({op.parameter(d.PDUR_DEST_HANDLE) for op in added}) == 2
    validate_generated_output(document, plan)


@pytest.mark.parametrize("conflict", ["queue", "frame_property", "ambiguous_frame"])
def test_skip_never_downgrades_existing_configuration_conflict(doip_case, conflict):
    from copy import deepcopy
    from davinci_gw.arxml.namespace import qualified
    document, workbook, container, group = doip_case
    index = document.build_index()
    if conflict == "queue":
        queue = index.find_by_short_name("Queue")[0]
        other = deepcopy(queue)
        other.find(qualified(document.namespace, "SHORT-NAME")).text = "ConflictingQueue"
        queue.getparent().append(other)
        dest = index.find_by_short_name("CanTarget")[0]
        for value in dest.iter():
            if definition_ref(value, document.namespace) == d.PDUR_QUEUE_REF:
                value.find(qualified(document.namespace, "VALUE-REF")).text = autosar_path(other, document.namespace)
    else:
        frame = index.find_by_short_name("FrameDST_CANRx")[0]
        if conflict == "ambiguous_frame":
            other = deepcopy(frame)
            other.find(qualified(document.namespace, "SHORT-NAME")).text = "AmbiguousFrame"
            frame.getparent().append(other)
        else:
            for value in frame.iter():
                if definition_ref(value, document.namespace) == d.CANIF_RX_DLC_CHECK:
                    value.find(qualified(document.namespace, "VALUE")).text = "true"
    plan = TransactionCoordinator(document, workbook, doip_missing_policy="SKIP").plan_and_apply()
    assert any(i.code == "DOIP_PREREQUISITE_CONFLICT" for i in plan.errors), plan.issues
    assert not any(i.code == "DOIP_PREREQUISITE_SKIPPED" for i in plan.issues)
    assert not plan.operations


@pytest.mark.parametrize("invalid_sample", ["path_type", "source_direction"])
def test_invalid_request_sample_blocks_new_channel(invalid_sample, doip_case):
    from davinci_gw.arxml.namespace import qualified
    document, workbook, _, _ = doip_case
    path = document.build_index().find_by_short_name("ExistingDoipRequest")[0]
    field = d.PDUR_PATH_COMM_TYPE if invalid_sample == "path_type" else d.PDUR_SRC_DIRECTION
    for value in path.iter():
        if definition_ref(value, document.namespace) == field:
            value.find(qualified(document.namespace, "VALUE")).text = "COMMUNICATION_INTERFACE" if invalid_sample == "path_type" else "TRANSMIT"
    plan = TransactionCoordinator(document, replace(workbook, doip_mappings=(mapping(0x101),)),
                                  doip_missing_policy="SKIP").plan_and_apply()
    assert any(i.code == "DOIP_CONFIGURATION_CONFLICT" for i in plan.errors), plan.issues
    assert not document.build_index().find_by_short_name("GWT_DoIPChannel_OBD_DoIp101")


@pytest.mark.parametrize("functional", [False, True])
@pytest.mark.parametrize("ambiguity", ["duplicate_container", "multiple_references"])
def test_skip_blocks_ambiguous_cantp_npdu(doip_case, functional, ambiguity):
    from copy import deepcopy
    from davinci_gw.arxml.namespace import qualified
    if functional:
        document, workbook = functional_sample(doip_case)
        workbook = replace(workbook, doip_mappings=(mapping(0xE400, "DST_CAN", 0x7DF, None, True),))
        owner = document.build_index().find_by_short_name("DST_CANFuncSdu")[0]
    else:
        document, workbook, _, _ = doip_case
        channel = document.build_index().find_by_short_name("MisleadingNameFalse")[0]
        owner = next(n for n in channel.iter() if definition_ref(n, document.namespace) == d.CANTP_TX)
    npdu = next(n for n in owner.iter() if definition_ref(n, document.namespace) == d.CANTP_TX_NPDU)
    if ambiguity == "duplicate_container":
        duplicate = deepcopy(npdu)
        duplicate.find(qualified(document.namespace, "SHORT-NAME")).text = "AmbiguousNPdu"
        npdu.getparent().append(duplicate)
    else:
        references = npdu.find(qualified(document.namespace, "REFERENCE-VALUES"))
        references.append(deepcopy(references[0]))
    plan = TransactionCoordinator(document, workbook, doip_missing_policy="SKIP").plan_and_apply()
    assert any(i.code == "DOIP_PREREQUISITE_CONFLICT" for i in plan.errors), plan.issues
    assert not any(i.code == "DOIP_PREREQUISITE_SKIPPED" for i in plan.issues)
    assert not plan.operations


def standard_input(workbook_factory, rows):
    path = workbook_factory(direct_rows=(), signal_rows=())
    book = load_workbook(path)
    sheet = book.create_sheet("DoIP_to_CAN")
    sheet.append(DOIP_HEADERS)
    for route in rows:
        sheet.append([route.tester, route.logical_address, route.router_type, route.can_bus,
                      route.request_id, route.response_id, "YES" if route.functional else None])
    book.save(path)
    return path


def test_facade_does_not_create_session_until_choice_and_keeps_skip_for_commit(doip_case, workbook_factory, tmp_path):
    from davinci_gw.application.facade import GatewayFacade
    from davinci_gw.contracts import OperationStatus, UpdateRequestDto
    document, _, _, _ = doip_case
    baseline = tmp_path / "baseline.arxml"
    document.tree.write(str(baseline), encoding="UTF-8", xml_declaration=True)
    config = standard_input(workbook_factory, (mapping(0x101), mapping(0x102, request=0x700, response=0x708)))
    facade = GatewayFacade()
    undecided = facade.prepare(UpdateRequestDto(config, baseline))
    assert undecided.status is OperationStatus.VALIDATION_FAILED
    assert undecided.session_id is None
    assert any(i.code == "DOIP_PREREQUISITE_CHOICE_REQUIRED" for i in undecided.issues)
    selected = facade.prepare(UpdateRequestDto(config, baseline, doip_missing_policy="SKIP"))
    assert selected.status is OperationStatus.SUCCESS, selected.issues
    record = facade.sessions._records[selected.session_id]
    assert record.request.doip_missing_policy == "SKIP"
    output = tmp_path / "output.arxml"
    result = facade.commit_prepared(selected.session_id, str(output))
    assert result.status is OperationStatus.SUCCESS, result.issues
    assert output.is_file()
    assert any(i.code == "DOIP_PREREQUISITE_SKIPPED" for i in selected.preview.issues)


def test_no_com_is_allowed_without_signal_and_remains_required_with_signal(doip_case, workbook_factory, tmp_path):
    from davinci_gw.application.generate import prepare_transaction
    from davinci_gw.application.validate import validate_inputs
    document, _, _, _ = doip_case
    module = document.build_index().find_by_short_name("Com")[0]
    module.getparent().remove(module)
    baseline = tmp_path / "without_com.arxml"
    document.tree.write(str(baseline), encoding="UTF-8", xml_declaration=True)
    config = standard_input(workbook_factory, (mapping(),))
    prepared = prepare_transaction(config, baseline)
    assert prepared.validation.is_valid, prepared.validation.issues
    assert not prepared.plan.errors, prepared.plan.issues
    output = tmp_path / "without_com_output.arxml"
    prepared.document.write_atomic(output, validator=lambda doc: validate_generated_output(doc, prepared.plan))
    assert output.is_file()
    # 默认信号执行行仍要求 Com，不会放宽信号路由的必要模块检查。
    signal_config = workbook_factory(filename="signals_v1.0.xlsx", direct_rows=())
    validation = validate_inputs(signal_config, baseline)
    assert any(i.code == "ARXML_STRUCTURE_INVALID" and "Com" in i.message for i in validation.errors)


def test_gui_asks_choice_and_clears_it_when_input_changes(qtbot, workbook_factory, arxml_factory):
    from UI.view_model import GatewayViewModel
    from tests.unit.test_gui import _FakeRunner
    from davinci_gw.application.sessions import fingerprint_file
    from davinci_gw.contracts import InputFileDto, IssueDto, OperationStatus, PreparedSessionDto, PreviewResultDto
    from UI.state import GuiState
    config, baseline = workbook_factory(), arxml_factory()
    runner = _FakeRunner()
    calls = []
    runner.start_preview = lambda *args: calls.append(args)
    vm = GatewayViewModel(runner)
    vm.set_config_path(str(config))
    vm.set_baseline_path(str(baseline))
    issue = IssueDto("DOIP_PREREQUISITE_CHOICE_REQUIRED", "请选择")
    inputs = (InputFileDto("CONFIG", str(config), fingerprint_file(config)),
              InputFileDto("BASELINE", str(baseline), fingerprint_file(baseline)))
    result = PreparedSessionDto("id", OperationStatus.VALIDATION_FAILED,
        preview=PreviewResultDto("id", OperationStatus.VALIDATION_FAILED, issues=(issue,), input_files=inputs))
    with qtbot.waitSignal(vm.doip_choice_required):
        runner.preview_finished.emit(result)
    assert not vm.controls.generate_enabled
    vm.choose_doip_missing_policy("SKIP")
    assert calls[-1][-1] == "SKIP"
    vm.state = GuiState.PREVIEW_INVALID
    config.write_bytes(config.read_bytes() + b"changed")
    vm.start_preview()
    assert len(calls[-1]) == 3
    assert vm.doip_missing_policy is None


def test_mcp_preview_requires_choice_then_passes_explicit_skip(doip_case, workbook_factory, tmp_path):
    import asyncio
    from davinci_gw.mcp.service import GatewayMcpService
    document, _, _, _ = doip_case
    baseline = tmp_path / "baseline.arxml"
    document.tree.write(str(baseline), encoding="UTF-8", xml_declaration=True)
    config = standard_input(workbook_factory, (mapping(), mapping(0x102, request=0x700, response=0x708)))
    async def run():
        service = GatewayMcpService()
        try:
            undecided = await service.preview_gateway_update(str(config), str(baseline))
            assert undecided["requires_doip_prerequisite_choice"]
            assert undecided["preparation_id"] is None
            selected = await service.preview_gateway_update(str(config), str(baseline), "SKIP")
            assert selected["success"], selected
            assert selected["preparation_id"]
            assert selected["doip_missing_policy"] == "SKIP"
        finally:
            await service.close()
    asyncio.run(run())


def test_cli_blocks_undecided_generation_and_reports_explicit_skip(doip_case, workbook_factory, tmp_path, capsys):
    from davinci_gw.cli import main
    document, _, _, _ = doip_case
    baseline = tmp_path / "baseline.arxml"
    document.tree.write(str(baseline), encoding="UTF-8", xml_declaration=True)
    config = standard_input(workbook_factory, (mapping(0x101), mapping(0x102, request=0x700, response=0x708)))
    output = tmp_path / "output.arxml"
    arguments = ["--config", str(config), "--baseline", str(baseline)]
    assert main(["generate", *arguments, "--output", str(output)]) == 2
    assert not output.exists()
    assert "未生成输出" in capsys.readouterr().out
    assert main(["generate", *arguments, "--output", str(output), "--doip-missing-policy", "SKIP"]) == 0
    text = capsys.readouterr().out
    assert "DoIP→CAN：新增或补全 1，已有 0，跳过 1" in text
    assert output.is_file()
