"""测试夹具：动态生成小型标准工作簿和脱敏 ARXML。"""

from __future__ import annotations

from collections.abc import Callable, Iterable, Mapping, Sequence
from copy import copy, deepcopy
from pathlib import Path

import pytest
from lxml import etree
from openpyxl import Workbook

from davinci_gw.modules import definitions as defs
from davinci_gw.application.generate import generate_inputs


DIRECT_HEADERS = (
    "源网段报文名称", "源网段报文CANID", "源网段报文Length", "源网段报文类型",
    "源网段CAN通道", "源网段RxIndicationUL", "源网段报文Checksum使能",
    "源网段报文Dlc Check使能", "目标网段报文名称", "目标网段报文CANID",
    "目标网段报文Length", "目标网段报文类型", "目标网段CAN通道",
    "目标网段报文Checksum使能", "目标网段报文PnFilter使能",
    "目标网段报文Truncation使能", "路由Length Strategy功能选择", "操作类型",
)
SIGNAL_HEADERS = (
    "源网段", "源报文名", "源信号名", "字节序", "超时值", "超时时间",
    "源信号组合名", "目标网段", "目标报文名", "目标信号名", "操作类型",
)
REFERENCE_HEADERS = ("CAN通道名称", "CanIfTxBuffer名称", "CanIfHrh名称")


def direct_row(**changes: object) -> dict[str, object]:
    """返回一条完整、有效的直接报文 ADD 数据。"""
    row: dict[str, object] = {
        "源网段报文名称": "SRC_MSG", "源网段报文CANID": "0x100",
        "源网段报文Length": 8, "源网段报文类型": "STANDARD_CAN",
        "源网段CAN通道": "SRC_CAN", "源网段RxIndicationUL": "PDUR",
        "源网段报文Checksum使能": None, "源网段报文Dlc Check使能": "Enable",
        "目标网段报文名称": "DST_MSG", "目标网段报文CANID": 0x200,
        "目标网段报文Length": 8, "目标网段报文类型": "STANDARD_FD_CAN",
        "目标网段CAN通道": "DST_CAN", "目标网段报文Checksum使能": None,
        "目标网段报文PnFilter使能": None, "目标网段报文Truncation使能": "Enable",
        "路由Length Strategy功能选择": "IGNORE", "操作类型": "ADD",
    }
    row.update(changes)
    return row


def signal_row(**changes: object) -> dict[str, object]:
    """返回一条完整、有效的信号 ADD 数据。"""
    row: dict[str, object] = {
        "源网段": "SRC", "源报文名": "SRC_MSG", "源信号名": "SRC_SIG",
        "字节序": "Motorola LSB", "超时值": None, "超时时间": 2,
        "源信号组合名": None, "目标网段": "DST", "目标报文名": "DST_MSG",
        "目标信号名": "DST_SIG", "操作类型": "ADD",
    }
    row.update(changes)
    return row


def _append_mapping(sheet: object, headers: Sequence[str], row: Mapping[str, object]) -> None:
    sheet.append([row.get(header) for header in headers])


@pytest.fixture
def workbook_factory(tmp_path: Path) -> Callable[..., Path]:
    """返回可按场景生成工作簿的工厂。"""
    def create(
        *,
        filename: str = "gateway_v4.84.xlsx",
        direct_rows: Iterable[Mapping[str, object]] | None = None,
        signal_rows: Iterable[Mapping[str, object]] | None = None,
        reference_rows: Iterable[Mapping[str, object]] | None = None,
        direct_headers: Sequence[str] = DIRECT_HEADERS,
        signal_headers: Sequence[str] = SIGNAL_HEADERS,
        reference_headers: Sequence[str] = REFERENCE_HEADERS,
        omitted_sheets: Iterable[str] = (),
        formatted_blank_rows: int = 0,
    ) -> Path:
        path = tmp_path / filename
        workbook = Workbook()
        workbook.remove(workbook.active)
        omitted = set(omitted_sheets)
        if "引用数据" not in omitted:
            sheet = workbook.create_sheet("引用数据")
            sheet.append(list(reference_headers))
            rows = reference_rows if reference_rows is not None else (
                {"CAN通道名称": "SRC_CAN", "CanIfTxBuffer名称": "TX", "CanIfHrh名称": "HRH"},
                {"CAN通道名称": "DST_CAN", "CanIfTxBuffer名称": "TX2", "CanIfHrh名称": "HRH2"},
            )
            for row in rows:
                _append_mapping(sheet, reference_headers, row)
        if "直接报文路由" not in omitted:
            sheet = workbook.create_sheet("直接报文路由")
            sheet.append(list(direct_headers))
            for row in (direct_rows if direct_rows is not None else (direct_row(),)):
                _append_mapping(sheet, direct_headers, row)
            for index in range(formatted_blank_rows):
                sheet.cell(row=sheet.max_row + 1, column=1).fill = copy(sheet["A1"].fill)
        if "信号路由" not in omitted:
            sheet = workbook.create_sheet("信号路由")
            sheet.append(list(signal_headers))
            for row in (signal_rows if signal_rows is not None else (signal_row(),)):
                _append_mapping(sheet, signal_headers, row)
        version = workbook.create_sheet("版本")
        version.append(["V01", "V02"])
        workbook.save(path)
        return path
    return create


@pytest.fixture
def arxml_factory(tmp_path: Path) -> Callable[..., Path]:
    """返回可生成模块顺序及异常场景的 ARXML 工厂。"""
    def create(
        *,
        filename: str = "baseline.arxml",
        order: Sequence[str] = ("CanIf", "Com", "EcuC", "PduR"),
        missing: str | None = None,
        duplicate: str | None = None,
        extra: bool = False,
        routing_groups: Mapping[str, Sequence[str]] | None = None,
    ) -> Path:
        ns = "http://autosar.org/schema/r4.0"
        xsi = "http://www.w3.org/2001/XMLSchema-instance"
        root = etree.Element(f"{{{ns}}}AUTOSAR", nsmap={None: ns, "xsi": xsi})
        root.set(f"{{{xsi}}}schemaLocation", f"{ns} AUTOSAR_00049.xsd")
        root.append(etree.Comment("必须保留的注释"))
        packages = etree.SubElement(root, f"{{{ns}}}AR-PACKAGES")
        package = etree.SubElement(packages, f"{{{ns}}}AR-PACKAGE")
        etree.SubElement(package, f"{{{ns}}}SHORT-NAME").text = "Cfg"
        elements = etree.SubElement(package, f"{{{ns}}}ELEMENTS")
        uuid_counter = 0

        def new_uuid() -> str:
            nonlocal uuid_counter
            uuid_counter += 1
            return f"00000000-0000-0000-0000-{uuid_counter:012d}"

        def add_parameter(
            node: etree._Element, definition: str, value: object, *, kind: str = "TEXTUAL",
        ) -> None:
            group = node.find(f"{{{ns}}}PARAMETER-VALUES")
            if group is None:
                group = etree.SubElement(node, f"{{{ns}}}PARAMETER-VALUES")
            entry = etree.SubElement(group, f"{{{ns}}}ECUC-{kind}-PARAM-VALUE")
            etree.SubElement(entry, f"{{{ns}}}DEFINITION-REF").text = definition
            etree.SubElement(entry, f"{{{ns}}}VALUE").text = str(value)

        def add_reference(node: etree._Element, definition: str, value: str) -> None:
            group = node.find(f"{{{ns}}}REFERENCE-VALUES")
            if group is None:
                group = etree.SubElement(node, f"{{{ns}}}REFERENCE-VALUES")
            entry = etree.SubElement(group, f"{{{ns}}}ECUC-REFERENCE-VALUE")
            definition_node = etree.SubElement(entry, f"{{{ns}}}DEFINITION-REF")
            definition_node.set("DEST", "ECUC-REFERENCE-DEF")
            definition_node.text = definition
            value_node = etree.SubElement(entry, f"{{{ns}}}VALUE-REF")
            value_node.set("DEST", "ECUC-CONTAINER-VALUE")
            value_node.text = value

        def add_container(
            group: etree._Element, short_name: str, definition: str,
            parameters: Sequence[tuple[str, object, str]] = (),
            references: Sequence[tuple[str, str]] = (),
        ) -> etree._Element:
            node = etree.SubElement(group, f"{{{ns}}}ECUC-CONTAINER-VALUE", UUID=new_uuid())
            etree.SubElement(node, f"{{{ns}}}SHORT-NAME").text = short_name
            etree.SubElement(node, f"{{{ns}}}DEFINITION-REF").text = definition
            for parameter, value, kind in parameters:
                add_parameter(node, parameter, value, kind=kind)
            for reference, value in references:
                add_reference(node, reference, value)
            return node

        def subcontainers(node: etree._Element) -> etree._Element:
            return etree.SubElement(node, f"{{{ns}}}SUB-CONTAINERS")

        def add_routing_structures(name: str, containers: etree._Element) -> None:
            if name == "EcuC":
                collection = add_container(containers, "EcucPduCollection", "/MICROSAR/EcuC/EcucPduCollection")
                add_container(
                    subcontainers(collection), "ExistingPdu_Rx", defs.ECUC_PDU,
                    ((defs.ECUC_PDU_LENGTH, 8, "NUMERICAL"), (defs.ECUC_PDU_J1939, "false", "TEXTUAL")),
                )
            elif name == "CanIf":
                init = add_container(containers, "CanIfInitCfg", "/MICROSAR/CanIf/CanIfInitCfg")
                init_subs = subcontainers(init)
                hoh = add_container(init_subs, "CanIfInitHohCfg", "/MICROSAR/CanIf/CanIfInitCfg/CanIfInitHohCfg")
                hoh_subs = subcontainers(hoh)
                add_container(hoh_subs, "HRH", defs.CANIF_HRH)
                add_container(hoh_subs, "HRH2", defs.CANIF_HRH)
                add_container(init_subs, "TX", defs.CANIF_BUFFER)
                add_container(init_subs, "TX2", defs.CANIF_BUFFER)
                rx_params = (
                    (defs.CANIF_RX_CAN_ID, 1, "NUMERICAL"),
                    (defs.CANIF_RX_CAN_ID_TYPE, "STANDARD_CAN", "TEXTUAL"),
                    (defs.CANIF_RX_DLC, 8, "NUMERICAL"),
                    (defs.CANIF_RX_INDICATION_NAME, "PduR_CanIfRxIndication", "TEXTUAL"),
                    (defs.CANIF_RX_INDICATION_UL, "PDUR", "TEXTUAL"),
                    (defs.CANIF_RX_HANDLE, 0, "NUMERICAL"),
                    (defs.CANIF_RX_READ_DATA, "false", "TEXTUAL"),
                    (defs.CANIF_RX_READ_NOTIFY, "false", "TEXTUAL"),
                    (defs.CANIF_RX_DLC_CHECK, "true", "TEXTUAL"),
                    (defs.CANIF_RX_TYPE, "STATIC", "TEXTUAL"),
                )
                add_container(
                    init_subs, "ExistingRx", defs.CANIF_RX, rx_params,
                    ((defs.CANIF_RX_HRH_REF, "/Cfg/CanIf/CanIfInitCfg/CanIfInitHohCfg/HRH"),
                     (defs.CANIF_RX_PDU_REF, "/Cfg/EcuC/EcucPduCollection/ExistingPdu_Rx")),
                )
                tx_params = (
                    (defs.CANIF_TX_CAN_ID, 2, "NUMERICAL"),
                    (defs.CANIF_TX_CAN_ID_TYPE, "STANDARD_FD_CAN", "TEXTUAL"),
                    (defs.CANIF_TX_DLC, 8, "NUMERICAL"),
                    (defs.CANIF_TX_CONFIRM_NAME, "NULL_PTR", "TEXTUAL"),
                    (defs.CANIF_TX_CONFIRM_UL, "NONE", "TEXTUAL"),
                    (defs.CANIF_TX_HANDLE, 0, "NUMERICAL"),
                    (defs.CANIF_TX_READ_NOTIFY, "false", "TEXTUAL"),
                    (defs.CANIF_TX_TYPE, "STATIC", "TEXTUAL"),
                    (defs.CANIF_TX_TRUNCATION, "true", "TEXTUAL"),
                )
                add_container(
                    init_subs, "ExistingTx", defs.CANIF_TX, tx_params,
                    ((defs.CANIF_TX_BUFFER_REF, "/Cfg/CanIf/CanIfInitCfg/TX2"),
                     (defs.CANIF_TX_PDU_REF, "/Cfg/EcuC/EcucPduCollection/ExistingPdu_Rx")),
                )
            elif name == "PduR":
                add_container(
                    containers, "CanIf", defs.PDUR_BSW_MODULE,
                    references=((defs.PDUR_BSW_MODULE_REF, "/Cfg/CanIf"),),
                )
                tables = add_container(containers, "PduRRoutingTables", "/MICROSAR/PduR/PduRRoutingTables")
                tables_subs = subcontainers(tables)
                add_container(tables_subs, "Lock", "/MICROSAR/PduR/PduRRoutingTables/PduRLock")
                table = add_container(tables_subs, "PduRRoutingTable", "/MICROSAR/PduR/PduRRoutingTables/PduRRoutingTable")
                path = add_container(
                    subcontainers(table), "ExistingPath", defs.PDUR_PATH,
                    ((defs.PDUR_PATH_COMM_TYPE, "COMMUNICATION_INTERFACE", "TEXTUAL"),
                     (defs.PDUR_PATH_MULTICORE, "false", "TEXTUAL")),
                    ((defs.PDUR_PATH_LOCK_REF, "/Cfg/PduR/PduRRoutingTables/Lock"),),
                )
                path_subs = subcontainers(path)
                add_container(
                    path_subs, "ExistingSrc", defs.PDUR_SRC,
                    ((defs.PDUR_SRC_HANDLE, 0, "NUMERICAL"),
                     (defs.PDUR_SRC_DIRECTION, "RECEIVE", "TEXTUAL")),
                    ((defs.PDUR_SRC_PDU_REF, "/Cfg/EcuC/EcucPduCollection/ExistingPdu_Rx"),
                     (defs.PDUR_SRC_MODULE_REF, "/Cfg/PduR/CanIf")),
                )
                add_container(
                    path_subs, "ExistingDest", defs.PDUR_DEST,
                    ((defs.PDUR_DEST_HANDLE, 0, "NUMERICAL"),
                     (defs.PDUR_DEST_DIRECTION, "TRANSMIT", "TEXTUAL"),
                     (defs.PDUR_DEST_ROUTING_TYPE, "GATEWAY_ROUTING", "TEXTUAL"),
                     (defs.PDUR_DEST_PROCESSING, "IMMEDIATE", "TEXTUAL"),
                     (defs.PDUR_DEST_LENGTH_STRATEGY, "IGNORE", "TEXTUAL"),
                     (defs.PDUR_DEST_CROSS_PARTITION, "false", "TEXTUAL"),
                     (defs.PDUR_DEST_DATA_PROVISION, "PDUR_DIRECT", "TEXTUAL"),
                     (defs.PDUR_DEST_CONFIRMATION, "false", "TEXTUAL")),
                    ((defs.PDUR_DEST_PDU_REF, "/Cfg/EcuC/EcucPduCollection/ExistingPdu_Rx"),
                     (defs.PDUR_DEST_MODULE_REF, "/Cfg/PduR/CanIf")),
                )
                group_members = routing_groups if routing_groups is not None else {
                    "DefaultRoutingGroup": (
                        "/Cfg/PduR/PduRRoutingTables/PduRRoutingTable/ExistingPath/ExistingDest",
                    ),
                }
                for group_name, members in group_members.items():
                    routing_group = add_container(
                        tables_subs, group_name, defs.PDUR_ROUTING_GROUP,
                        ((defs.PDUR_ROUTING_GROUP_ENABLED, "true", "NUMERICAL"),
                         (defs.PDUR_ROUTING_GROUP_ID, 1, "NUMERICAL")),
                    )
                    for member in members:
                        add_reference(routing_group, defs.PDUR_ROUTING_GROUP_DEST_REF, member)
            elif name == "Com":
                config = add_container(containers, "ComConfig", "/MICROSAR/Com/ComConfig")
                config_subs = subcontainers(config)
                src_signal = add_container(config_subs, "SRC_SIG_oSRC_MSG_oSRC_Rx", defs.COM_SIGNAL)
                add_parameter(src_signal, f"{defs.COM_SIGNAL}/ComBitPosition", 0, kind="NUMERICAL")
                add_parameter(src_signal, defs.COM_SIGNAL_ACCESS, "ACCESS_UNCLEAR")
                dst_signal = add_container(config_subs, "DST_SIG_oDST_MSG_oDST_Tx", defs.COM_SIGNAL)
                add_parameter(dst_signal, f"{defs.COM_SIGNAL}/ComBitPosition", 0, kind="NUMERICAL")
                add_parameter(dst_signal, defs.COM_SIGNAL_ACCESS, "ACCESS_UNCLEAR")
                dst_signal_2 = add_container(config_subs, "DST_SIG_2_oDST_MSG_oDST_Tx", defs.COM_SIGNAL)
                add_parameter(dst_signal_2, f"{defs.COM_SIGNAL}/ComBitPosition", 8, kind="NUMERICAL")
                timeout_template = add_container(config_subs, "TimeoutTemplate", defs.COM_SIGNAL)
                add_parameter(timeout_template, defs.COM_TIMEOUT_ACTION, "REPLACE")
                add_parameter(timeout_template, defs.COM_TIMEOUT, "0.5", kind="NUMERICAL")
                add_parameter(timeout_template, defs.COM_TIMEOUT_SUBSTITUTION, "0")
                src_ipdu = add_container(config_subs, "SRC_MSG_oSRC_Rx", defs.COM_IPDU)
                add_reference(src_ipdu, defs.COM_IPDU_SIGNAL_REF, "/Cfg/Com/ComConfig/SRC_SIG_oSRC_MSG_oSRC_Rx")
                dst_ipdu = add_container(config_subs, "DST_MSG_oDST_Tx", defs.COM_IPDU)
                add_reference(dst_ipdu, defs.COM_IPDU_SIGNAL_REF, "/Cfg/Com/ComConfig/DST_SIG_oDST_MSG_oDST_Tx")
                add_reference(dst_ipdu, defs.COM_IPDU_SIGNAL_REF, "/Cfg/Com/ComConfig/DST_SIG_2_oDST_MSG_oDST_Tx")
                mapping = add_container(config_subs, "TemplateMapping", defs.COM_GW_MAPPING)
                mapping_subs = subcontainers(mapping)
                source = add_container(mapping_subs, "ComGwSource", defs.COM_GW_SOURCE)
                source_inner = add_container(subcontainers(source), "ComGwSignal", defs.COM_GW_SOURCE_SIGNAL)
                add_reference(source_inner, defs.COM_GW_SOURCE_SIGNAL_REF, "/Cfg/Com/ComConfig/TimeoutTemplate")
                destination = add_container(mapping_subs, "ComGwDestination", defs.COM_GW_DEST)
                destination_inner = add_container(subcontainers(destination), "ComGwSignal", defs.COM_GW_DEST_SIGNAL)
                add_reference(destination_inner, defs.COM_GW_DEST_SIGNAL_REF, "/Cfg/Com/ComConfig/TimeoutTemplate")

        def add_module(name: str) -> None:
            node = etree.SubElement(elements, f"{{{ns}}}ECUC-MODULE-CONFIGURATION-VALUES")
            node.set("VENDOR-EXT", "keep")
            etree.SubElement(node, f"{{{ns}}}SHORT-NAME").text = name
            ref = etree.SubElement(node, f"{{{ns}}}DEFINITION-REF")
            ref.set("DEST", "ECUC-MODULE-DEF")
            ref.text = f"/MICROSAR/{name}"
            containers = etree.SubElement(node, f"{{{ns}}}CONTAINERS")
            child = etree.SubElement(containers, f"{{{ns}}}ECUC-CONTAINER-VALUE")
            etree.SubElement(child, f"{{{ns}}}SHORT-NAME").text = "SharedName"
            values = etree.SubElement(child, f"{{{ns}}}REFERENCE-VALUES")
            value = etree.SubElement(values, f"{{{ns}}}ECUC-REFERENCE-VALUE")
            etree.SubElement(value, f"{{{ns}}}DEFINITION-REF").text = "/Defs/Reference"
            etree.SubElement(value, f"{{{ns}}}VALUE-REF").text = "/Cfg/CanIf"
            add_routing_structures(name, containers)

        for module in order:
            if module != missing:
                add_module(module)
                if module == duplicate:
                    add_module(module)
        if extra:
            add_module("CanSM")
        etree.SubElement(elements, f"{{{ns}}}VENDOR-UNKNOWN").text = "preserve"
        path = tmp_path / filename
        etree.ElementTree(root).write(path, encoding="UTF-8", xml_declaration=True, pretty_print=True)
        return path
    return create


@pytest.fixture
def lin_target_arxml_factory(
    tmp_path: Path, arxml_factory: Callable[..., Path],
) -> Callable[..., Path]:
    """生成逻辑 LIN 节点名与物理通道名无固定映射的目标端 DBC 基线。"""
    counter = 0

    def create(*, channels: Sequence[str] = ("LIN04",)) -> Path:
        nonlocal counter
        counter += 1
        path = arxml_factory(filename=f"lin_target_{counter}.arxml")
        tree = etree.parse(str(path))
        namespace = etree.QName(tree.getroot()).namespace

        def named(short_name: str) -> etree._Element:
            return next(
                node for node in tree.getroot().iter()
                if node.findtext(f"{{{namespace}}}SHORT-NAME") == short_name
            )

        original_ipdu = named("DST_MSG_oDST_Tx")
        original_signals = (named("DST_SIG_oDST_MSG_oDST_Tx"),
                            named("DST_SIG_2_oDST_MSG_oDST_Tx"))
        parent = original_ipdu.getparent()
        assert parent is not None and all(signal.getparent() is parent for signal in original_signals)
        parent.remove(original_ipdu)
        for signal in original_signals:
            parent.remove(signal)

        for channel in channels:
            renamed_paths: dict[str, str] = {}
            clones = [deepcopy(signal) for signal in original_signals]
            for clone, signal_name in zip(clones, ("DST_SIG", "DST_SIG_2"), strict=True):
                old_path = f"/Cfg/Com/ComConfig/{signal_name}_oDST_MSG_oDST_Tx"
                new_name = f"{signal_name}_oDST_MSG_o{channel}_Tx"
                clone.find(f"{{{namespace}}}SHORT-NAME").text = new_name
                renamed_paths[old_path] = f"/Cfg/Com/ComConfig/{new_name}"
                parent.append(clone)
            ipdu = deepcopy(original_ipdu)
            ipdu.find(f"{{{namespace}}}SHORT-NAME").text = f"DST_MSG_o{channel}_Tx"
            for value_ref in ipdu.iter(f"{{{namespace}}}VALUE-REF"):
                if value_ref.text in renamed_paths:
                    value_ref.text = renamed_paths[value_ref.text]
            parent.append(ipdu)

        tree.write(str(path), encoding="UTF-8", xml_declaration=True, pretty_print=True)
        return path

    return create


@pytest.fixture
def generated_arxml_factory(
    tmp_path: Path, workbook_factory: Callable[..., Path], arxml_factory: Callable[..., Path],
) -> Callable[..., Path]:
    """先用第02轮 ADD 生成结构完整的删除测试基线，避免手写被测路由 XML。"""
    counter = 0

    def create(
        *,
        direct_rows: Iterable[Mapping[str, object]] = (),
        signal_rows: Iterable[Mapping[str, object]] = (),
        reference_rows: Iterable[Mapping[str, object]] | None = None,
    ) -> Path:
        nonlocal counter
        counter += 1
        config = workbook_factory(
            filename=f"delete_baseline_{counter}_v4.84.xlsx",
            direct_rows=direct_rows,
            signal_rows=signal_rows,
            reference_rows=reference_rows,
        )
        baseline = arxml_factory(filename=f"delete_source_{counter}.arxml")
        output = tmp_path / f"delete_generated_{counter}.arxml"
        report = generate_inputs(config, baseline, output)
        assert report.is_success, [issue.message for issue in report.all_issues]
        return output

    return create
