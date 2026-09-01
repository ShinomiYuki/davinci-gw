"""使用本地真实输入执行的慢速集成测试。"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
import warnings
from openpyxl import load_workbook

from davinci_gw.application.facade import GatewayFacade
from davinci_gw.application.generate import generate_inputs
from davinci_gw.application.preview import preview_inputs
from davinci_gw.application.validate import inspect_baseline
from davinci_gw.arxml.document import ArxmlDocument
from davinci_gw.contracts import OperationStatus, UpdateRequestDto
from davinci_gw.input.workbook_reader import read_workbook
from davinci_gw.modules import definitions as defs
from davinci_gw.modules.common import semantic_values
from davinci_gw.routing.routing_group_membership import RoutingGroupMembershipService


ROOT = Path(__file__).resolve().parents[2]
CONFIG = ROOT / "input" / "网关路由配置表_v4.84.xlsx"
BASELINE = ROOT / "input" / "T13J.arxml"


def fingerprint(path: Path) -> tuple[int, int, str]:
    """计算基线大小、纳秒时间戳和 SHA-256。"""
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    stat = path.stat()
    return stat.st_size, stat.st_mtime_ns, digest.hexdigest()


def make_add_only_copy(source: Path, target: Path) -> Path:
    """在 tmp_path 中删除 DELETE 数据行，不修改或另存真实输入。"""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported.*")
        workbook = load_workbook(source)
    try:
        for sheet_name in ("直接报文路由", "信号路由"):
            sheet = workbook[sheet_name]
            headers = {cell.value: cell.column for cell in sheet[1]}
            operation_column = headers["操作类型"]
            delete_rows = [row for row in range(2, sheet.max_row + 1)
                           if str(sheet.cell(row=row, column=operation_column).value or "").strip().upper() == "DELETE"]
            for row in reversed(delete_rows):
                sheet.delete_rows(row)
        workbook.save(target)
    finally:
        workbook.close()
    return target


@pytest.mark.slow
@pytest.mark.skipif(not CONFIG.exists() or not BASELINE.exists(), reason="本地真实输入不存在")
def test_real_inputs_preview_and_full_transaction_generation(tmp_path: Path) -> None:
    baseline_before = fingerprint(BASELINE)
    config_before = fingerprint(CONFIG)
    preview = preview_inputs(CONFIG, BASELINE)
    assert preview.is_valid, [issue.message for issue in preview.validation.issues]
    assert preview.target_version == "4.84"
    assert preview.reference_count == 12
    assert (preview.direct_add_count, preview.direct_delete_count) == (9, 1)
    assert (preview.signal_add_count, preview.signal_delete_count) == (2, 2)

    workbook = read_workbook(CONFIG)
    service = RoutingGroupMembershipService(
        ArxmlDocument.load(BASELINE), reference_data=workbook.data.reference_data,
    )
    mapping, mapping_problems = service.application_group_mapping()
    assert {channel: value[0] for channel, value in mapping.items()} == {
        "BDCAN": "PduRRoutingPathGroup_BDCanApp",
        "CHCAN": "PduRRoutingPathGroup_CHCanApp",
        "DACAN": "PduRRoutingPathGroup_DACanApp",
        "DKCAN": "PduRRoutingPathGroup_DKCanApp",
        "DMCAN": "PduRRoutingPathGroup_DMCanApp",
        "EPCAN": "PduRRoutingPathGroup_EPCanApp",
        "GLCAN": "PduRRoutingPathGroup_GLCanApp",
        "ICCAN": "PduRRoutingPathGroup_ICCanApp",
        "LCCAN": "PduRRoutingPathGroup_LCCanApp",
        "PTCAN": "PduRRoutingPathGroup_PTCanApp",
        "SUCAN": "PduRRoutingPathGroup_SUCanApp",
    }
    assert "DGCAN" not in mapping
    assert not [problem for problem in mapping_problems if not problem.warning]
    warning_text = "\n".join(problem.message for problem in mapping_problems if problem.warning)
    assert "FRZCU_6_DACAN" in warning_text and "BMS_2_G_50B_ICCAN" in warning_text

    inspection = inspect_baseline(BASELINE)
    assert inspection.namespace_uri == "http://autosar.org/schema/r4.0"
    assert inspection.schema_filename == "AUTOSAR_00049.xsd"
    assert {name: info.definition_ref for name, info in inspection.modules.items()} == {
        "CanIf": "/MICROSAR/CanIf", "Com": "/MICROSAR/Com",
        "EcuC": "/MICROSAR/EcuC", "PduR": "/MICROSAR/PduR",
    }
    output = tmp_path / "real_full_transaction.arxml"
    facade = GatewayFacade()
    prepared = facade.prepare(UpdateRequestDto(str(CONFIG), str(BASELINE)))
    assert prepared.status is OperationStatus.SUCCESS, [issue.message for issue in prepared.issues]
    assert prepared.preview is not None
    baseline_warnings = [
        issue for issue in prepared.preview.issues
        if issue.code == "PDUR_ROUTING_GROUP_NON_MAIN_MEMBER"
    ]
    assert len(baseline_warnings) == 2
    assert all(
        issue.file_path == str(BASELINE) and issue.sheet_name is None
        and issue.row_number is None and not issue.messages
        for issue in baseline_warnings
    )
    assert all("目标 PduRDestPdu" not in issue.message for issue in baseline_warnings)
    features = {item.feature_id: item for item in prepared.preview.features}
    direct = {item.key: item.value for item in features["direct_message"].metrics}
    signal = {item.key: item.value for item in features["signal_route"].metrics}
    assert (direct["requested_add"], direct["requested_delete"]) == (9, 1)
    assert (signal["requested_add"], signal["requested_delete"]) == (2, 2)
    result = facade.commit_prepared(prepared.session_id, str(output))
    assert result.status is OperationStatus.SUCCESS, [issue.message for issue in result.issues]
    assert inspect_baseline(output).schema_filename == "AUTOSAR_00049.xsd"
    output_document = ArxmlDocument.load(output)
    output_index = output_document.build_index()
    for signal_path in (
        "/ActiveEcuC/Com/ComConfig/TMS_ModeAdjustDisplaySts_oTMS_3_oE0X_PT_CarFLZCU_VCU_GLMessagelis_db480268_Rx",
        "/ActiveEcuC/Com/ComConfig/TMS_ModeAdjustDisplaySts_oFLZCU_44_oE0X_PT_CarFLZCU_VCU_BDMessagelis_a723ce90_Tx",
        "/ActiveEcuC/Com/ComConfig/TMS_ZoneSelectionDisplaySts_oTMS_3_oE0X_PT_CarFLZCU_VCU_GLMessagelis_d2e4d783_Rx",
        "/ActiveEcuC/Com/ComConfig/TMS_ZoneSelectionDisplaySts_oFLZCU_44_oE0X_PT_CarFLZCU_VCU_BDMessagelis_21bb2f41_Tx",
    ):
        parameters, _ = semantic_values(
            output_index.find_by_path(signal_path)[0], output_document.namespace,
        )
        assert parameters[defs.COM_SIGNAL_ACCESS] == ("ACCESS_NEEDED_BY_SWC_OR_COM",)
    assert fingerprint(BASELINE) == baseline_before
    assert fingerprint(CONFIG) == config_before


@pytest.mark.slow
@pytest.mark.skipif(not CONFIG.exists() or not BASELINE.exists(), reason="本地真实输入不存在")
def test_real_add_only_generation_processes_standard_routes(tmp_path: Path) -> None:
    before = fingerprint(BASELINE)
    add_only = make_add_only_copy(CONFIG, tmp_path / "真实路由_add_only_v4.84.xlsx")
    output = tmp_path / "real_generated.arxml"
    report = generate_inputs(add_only, BASELINE, output)
    assert report.is_success, [issue.message for issue in report.all_issues]
    assert report.plan.direct_added_count + report.plan.direct_existing_count + report.plan.direct_skipped_count == 9
    assert report.plan.signal_added_count + report.plan.signal_existing_count + report.plan.signal_skipped_count == 2
    assert report.plan.signal_added_count + report.plan.signal_existing_count == 2
    assert report.plan.signal_skipped_count == 0
    assert inspect_baseline(output).schema_filename == "AUTOSAR_00049.xsd"
    assert fingerprint(BASELINE) == before
